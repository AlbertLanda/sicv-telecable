"""El historial de pagos del abonado, exportado a Excel o PDF.

Exporta **todas** las filas del historial, no solo la página que se ve: el
botón dice «Exportar todas las filas», y un archivo con quince filas de
ochenta se leería como si el abonado hubiera pagado quince veces.

Lleva las columnas de la tabla en su mismo orden. Dos datos que en pantalla
viven a un pase del cursor -el número de operación y la observación o el
motivo de la anulación- van aquí en columnas propias, porque en un archivo no
hay cursor que los muestre.
"""

from io import BytesIO
from xml.sax.saxutils import escape

from django.utils import timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


# El azul del SICV, el mismo de los demás reportes impresos.
TC_BLUE = "0F4C7F"

DATE_FORMAT = "%d/%m/%Y"
STAMP_FORMAT = "%d/%m/%Y %H:%M"

TITLE = "Historial de pagos"

# Los archivos se abren normalmente en Excel. Referencias, observaciones y
# nombres vienen de datos operativos y no deben convertirse en fórmulas si
# empiezan por =, +, -, @ o caracteres de control.
_EXCEL_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r", "\n")


def safe_excel_text(value):
    """Fuerza una cadena a texto seguro para una celda de Excel."""
    text = "" if value is None else str(value)

    if text.startswith(_EXCEL_FORMULA_PREFIXES):
        return "'" + text

    return text


# Las columnas de la tabla, en su orden, más referencia y observación.
COLUMNS = [
    {"key": "date", "label": "Fecha", "width": 11},
    {"key": "detail", "label": "Detalle", "width": 30},
    {"key": "period", "label": "Periodo", "width": 12},
    {"key": "method", "label": "Método", "width": 13},
    {"key": "reference", "label": "Referencia", "width": 14},
    {"key": "amount", "label": "Monto", "width": 11, "numeric": True},
    {"key": "due_date", "label": "Vencimiento", "width": 12},
    {"key": "status", "label": "Estado", "width": 10},
    {"key": "receipt", "label": "Comprobante", "width": 15},
    {"key": "observation", "label": "Observación", "width": 30},
]

FORMATS = ("excel", "pdf")

CONTENT_TYPES = {
    "excel": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}


def _format_date(value):
    if not value:
        return ""

    if hasattr(value, "hour"):
        value = timezone.localtime(value)

    return value.strftime(DATE_FORMAT)


def export_row(row):
    """Una fila de la tabla del historial, lista para el archivo."""
    payment = row["payment"]
    receipt = row["receipt"]

    return {
        "date": _format_date(row["date"]),
        "detail": row["detail"],
        "period": row["period"] or "",
        "method": row["method"],
        "reference": payment.reference,
        "amount": row["amount"],
        "due_date": _format_date(row["due_date"]),
        "status": row["status"],
        "receipt": receipt.full_number if receipt is not None else "",
        "observation": row["status_title"],
    }


def filename(customer, extension):
    hoy = timezone.localdate().strftime("%Y%m%d")

    return f"historial_pagos_{customer.code}_{hoy}.{extension}"


def _stamp(printed_by):
    return f"{printed_by} {timezone.localtime().strftime(STAMP_FORMAT)}".strip()


def render_excel(customer, rows, printed_by=""):
    book = Workbook()
    sheet = book.active
    sheet.title = "Historial de pagos"

    last = len(COLUMNS)
    border = Border(
        left=Side(style="thin", color="D9E4EF"),
        right=Side(style="thin", color="D9E4EF"),
        top=Side(style="thin", color="D9E4EF"),
        bottom=Side(style="thin", color="D9E4EF"),
    )

    sheet.cell(row=1, column=1, value=safe_excel_text(_stamp(printed_by))).font = Font(bold=True, size=8)
    code = sheet.cell(row=1, column=last, value=safe_excel_text(customer.code))
    code.font = Font(bold=True, size=11)
    code.alignment = Alignment(horizontal="right")

    for row, (text, font) in enumerate(
        [
            (TITLE.upper(), Font(bold=True, size=12)),
            (f"Abonado: {customer}", Font(bold=True, size=10)),
        ],
        start=2,
    ):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last)
        cell = sheet.cell(row=row, column=1, value=safe_excel_text(text))
        cell.font = font
        cell.alignment = Alignment(horizontal="center")

    header_row = 5

    for index, column in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=header_row, column=index, value=column["label"])
        cell.fill = PatternFill("solid", fgColor=TC_BLUE)
        cell.font = Font(bold=True, color="FFFFFF", size=9)
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.column_dimensions[get_column_letter(index)].width = column["width"]

    for offset, row in enumerate(rows, start=header_row + 1):
        data = export_row(row)

        for index, column in enumerate(COLUMNS, start=1):
            value = data[column["key"]]

            # El monto se mantiene numérico para poder sumar; cualquier otro
            # valor se fuerza a texto para que Excel no lo evalúe como fórmula.
            if column["key"] == "amount":
                value = float(value)
            else:
                value = safe_excel_text(value)

            cell = sheet.cell(row=offset, column=index, value=value)
            cell.font = Font(size=9)
            cell.border = border
            cell.alignment = Alignment(
                horizontal="right" if column.get("numeric") else "left",
                vertical="top",
                wrap_text=column["key"] in ("detail", "observation"),
            )

            if column["key"] == "amount":
                cell.number_format = "#,##0.00"

    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
    sheet.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(last)}{header_row + len(rows)}"
    )

    buffer = BytesIO()
    book.save(buffer)
    buffer.seek(0)

    return buffer, filename(customer, "xlsx")


def render_pdf(customer, rows, printed_by=""):
    """En A4 apaisado: son diez columnas."""
    page = landscape(A4)
    margin = 10 * mm
    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=page,
        leftMargin=margin,
        rightMargin=margin,
        topMargin=margin,
        bottomMargin=margin,
        title=f"{TITLE} - {customer}",
    )
    width = page[0] - margin * 2

    bold = ParagraphStyle("bold", fontName="Helvetica-Bold", fontSize=10)
    cell = ParagraphStyle("cell", fontName="Helvetica", fontSize=7, leading=8.5)
    head = ParagraphStyle(
        "head", fontName="Helvetica-Bold", fontSize=7, leading=8.5,
        textColor=colors.white,
    )

    strip = Table(
        [[
            Paragraph(escape(_stamp(printed_by)), ParagraphStyle("stamp", fontSize=7.5)),
            Paragraph(f"<para align='center'>{TITLE}</para>", bold),
            Paragraph(f"<para align='right'>{escape(customer.code)}</para>", bold),
        ]],
        colWidths=[width * 0.25, width * 0.5, width * 0.25],
    )
    strip.setStyle(TableStyle([
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
    ]))

    total_weight = sum(column["width"] for column in COLUMNS)
    widths = [width * column["width"] / total_weight for column in COLUMNS]

    data = [[Paragraph(column["label"], head) for column in COLUMNS]]

    for row in rows:
        values = export_row(row)
        data.append([
            Paragraph(
                escape(
                    f"S/ {values['amount']:,.2f}"
                    if column["key"] == "amount"
                    else str(values[column["key"]] or "")
                ),
                cell,
            )
            for column in COLUMNS
        ])

    table = Table(data, colWidths=widths, repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{TC_BLUE}")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D9E4EF")),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]

    for index in range(2, len(data), 2):
        style.append(
            ("BACKGROUND", (0, index), (-1, index), colors.HexColor("#F5F9FD"))
        )

    table.setStyle(TableStyle(style))

    document.build([
        strip,
        Spacer(1, 4),
        Paragraph(escape(f"Abonado: {customer}"), bold),
        Spacer(1, 8),
        table,
    ])
    buffer.seek(0)

    return buffer, filename(customer, "pdf")


RENDERERS = {
    "excel": render_excel,
    "pdf": render_pdf,
}
