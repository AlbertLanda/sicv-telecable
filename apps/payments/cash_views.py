from io import BytesIO
from decimal import Decimal

from django.contrib import messages
from django.contrib.auth.mixins import LoginRequiredMixin
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.paginator import Paginator
from django.http import FileResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views import View
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from apps.organization.context_processors import get_active_branch, get_active_office
from .cash import (add_movement, open_session, operate_offices, reverse_movement, review_close,
                   review_offices, snapshot_session, submit_close, visible_sessions)
from .cash_forms import CashCountForm, CashMovementForm, CashReversalForm, CashReviewForm, OpenCashForm
from .models import CashSession, Payment


class CashAccessMixin(LoginRequiredMixin):
    def dispatch(self, request, *args, **kwargs):
        if request.user.is_authenticated and not (
            request.user.has_perm("payments.operate_cash") or request.user.has_perm("payments.review_cash")
        ):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)


class CashWorkspaceView(CashAccessMixin, View):
    def get(self, request):
        return self.display(request)

    def display(self, request, form=None):
        office = get_active_office(request)
        sessions = visible_sessions(request.user, get_active_branch(request))
        status = request.GET.get("status", "")
        if status in CashSession.Status.values:
            sessions = sessions.filter(status=status)
        previous = CashSession.objects.filter(office=office, cashier=request.user, status="APPROVED").first()
        previous_close = previous.closes.first() if previous else None
        form = form or OpenCashForm(initial={"opening_amount": previous_close.declared_cash if previous_close else "0.00"})
        return render(request, "payments/cash/workspace.html", {
            "page": Paginator(sessions, 30).get_page(request.GET.get("page")), "form": form, "office": office,
            "can_open": office and operate_offices(request.user, get_active_branch(request)).filter(pk=office.pk).exists(),
            "statuses": CashSession.Status.choices, "status_filter": status, "today": timezone.localdate(),
        })

    def post(self, request):
        if not request.user.has_perm("payments.operate_cash"):
            raise PermissionDenied
        office = get_active_office(request)
        form = OpenCashForm(request.POST)
        if office is None or office.is_deposit:
            messages.error(request, "Seleccione una oficina física autorizada en la barra superior.")
        elif form.is_valid():
            try:
                session = open_session(actor=request.user, office=office, **form.cleaned_data)
            except ValidationError as exc:
                form.add_error(None, exc)
            else:
                messages.success(request, "Caja abierta. Los cobros de hoy de esta oficina y cajero están incorporados.")
                return redirect("payments:cash_detail", pk=session.pk)
        return self.display(request, form)


class CashDetailView(CashAccessMixin, View):
    def session(self, request, pk):
        return get_object_or_404(visible_sessions(request.user, get_active_branch(request)), pk=pk)

    def get(self, request, pk):
        return self.display(request, self.session(request, pk))

    def display(self, request, session, bound_action=None, bound_form=None):
        latest = session.closes.first()
        snapshot = latest.snapshot if latest and session.status != "OPEN" else snapshot_session(session)
        forms = {
            "movement": CashMovementForm(office=session.office), "count": CashCountForm(),
            "review": CashReviewForm(initial={"close_id": latest.pk if latest else None}), "reverse": CashReversalForm(),
        }
        if bound_form is not None:
            forms[bound_action] = bound_form
        can_operate = session.cashier_id == request.user.pk and request.user.has_perm("payments.operate_cash")
        can_review = session.cashier_id != request.user.pk and review_offices(request.user).filter(pk=session.office_id).exists()
        methods = [{"name": dict(Payment.Method.choices).get(key, key), "total": value} for key, value in snapshot["methods"].items()]
        return render(request, "payments/cash/detail.html", {
            "cash": session, "snapshot": snapshot, "latest": latest, "forms": forms,
            "can_operate": can_operate and session.status == "OPEN", "can_review": can_review and session.status == "SUBMITTED",
            "entry_page": Paginator(snapshot["entries"], 50).get_page(request.GET.get("page")),
            "methods": methods, "events": session.events.select_related("actor", "close"), "closes": session.closes.all(),
        })

    def post(self, request, pk):
        session = self.session(request, pk)
        action = request.POST.get("action")
        if action == "movement":
            form = CashMovementForm(request.POST, office=session.office)
        elif action == "count":
            form = CashCountForm(request.POST)
        elif action == "review":
            form = CashReviewForm(request.POST)
        elif action == "reverse":
            form = CashReversalForm(request.POST)
        else:
            raise PermissionDenied
        if form.is_valid():
            try:
                values = dict(form.cleaned_data)
                if action == "movement":
                    values["issuer_id"] = values["issuer_id"].pk
                    add_movement(session_id=pk, actor=request.user, **values)
                elif action == "count":
                    submit_close(session_id=pk, actor=request.user, counts=form.counts(), explanation=values["explanation"])
                elif action == "reverse":
                    reverse_movement(session_id=pk, actor=request.user, **values)
                else:
                    review_close(session_id=pk, actor=request.user, **values)
            except ValidationError as exc:
                form.add_error(None, exc)
            else:
                messages.success(request, "Operación de caja registrada.")
                return redirect("payments:cash_detail", pk=pk)
        return self.display(request, session, action, form)


class CashExportView(CashDetailView):
    http_method_names = ["get", "head", "options"]

    def get(self, request, pk):
        session = self.session(request, pk)
        raw = request.GET.get("revision")
        if raw:
            try:
                revision = int(raw)
            except ValueError:
                revision = 0
            close = get_object_or_404(session.closes, revision=revision)
        else:
            close = session.closes.first() if session.status != "OPEN" else None
        snapshot = close.snapshot if close else snapshot_session(session)
        book = Workbook()
        sheet = book.active
        sheet.title = "Caja"
        # Todo texto es literal, incluso motivos/referencias que empiezan por =.
        def row(sheet, values):
            sheet.append(values)
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        for values in [
            ["Caja operativa PEN", snapshot["session_id"]], ["Fecha de caja", snapshot["date"]],
            ["Oficina", snapshot["office"]], ["Cajero", snapshot["cashier"]],
            ["Evidencia", f"Versión enviada {close.revision}" if close else "Provisional: caja abierta"],
            ["Huella de la versión", close.digest if close else ""],
            ["Fondo inicial", Decimal(snapshot["opening"])], ["Origen de apertura", snapshot["opening_reason"]],
            ["Efectivo esperado", Decimal(snapshot["expected_cash"])], ["Efectivo contado", Decimal(snapshot["declared_cash"]) if close else "Sin arqueo"],
            ["Diferencia", Decimal(snapshot["difference"]) if close else "Sin arqueo"], ["Explicación", snapshot.get("explanation", "")],
            ["Alcance", "Caja física; otros medios no se suman al efectivo. No acredita conciliación bancaria ni fiscal."],
        ]:
            row(sheet, values)
        movements = book.create_sheet("Movimientos")
        columns = [("id", "ID"), ("kind", "Tipo"), ("occurred_at", "Fecha efectiva"), ("recorded_at", "Fecha registro"),
                   ("amount", "Importe firmado PEN"), ("method", "Medio"), ("issuer", "Empresa"), ("ruc", "RUC"),
                   ("description", "Motivo"), ("reference", "Referencia"), ("bank", "Banco"), ("account", "Cuenta"),
                   ("document_type", "Tipo sustento"), ("document_number", "Número sustento"), ("actor", "Registrado por"),
                   ("payment_id", "Pago ERP"), ("reversal_of", "Revierte ID"), ("cash_balance", "Saldo efectivo")]
        row(movements, [label for _, label in columns])
        for entry in snapshot["entries"]:
            row(movements, [Decimal(entry[key]) if key in ("amount", "cash_balance") else entry.get(key, "") for key, _ in columns])
        count = book.create_sheet("Arqueo")
        row(count, ["Denominación PEN", "Cantidad"])
        for value, quantity in snapshot.get("counts", {}).items():
            row(count, [Decimal(value), quantity])
        methods = book.create_sheet("Medios")
        row(methods, ["Medio", "Movimiento neto PEN (sin fondo inicial)"])
        for method, total in snapshot["methods"].items():
            row(methods, [dict(Payment.Method.choices).get(method, method), Decimal(total)])
        issuers = book.create_sheet("Empresas")
        row(issuers, ["Empresa", "RUC", "Movimiento neto PEN (sin fondo inicial)"])
        for item in snapshot["issuers"]:
            row(issuers, [item["name"], item["ruc"], Decimal(item["net"])])
        audit = book.create_sheet("Revisiones")
        row(audit, ["Fecha", "Acción", "Persona", "Versión", "Comentario"])
        for event in session.events.select_related("actor", "close"):
            row(audit, [event.created_at.isoformat(), event.action, str(event.actor), event.close.revision if event.close else "", event.reason])
        for sheet in book:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="174C5B")
            for column in sheet.columns:
                sheet.column_dimensions[column[0].column_letter].width = min(55, max(18, max(len(str(c.value or "")) for c in column) + 2))
                for cell in column:
                    if isinstance(cell.value, Decimal):
                        cell.number_format = '#,##0.00;[Red]-#,##0.00'
        output = BytesIO()
        book.save(output)
        output.seek(0)
        return FileResponse(output, as_attachment=True, filename=f"caja-{session.pk}-v{close.revision if close else 'provisional'}.xlsx")
