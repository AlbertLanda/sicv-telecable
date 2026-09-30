"""El botón Exportar del historial de pagos: Excel y PDF con todas las filas."""

from decimal import Decimal
from io import BytesIO

from django.utils import timezone

from openpyxl import load_workbook

from apps.payments.models import Payment
from apps.payments.services import register_payment
from apps.reports.tests.pdf_text import pdf_text

from .test_historial_y_comprobante import HistorialBase


class HistoryExportTests(HistorialBase):
    def pay_october(self, **overrides):
        data = dict(
            customer=self.customer,
            amount=Decimal("65.00"),
            method=Payment.Method.YAPE,
            reference="OP-5521",
            branch=self.branch,
            user=self.cashier,
            allocations=[(self.octubre, Decimal("65.00"))],
        )
        data.update(overrides)

        return register_payment(**data)

    def advance(self, amount="10.00"):
        """Un pago a cuenta, sin deuda: una fila del historial."""
        return Payment.objects.create(
            customer=self.customer,
            amount=Decimal(amount),
            method=Payment.Method.CASH,
            branch=self.branch,
            received_by=self.cashier,
            paid_at=timezone.now(),
        )

    def export(self, export_format, **extra):
        self.login(self.viewer)

        return self.client.get(self.url(), {"exportar": export_format, **extra})

    def sheet(self, response):
        book = load_workbook(BytesIO(b"".join(response.streaming_content)))

        return book["Historial de pagos"]

    def test_the_button_offers_excel_and_pdf(self):
        self.pay_october()
        self.login(self.viewer)

        response = self.client.get(self.url())

        self.assertContains(response, "Exportar todas las filas")
        self.assertContains(response, "exportar=excel")
        self.assertContains(response, "exportar=pdf")

    def test_without_payments_there_is_nothing_to_export(self):
        self.login(self.viewer)

        response = self.client.get(self.url())

        self.assertNotContains(response, "Exportar todas las filas")

    def test_excel_carries_the_columns_of_the_table(self):
        self.pay_october()

        response = self.export("excel")

        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("attachment", response["Content-Disposition"])
        self.assertIn(f"historial_pagos_{self.customer.code}_", response["Content-Disposition"])

        sheet = self.sheet(response)
        headers = [cell.value for cell in sheet[5]]

        self.assertEqual(
            headers,
            [
                "Fecha", "Detalle", "Periodo", "Método", "Referencia", "Monto",
                "Vencimiento", "Estado", "Comprobante", "Observación",
            ],
        )

        row = {header: cell.value for header, cell in zip(headers, sheet[6])}

        self.assertEqual(row["Método"], "Yape")
        self.assertEqual(row["Referencia"], "OP-5521")
        self.assertEqual(row["Monto"], 65.0)
        self.assertEqual(row["Estado"], "Pagado")
        self.assertTrue(row["Comprobante"])

    def test_excel_exports_every_row_not_just_the_page(self):
        for _ in range(20):
            self.advance()

        self.login(self.viewer)
        page = self.client.get(self.url())
        self.assertEqual(len(page.context["rows"]), 15)

        sheet = self.sheet(self.export("excel", page="2"))
        amounts = [
            sheet.cell(row=row, column=6).value
            for row in range(6, sheet.max_row + 1)
            if sheet.cell(row=row, column=6).value is not None
        ]

        self.assertEqual(len(amounts), 20)

    def test_a_voided_row_carries_its_reason(self):
        payment, _receipt = self.pay_october()
        payment.void(self.cashier, "Cobro duplicado")

        sheet = self.sheet(self.export("excel"))
        headers = [cell.value for cell in sheet[5]]
        row = {header: cell.value for header, cell in zip(headers, sheet[6])}

        self.assertEqual(row["Estado"], "Anulado")
        self.assertIn("Cobro duplicado", row["Observación"])

    def test_pdf_opens_in_the_browser_with_the_rows(self):
        self.pay_october()

        response = self.export("pdf")
        data = b"".join(response.streaming_content)

        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertIn("inline", response["Content-Disposition"])
        self.assertTrue(data.startswith(b"%PDF"))

        text = pdf_text(data)
        self.assertIn("Historial de pagos", text)
        self.assertIn(self.customer.code, text)
        self.assertIn("OP-5521", text)

    def test_who_cannot_see_the_history_cannot_export_it(self):
        self.pay_october()
        stranger = self.make_user("sinpermiso9")
        self.login(stranger)

        response = self.client.get(self.url(), {"exportar": "excel"})

        self.assertEqual(response.status_code, 403)

    def test_an_unknown_format_shows_the_history(self):
        self.pay_october()

        response = self.export("word")

        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "payments/customer_payment_history.html")
