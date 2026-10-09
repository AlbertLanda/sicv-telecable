"""Invariantes monetarias, evidencia, límites de acceso y carreras de caja."""
from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta
from decimal import Decimal
from io import BytesIO
from threading import Barrier
from unittest import skipUnless
from unittest.mock import patch
import uuid

from django.contrib import admin
from django.contrib.auth import get_user_model
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import close_old_connections, connection, connections
from django.test import TransactionTestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.customers.models import Customer
from apps.organization.context_processors import ACTIVE_BRANCH_SESSION_KEY, ACTIVE_OFFICE_SESSION_KEY
from apps.organization.models import Branch, Office
from apps.payments.cash import (DENOMINATIONS, add_movement, open_session, reverse_movement,
                               review_close, snapshot_session, submit_close, visible_sessions)
from apps.payments.models import CashClose, CashEntry, CashEvent, CashSession, Issuer, OfficeSequence, Payment, Receipt, ReceiptSequence
from apps.payments.services import register_payment
from apps.payments.tests.base import PaymentsTestCase


class CashFixture:
    def setup_cash(self):
        self.office = Office.objects.create(branch=self.branch, code="CASH-OFF", name="Oficina de prueba")
        self.actor = get_user_model().objects.create_user(username="cash-operator", role="ATC", branch=self.branch, office=self.office)
        self.reviewer = get_user_model().objects.create_user(username="cash-reviewer", role="ACCOUNTING", branch=self.branch)
        self.issuer = Issuer.objects.create(code="CASH-TEST", business_name="Emisor de prueba", ruc="20000000001")
        self.sequence = ReceiptSequence.objects.create(code="CASH-TEST", series="TEST", label="Prueba", issuer=self.issuer, autonumber=True)
        OfficeSequence.objects.create(office=self.office, sequence=self.sequence)

    def open(self, **overrides):
        values = dict(actor=self.actor, office=self.office, opening_amount=Decimal("100.00"), reason="Fondo entregado para prueba")
        values.update(overrides)
        return open_session(**values)

    def payment(self, **overrides):
        values = dict(customer=self.customer, amount=Decimal("80"), method="CASH", branch=self.branch,
                      user=self.actor, series=self.sequence, office=self.office)
        values.update(overrides)
        return register_payment(**values)[0]

    def movement(self, session, **overrides):
        values = dict(session_id=session.pk, actor=self.actor, request_key=uuid.uuid4(), kind="EXPENSE", amount=Decimal("10"),
                      issuer_id=self.issuer.pk, description="Movilidad de prueba", document_type="Sustento interno", document_number="TEST-1")
        values.update(overrides)
        return add_movement(**values)

    def submit(self, session, counts=None, **overrides):
        values = dict(session_id=session.pk, actor=self.actor, counts=counts or {value: (1 if value == "100.00" else 0) for value in DENOMINATIONS})
        values.update(overrides)
        return submit_close(**values)

    def approve(self, session, close):
        return review_close(session_id=session.pk, close_id=close.pk, actor=self.reviewer, decision="APPROVED", reason="Revisado contra efectivo y sustentos")


class CashSessionTests(CashFixture, PaymentsTestCase):
    def setUp(self):
        super().setUp()
        self.setup_cash()

    def test_cash_and_non_cash_are_separated_and_deposit_is_an_outflow(self):
        session = self.open()
        self.payment()
        self.payment(amount=30, method="YAPE", reference="TEST-YAPE")
        self.movement(session)
        self.movement(session, kind="DEPOSIT", amount=50, bank="Banco de prueba", account="Cuenta de prueba", reference="OP-1")
        counts = {value: (1 if value in ("100.00", "20.00") else 0) for value in DENOMINATIONS}
        close = self.submit(session, counts=counts)
        self.assertEqual((close.expected_cash, close.declared_cash, close.difference), (120, 120, 0))
        self.assertEqual(Decimal(close.snapshot["methods"]["YAPE"]), 30)
        self.assertEqual(len(close.snapshot["entries"]), 4)
        self.approve(session, close)
        session.refresh_from_db()
        self.assertEqual(session.status, "APPROVED")

    def test_open_adopts_today_only_and_excludes_pending_voided_and_other_cashier(self):
        previous = self.payment(paid_at=timezone.now() - timedelta(days=1))
        paid = self.payment()
        pending = self.payment(settled=False)
        voided = self.payment()
        voided.void(user=self.actor, reason="Prueba")
        other = self.payment(user=self.cashier)
        session = self.open()
        self.assertEqual(list(session.entries.values_list("payment_id", flat=True)), [paid.pk])
        self.assertEqual(Decimal(snapshot_session(session)["expected_cash"]), 180)
        pending.confirm(actor=self.actor)
        self.assertEqual(session.entries.count(), 2)
        self.assertFalse(CashEntry.objects.filter(payment__in=[previous, voided, other]).exists())

    def test_pending_not_counted_until_confirmation_and_void_reverses_once(self):
        session = self.open()
        payment = self.payment(settled=False)
        self.assertFalse(session.entries.exists())
        payment.confirm(actor=self.actor)
        payment.void(user=self.actor, reason="Cobro duplicado")
        self.assertEqual(Decimal(snapshot_session(session)["expected_cash"]), 100)
        self.assertEqual(list(session.entries.values_list("kind", flat=True)), ["PAYMENT", "VOID"])
        with self.assertRaises(ValidationError):
            payment.void(user=self.actor, reason="Reintento")
        self.assertEqual(session.entries.count(), 2)

    def test_submitted_cash_blocks_new_payment_confirm_void_and_manual_edits(self):
        session = self.open()
        paid = self.payment()
        pending = self.payment(settled=False)
        self.submit(session, explanation="Conteo de prueba con diferencia")
        number = self.sequence.receipts.count()
        for action in [self.payment, lambda: pending.confirm(actor=self.actor),
                       lambda: paid.void(user=self.actor, reason="Prueba"), lambda: self.movement(session)]:
            with self.assertRaises(ValidationError):
                action()
        self.assertEqual(self.sequence.receipts.count(), number)
        paid.refresh_from_db()
        pending.refresh_from_db()
        self.assertEqual(paid.status, "REGISTERED")
        self.assertEqual(pending.status, "PENDING")
        self.assertEqual(session.entries.count(), 1)

    def test_return_and_resubmit_retains_old_evidence_and_stale_review_is_rejected(self):
        session = self.open()
        first = self.submit(session)
        original = first.snapshot
        review_close(session_id=session.pk, close_id=first.pk, actor=self.reviewer, decision="RETURNED", reason="Falta el gasto")
        self.movement(session)
        second = self.submit(session, explanation="Se contó el fondo sin descontar el gasto")
        with self.assertRaises(ValidationError):
            self.approve(session, first)
        self.approve(session, second)
        first.refresh_from_db()
        self.assertEqual(first.snapshot, original)
        self.assertEqual((first.revision, second.revision), (1, 2))
        self.assertEqual(first.snapshot["entries"], [])
        self.assertEqual(len(second.snapshot["entries"]), 1)

    def test_no_self_approval_even_superuser(self):
        session = self.open()
        close = self.submit(session)
        self.actor.is_superuser = True
        self.actor.save()
        with self.assertRaises(PermissionDenied):
            review_close(session_id=session.pk, close_id=close.pk, actor=self.actor, decision="APPROVED", reason="Autorrevisión")

    def test_difference_requires_explanation_and_count_is_strict(self):
        session = self.open()
        self.payment()
        with self.assertRaises(ValidationError):
            self.submit(session)
        for bad in (-1, True, 1.5, 1000001):
            counts = dict.fromkeys(DENOMINATIONS, 0)
            counts["10.00"] = bad
            with self.assertRaises(ValidationError):
                self.submit(session, counts=counts, explanation="Prueba")
        self.assertEqual(CashClose.objects.count(), 0)

    def test_movement_replay_and_conflicting_reuse(self):
        session = self.open()
        key = uuid.uuid4()
        entry = self.movement(session, request_key=key)
        self.assertEqual(self.movement(session, request_key=key).pk, entry.pk)
        with self.assertRaises(ValidationError):
            self.movement(session, request_key=key, amount=20)
        self.assertEqual(session.entries.count(), 1)

    def test_duplicate_deposit_and_missing_sustento_are_rejected(self):
        session = self.open()
        for values in [dict(document_number=""), dict(document_type=""), dict(kind="DEPOSIT"), dict(amount="NaN"), dict(amount="1.001")]:
            with self.assertRaises(ValidationError):
                self.movement(session, **values)
        data = dict(kind="DEPOSIT", bank="Banco", account="12345", reference="OP-2")
        self.movement(session, **data)
        with self.assertRaises(ValidationError):
            self.movement(session, **data)
        self.assertEqual(session.entries.count(), 1)

    def test_cannot_use_another_office_issuer(self):
        session = self.open()
        issuer = Issuer.objects.create(code="OTHER", business_name="Otra empresa", ruc="20000000002")
        with self.assertRaises(ValidationError):
            self.movement(session, issuer_id=issuer.pk)

    def test_manual_reversal_keeps_original_and_cannot_reverse_payment(self):
        session = self.open()
        entry = self.movement(session)
        payment = self.payment()
        reverse_movement(session_id=session.pk, entry_id=entry.pk, actor=self.actor, reason="Dinero repuesto")
        self.assertEqual(Decimal(snapshot_session(session)["expected_cash"]), 180)
        entry.refresh_from_db()
        self.assertEqual(entry.amount, -10)
        for entry_id in [entry.pk, CashEntry.objects.get(payment=payment).pk]:
            with self.assertRaises(ValidationError):
                reverse_movement(session_id=session.pk, entry_id=entry_id, actor=self.actor, reason="Reintento")

    def test_next_day_carries_counted_balance_and_reversal_does_not_change_old_close(self):
        session = self.open()
        entry = self.movement(session)
        close = self.submit(session, explanation="Diferencia reconocida de prueba")
        self.approve(session, close)
        snapshot = close.snapshot
        future = timezone.now() + timedelta(days=1)
        with patch("django.utils.timezone.now", return_value=future):
            with self.assertRaises(ValidationError):
                self.open(opening_amount=99, reason="")
            next_session = self.open(reason="")
            self.assertEqual(next_session.previous_id, session.pk)
            reverse_movement(session_id=next_session.pk, entry_id=entry.pk, actor=self.actor, reason="Reintegro efectivo del gasto")
            self.assertEqual(Decimal(snapshot_session(next_session)["expected_cash"]), 110)
        close.refresh_from_db()
        self.assertEqual(close.snapshot, snapshot)

    def test_opening_once_and_no_new_day_before_approval(self):
        session = self.open()
        with self.assertRaises(ValidationError):
            self.open()
        with patch("django.utils.timezone.now", return_value=timezone.now() + timedelta(days=1)):
            with self.assertRaises(ValidationError):
                self.open()
            with self.assertRaises(ValidationError):
                self.payment()
        self.assertEqual(CashSession.objects.count(), 1)
        self.assertEqual(session.entries.count(), 0)

    def test_backdating_cannot_bypass_control_after_activation(self):
        self.open()
        with self.assertRaises(ValidationError):
            self.payment(paid_at=timezone.now() - timedelta(days=1))

    def test_open_requires_authorized_physical_office_and_operational_permission(self):
        self.actor.office = None
        self.actor.save()
        with self.assertRaises(PermissionDenied):
            self.open()
        self.actor.allowed_offices.add(self.office)
        self.open()
        deposit = Office.objects.create(branch=self.branch, name="Depósito", code="BANK", is_deposit=True)
        for values in [dict(office=deposit), dict(actor=self.reviewer)]:
            with self.assertRaises(PermissionDenied):
                self.open(**values)

    def test_other_cashier_and_foreign_reviewer_cannot_read_or_write(self):
        session = self.open()
        other = self.make_user("other-cash", role="ATC")
        other.allowed_offices.add(self.office)
        self.assertFalse(visible_sessions(other, self.branch).exists())
        with self.assertRaises(PermissionDenied):
            self.movement(session, actor=other)
        foreign = Branch.objects.create(code="FOREIGN", name="Otra sede")
        self.reviewer.branch = foreign
        self.reviewer.save()
        self.assertFalse(visible_sessions(self.reviewer, self.branch).exists())
        close = self.submit(session)
        with self.assertRaises(PermissionDenied):
            self.approve(session, close)
        self.reviewer.allowed_offices.add(self.office)
        self.approve(session, close)

    def test_explicit_operating_grant_still_requires_physical_office_assignment(self):
        operator = self.make_user("accounting-operator", permissions=["operate_cash"])
        with self.assertRaises(PermissionDenied):
            self.open(actor=operator)
        operator.allowed_offices.add(self.office)
        session = self.open(actor=operator)
        self.assertEqual(session.cashier_id, operator.pk)

    def test_roleless_user_and_revoked_office_cannot_access_cash(self):
        session = self.open()
        for role in ("TECHNICIAN", "WAREHOUSE", "SALES"):
            user = self.make_user(f"role-{role}", role=role)
            self.client.force_login(user)
            self.assertEqual(self.client.get(reverse("payments:cash_workspace")).status_code, 403)
        self.actor.office = None
        self.actor.save()
        self.assertFalse(visible_sessions(self.actor, self.branch).exists())
        with self.assertRaises(PermissionDenied):
            self.submit(session)

    def test_issuer_changes_do_not_rewrite_entry_or_closing_evidence(self):
        session = self.open()
        payment = self.payment()
        self.issuer.business_name = "Cambio posterior"
        self.issuer.save()
        self.sequence.issuer = None
        self.sequence.save()
        payment.void(user=self.actor, reason="Prueba con configuración posterior")
        close = self.submit(session, explanation="Prueba")
        self.assertEqual(close.snapshot["entries"][0]["issuer"], "Emisor de prueba")
        self.assertEqual(close.snapshot["entries"][0]["ruc"], "20000000001")
        self.assertEqual(close.snapshot["entries"][1]["issuer"], "Emisor de prueba")
        self.assertEqual(len(close.snapshot["issuers"]), 1)
        self.assertEqual(Decimal(close.snapshot["issuers"][0]["net"]), 0)

    def test_web_forms_review_export_and_formula_strings(self):
        self.client.force_login(self.actor)
        web = self.client.session
        web[ACTIVE_BRANCH_SESSION_KEY], web[ACTIVE_OFFICE_SESSION_KEY] = self.branch.pk, self.office.pk
        web.save()
        response = self.client.post(reverse("payments:cash_workspace"), {"opening_amount": "100", "reason": "Prueba"})
        self.assertEqual(response.status_code, 302)
        session = CashSession.objects.get(cashier=self.actor)
        url = reverse("payments:cash_detail", args=[session.pk])
        self.assertContains(self.client.get(url), "Arqueo y envío a revisión")
        self.assertContains(self.client.get(reverse("payments:cash_workspace")), "Caja operativa")
        self.assertEqual(self.client.post(url, {"action": "movement", "request_key": str(uuid.uuid4()), "issuer_id": self.issuer.pk,
                        "kind": "INCOME", "amount": "20", "description": "=HYPERLINK(\"test\")"}).status_code, 302)
        counts = {f"count_{value}": (1 if value in ("100.00", "20.00") else 0) for value in DENOMINATIONS}
        self.assertEqual(self.client.post(url, {"action": "count", **counts}).status_code, 302)
        close = session.closes.get()
        self.client.force_login(self.reviewer)
        response = self.client.get(url)
        self.assertContains(response, "Revisar cierre")
        self.assertEqual(self.client.post(url, {"action": "review", "close_id": close.pk, "decision": "APPROVED", "reason": "Revisado"}).status_code, 302)
        response = self.client.get(reverse("payments:cash_export", args=[session.pk]) + "?revision=1")
        self.assertEqual(response.status_code, 200)
        book = load_workbook(BytesIO(b"".join(response.streaming_content)))
        self.assertEqual(book["Movimientos"]["I2"].value, '=HYPERLINK("test")')
        self.assertEqual(book["Movimientos"]["I2"].data_type, "s")
        self.assertEqual(book["Arqueo"].max_row, len(DENOMINATIONS) + 1)
        self.assertEqual(self.client.get(reverse("payments:cash_export", args=[session.pk]) + "?revision=999").status_code, 404)

    def test_invalid_post_renders_errors_and_export_denies_foreign_cashier(self):
        session = self.open()
        self.client.force_login(self.actor)
        response = self.client.post(reverse("payments:cash_detail", args=[session.pk]), {"action": "movement"})
        self.assertEqual(response.status_code, 200)
        self.assertIn("amount", response.context["forms"]["movement"].errors)
        self.client.force_login(self.cashier)
        self.cashier.branch = None
        self.cashier.save()
        web = self.client.session
        web[ACTIVE_BRANCH_SESSION_KEY] = self.branch.pk
        web.save()
        self.assertEqual(self.client.get(reverse("payments:cash_export", args=[session.pk])).status_code, 404)

    def test_admin_has_no_edit_or_delete_path_for_evidence(self):
        request = type("Request", (), {"user": self.actor})()
        for model in (CashSession, CashClose, CashEntry, CashEvent):
            control = admin.site._registry[model]
            self.assertFalse(control.has_add_permission(request))
            self.assertFalse(control.has_change_permission(request))
            self.assertFalse(control.has_delete_permission(request))


@skipUnless(connection.vendor == "postgresql", "Bloqueos reales: ejecutado en CI PostgreSQL")
class ConcurrentCashTests(CashFixture, TransactionTestCase):
    def setUp(self):
        self.branch = Branch.objects.create(code="CASH-CON", name="Sede concurrencia")
        self.customer = Customer.objects.create(code="CASH-CON", branch=self.branch, first_name="Prueba", document_number="00000001")
        self.setup_cash()

    def race(self, *actions):
        barrier = Barrier(len(actions))
        def run(action):
            close_old_connections()
            try:
                with connection.cursor() as cursor:
                    cursor.execute("SET lock_timeout = '5s'")
                    cursor.execute("SET statement_timeout = '10s'")
                barrier.wait(timeout=5)
                try:
                    action()
                    return "ok"
                except ValidationError:
                    return "rejected"
            finally:
                connections.close_all()
        with ThreadPoolExecutor(max_workers=len(actions)) as pool:
            futures = [pool.submit(run, action) for action in actions]
            return [future.result(timeout=20) for future in futures]

    def test_simultaneous_opening_is_unique(self):
        self.assertEqual(sorted(self.race(self.open, self.open)), ["ok", "rejected"])
        self.assertEqual(CashSession.objects.count(), 1)

    def test_opening_and_payment_capture_payment_once(self):
        self.assertEqual(self.race(self.open, self.payment), ["ok", "ok"])
        self.assertEqual(CashEntry.objects.count(), 1)
        self.assertEqual(Decimal(snapshot_session(CashSession.objects.get())["expected_cash"]), 180)

    def test_closing_and_payment_do_not_lose_cash(self):
        session = self.open()
        results = self.race(lambda: self.submit(session, explanation="Conteo con revisión"), self.payment)
        self.assertEqual(results[0], "ok")
        close = CashClose.objects.get()
        payments = Payment.objects.count()
        self.assertEqual(close.expected_cash, 100 + 80 * payments)
        self.assertEqual(len(close.snapshot["entries"]), payments)

    def test_same_movement_request_is_written_once(self):
        session = self.open()
        key = uuid.uuid4()
        self.assertEqual(self.race(lambda: self.movement(session, request_key=key), lambda: self.movement(session, request_key=key)), ["ok", "ok"])
        self.assertEqual(CashEntry.objects.count(), 1)

    def test_void_and_close_cannot_mutate_closed_snapshot(self):
        session = self.open()
        payment = self.payment()
        results = self.race(lambda: self.submit(session, explanation="Prueba"), lambda: payment.void(user=self.actor, reason="Prueba"))
        self.assertEqual(results[0], "ok")
        payment.refresh_from_db()
        close = CashClose.objects.get()
        self.assertEqual(close.expected_cash, 100 if payment.status == "VOIDED" else 180)
        self.assertEqual(Decimal(close.snapshot["expected_cash"]), close.expected_cash)

    def test_two_reviewers_cannot_decide_twice(self):
        session = self.open()
        close = self.submit(session)
        self.assertEqual(sorted(self.race(lambda: self.approve(session, close), lambda: review_close(
            session_id=session.pk, close_id=close.pk, actor=self.reviewer, decision="RETURNED", reason="Corregir"))), ["ok", "rejected"])
        self.assertEqual(CashEvent.objects.filter(action__in=["APPROVED", "RETURNED"]).count(), 1)
