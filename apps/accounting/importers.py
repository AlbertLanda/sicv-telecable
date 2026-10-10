"""Bounded, all-or-nothing readers. Never evaluate Excel formulas or infer tax."""
import calendar
import csv
import hashlib
import re
import unicodedata
import zlib
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile
from xml.etree.ElementTree import ParseError

from django.core.exceptions import ValidationError
from django.db import transaction
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from apps.audit.models import AuditEvent
from apps.payments.models import Issuer
from .access import require_access
from .models import Document, ImportBatch, Line

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 15000
ZERO = Decimal("0.00")
CENT = Decimal("0.01")
LEGACY_HEADERS = ["TipoDocumento", "TipoDocumentoNombre", "FechaEmision", "ID_Comprobante", "",
                  "ClienteDocumento", "ClienteNombre", "Estado", "EstadoSunat", "ObservacionEstado",
                  "ItemDescripcion", "Concepto_OSIPTEL", "Tecnologia", "ReglaClasificacion",
                  "ItemCantidad", "ItemValorVenta", "ItemIGV", "ItemTotal"]
RVIE_HEADERS = ["ruc_emisor", "periodo", "tipo", "serie", "numero", "fecha_emision", "moneda",
                "base_imponible", "igv", "total", "estado_sunat", "documento_cliente", "nombre_cliente"]
FLAGS = {
    "NO_SIRE_STATE": "Sin código Est. Comp en el archivo SIRE",
    "TAX_MAPPING": "Base cero e IGV igual al total: verificar fuente",
    "UNKNOWN_CLASS": "Clasificación OSIPTEL por revisar",
    "LOCAL_PENDING": "Pendiente en SICV antiguo",
    "NO_SUNAT_STATE": "Sin estado SUNAT en el archivo",
    "MIXED_METADATA": "Datos o estados distintos dentro del mismo comprobante",
    "REPEATED_LINE": "Líneas idénticas conservadas: verificar si corresponden",
}


def normalize(value):
    return "".join(c for c in unicodedata.normalize("NFKD", str(value or "")).casefold() if c.isalnum())


def text_value(value, limit=1000):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    value = str(value).strip()
    if len(value) > limit or "\x00" in value:
        raise ValidationError("Un campo supera el tamaño permitido o contiene caracteres inválidos.")
    return value


def money(value, *, optional=False, places=2):
    if value in (None, "") and optional:
        return None
    try:
        # Numeric Excel floats sometimes carry binary noise; only that tiny
        # noise is removed. Text must use plain decimals, without separators.
        raw = str(value).strip()
        if isinstance(value, str) and not re.fullmatch(r"-?\d+(?:\.\d+)?", raw):
            raise ValueError
        amount = Decimal(raw)
        rounded = amount.quantize(Decimal(10) ** -places)
        if not amount.is_finite() or abs(amount) >= Decimal(10) ** (18 - places):
            raise ValueError
        if abs(amount - rounded) > (Decimal("0.00000001") if isinstance(value, float) else ZERO):
            raise ValueError
        return rounded
    except (InvalidOperation, ValueError, TypeError):
        raise ValidationError("Importe inválido. Use números decimales sin separadores de miles ni fórmulas.") from None


def parsed_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%Y %H:%M:%S"):
        try:
            return datetime.strptime(str(value), fmt).date()
        except ValueError:
            pass
    raise ValidationError("Fecha inválida. Use DD/MM/AAAA o AAAA-MM-DD.")


def document_key(kind, series, number):
    kind = text_value(kind, 2).zfill(2)
    series = text_value(series, 20).upper()
    number = text_value(number, 20)
    if not re.fullmatch(r"[0-9]{2}", kind) or not re.fullmatch(r"[A-Z0-9-]{1,20}", series):
        raise ValidationError("Tipo o serie del comprobante inválidos.")
    if not re.fullmatch(r"[0-9]{1,20}", number) or int(number) == 0:
        raise ValidationError("El correlativo debe ser un entero positivo.")
    return kind, series, str(int(number))


def checked_zip(raw):
    with ZipFile(BytesIO(raw)) as archive:
        infos = archive.infolist()
        if len(infos) > 250 or sum(i.file_size for i in infos) > 50 * 1024 * 1024:
            raise ValidationError("El Excel descomprimido supera el límite de lectura.")
        for info in infos:
            if info.flag_bits & 1 or "vbaproject" in info.filename.lower():
                raise ValidationError("No se admiten archivos cifrados ni macros.")
            if info.filename.endswith(".xml"):
                content = archive.read(info)
                if b"<!DOCTYPE" in content.upper() or b"<!ENTITY" in content.upper():
                    raise ValidationError("El Excel contiene declaraciones XML no admitidas.")


def legacy_filename(filename, ruc, period):
    match = re.fullmatch(r"OSIPTEL_(\d{11})_(\d{8})_(\d{8}).*\.xlsx", filename, flags=re.I)
    if not match:
        raise ValidationError("Conserve el nombre original OSIPTEL_RUC_AAAAMMDD_AAAAMMDD…xlsx para verificar empresa y periodo.")
    end = period.replace(day=calendar.monthrange(period.year, period.month)[1])
    if match.groups() != (ruc, period.strftime("%Y%m%d"), end.strftime("%Y%m%d")):
        raise ValidationError("El RUC o las fechas del nombre no coinciden. Importe un mes completo de la empresa elegida.")


def parse_osiptel(raw, filename, ruc, period):
    legacy_filename(filename, ruc, period)
    checked_zip(raw)
    workbook = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    documents, lines, metadata = {}, [], {"warnings": [], "currency": "PEN", "manual_totals": []}
    try:
        if "Detalle Ventas" not in workbook.sheetnames:
            raise ValidationError("No se encontró la hoja Detalle Ventas.")
        sheet = workbook["Detalle Ventas"]
        if sheet.max_row and sheet.max_row > MAX_ROWS + 1 or sheet.max_column and sheet.max_column > 64:
            raise ValidationError("El detalle supera 15,000 líneas o 64 columnas.")
        # Ignore stale dimension metadata; enforce the real stream limit too.
        sheet.reset_dimensions()
        rows = sheet.iter_rows(max_col=18, values_only=True)
        if [normalize(v) for v in next(rows, ())] != [normalize(v) for v in LEGACY_HEADERS]:
            raise ValidationError("Las columnas de Detalle Ventas no corresponden al formato OSIPTEL conocido.")
        seen_lines = Counter()
        for row_number, values in enumerate(rows, 2):
            if row_number > MAX_ROWS + 1:
                raise ValidationError("El archivo supera 15,000 líneas.")
            if not any(v is not None for v in values):
                continue
            try:
                key = document_key(values[0], values[3], values[4])
                emitted = parsed_date(values[2])
                if emitted.replace(day=1) != period:
                    raise ValidationError("La fecha de emisión está fuera del mes seleccionado.")
                base, tax, total = (money(v) for v in values[15:18])
                local_state, sunat_state = text_value(values[7], 60), text_value(values[8], 60)
                concept, tech = text_value(values[11], 160), text_value(values[12], 100)
                flags = []
                if base == ZERO and tax == total and total != ZERO:
                    flags.append("TAX_MAPPING")
                if not concept or not tech or "desconocido" in normalize(concept + tech):
                    flags.append("UNKNOWN_CLASS")
                if local_state == "1" or normalize(local_state) == "pendiente":
                    flags.append("LOCAL_PENDING")
                if not sunat_state:
                    flags.append("NO_SUNAT_STATE")
                identity = (emitted, text_value(values[5], 20), text_value(values[6], 250), local_state, sunat_state)
                doc = documents.setdefault(key, dict(
                    document_type=key[0], series=key[1], number=key[2], issue_date=emitted,
                    receiver_document=identity[1], receiver_name=identity[2], currency="PEN",
                    local_state=local_state, sunat_state=sunat_state, base=ZERO, tax=ZERO, total=ZERO,
                    flags=set(), identity=identity))
                if identity != doc["identity"]:
                    flags.append("MIXED_METADATA")
                signature = tuple(str(v) for v in values)
                if seen_lines[signature]:
                    flags.append("REPEATED_LINE")
                seen_lines[signature] += 1
                doc["base"] += base
                doc["tax"] += tax
                doc["total"] += total
                doc["flags"].update(flags)
                lines.append(dict(key=key, source_row=row_number, description=text_value(values[10]),
                                  concept=concept, technology=tech, quantity=money(values[14], places=5),
                                  base=base, tax=tax, total=total, flags=flags))
            except ValidationError as error:
                raise ValidationError(f"Detalle Ventas, fila {row_number}: {' '.join(error.messages)}") from None
        if "Notas Credito" in workbook.sheetnames:
            note_sheet = workbook["Notas Credito"]
            if note_sheet.max_row and note_sheet.max_row > MAX_ROWS + 1:
                raise ValidationError("La hoja Notas Credito supera el límite de lectura.")
            if note_sheet.max_row and note_sheet.max_row > 1:
                # Do not silently omit adjustments whose layout is not known.
                for values in note_sheet.iter_rows(min_row=2, max_row=MAX_ROWS + 2, max_col=30, values_only=True):
                    if any(v is not None for v in values):
                        raise ValidationError("Notas Credito contiene datos. Su estructura debe validarse antes de importar este archivo completo.")
        if "Resumen OSIPTEL" in workbook.sheetnames:
            cached = load_workbook(BytesIO(raw), read_only=True, data_only=True, keep_links=False)
            try:
                summary_rows = list(cached["Resumen OSIPTEL"].iter_rows(max_row=200, max_col=12))
                grouping_columns = [i for i in (7, 8) if any(normalize(row[i + 1].value) == "internet" for row in summary_rows)]
                for cells in summary_rows:
                    if normalize(cells[0].value) != "totalgeneral":
                        continue
                    source_total = money(cells[4].value)
                    computed = sum((v["total"] for v in documents.values()), ZERO)
                    metadata["summary_total"] = str(source_total)
                    if abs(source_total - computed) > CENT:
                        metadata["warnings"].append("El total del resumen no coincide con la suma del detalle.")
                    # Known legacy workbooks place manual grouping totals in H or I.
                    for i in grouping_columns:
                        cell = cells[i]
                        if cell.value is None:
                            continue
                        value = money(cell.value)
                        metadata["manual_totals"].append({"cell": cell.coordinate, "total": str(value),
                                                          "difference": str(value - computed)})
                        if abs(value - computed) > CENT:
                            metadata["warnings"].append(f"Total auxiliar {cell.coordinate}: diferencia de S/ {value - computed:.2f} respecto del detalle.")
                    break
            finally:
                cached.close()
        for doc in documents.values():
            doc.pop("identity")
            doc["flags"] = sorted(doc["flags"])
            for field in ("base", "tax", "total"):
                money(doc[field])
        return list(documents.values()), lines, metadata
    finally:
        workbook.close()


def parse_rvie(raw, filename, ruc, period):
    """Dispatch only recognized SICV and observed native RVIE schemas."""
    try:
        decoded = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise ValidationError("Guarde el archivo de comparación como CSV UTF-8.") from None
    if not decoded.strip():
        raise ValidationError("El CSV no contiene datos.")
    delimiter = ";" if ";" in decoded.splitlines()[0] else ","
    reader = csv.DictReader(StringIO(decoded), delimiter=delimiter)
    if reader.fieldnames and normalize(reader.fieldnames[0]) == "ruc":
        from .native_rvie import parse_native
        return parse_native(decoded, delimiter, ruc, period)
    if reader.fieldnames != RVIE_HEADERS:
        raise ValidationError("Use el CSV RVIE nativo de 40 columnas o la plantilla CSV de comparación de este panel.")
    docs, seen = [], set()
    for row_number, row in enumerate(reader, 2):
        if row_number > MAX_ROWS + 1:
            raise ValidationError("El CSV supera 15,000 comprobantes.")
        try:
            if None in row or any(v is None for v in row.values()):
                raise ValidationError("Cantidad de columnas incorrecta.")
            if row["ruc_emisor"] != ruc or row["periodo"] != period.strftime("%Y%m"):
                raise ValidationError("RUC o periodo distinto al seleccionado.")
            key = document_key(row["tipo"], row["serie"], row["numero"])
            if key in seen:
                raise ValidationError("Comprobante repetido. El archivo RVIE requiere una fila por comprobante.")
            seen.add(key)
            if row["moneda"] not in ("PEN", "USD"):
                raise ValidationError("Moneda no admitida: use PEN o USD.")
            docs.append(dict(document_type=key[0], series=key[1], number=key[2], issue_date=parsed_date(row["fecha_emision"]),
                             receiver_document=text_value(row["documento_cliente"], 20), receiver_name=text_value(row["nombre_cliente"], 250),
                             currency=row["moneda"], base=money(row["base_imponible"], optional=True),
                             tax=money(row["igv"], optional=True), total=money(row["total"]),
                             local_state="", sunat_state=text_value(row["estado_sunat"], 60), flags=[]))
        except ValidationError as error:
            raise ValidationError(f"CSV, fila {row_number}: {' '.join(error.messages)}") from None
    return docs, [], {"warnings": [], "format": "SICV-RVIE-1", "origin": "Archivo aportado por el usuario; sin consulta directa a SUNAT"}


def import_report(*, user, issuer, period, source, upload):
    require_access(user, issuer, "accounting.import_reports")
    if source not in ImportBatch.Source.values or period.day != 1:
        raise ValidationError("Fuente o periodo inválido.")
    if not re.fullmatch(r"[0-9]{11}", issuer.ruc):
        raise ValidationError("Configure un RUC de 11 dígitos para la empresa.")
    raw = upload.read(MAX_BYTES + 1)
    if not raw or len(raw) > MAX_BYTES:
        raise ValidationError("Seleccione un archivo de hasta 5 MB.")
    filename = Path(upload.name).name
    if len(filename) > 200:
        raise ValidationError("El nombre del archivo supera 200 caracteres.")
    extensions = (".xlsx",) if source == ImportBatch.Source.OSIPTEL else (".csv", ".zip")
    if not filename.lower().endswith(extensions):
        raise ValidationError(f"Esta fuente requiere un archivo {' o '.join(extensions)}.")
    digest = hashlib.sha256(raw).hexdigest()
    filters = dict(issuer=issuer, period=period, source=source, sha256=digest)
    existing = ImportBatch.objects.filter(**filters).first()
    if existing:
        if existing.issuer_ruc != issuer.ruc:
            raise ValidationError("El RUC configurado cambió desde esa carga. Revise la empresa y el archivo original.")
        return existing, False
    try:
        parser = parse_osiptel if source == ImportBatch.Source.OSIPTEL else parse_rvie
        content, inner_name = raw, filename
        if source == ImportBatch.Source.RVIE and filename.lower().endswith(".zip"):
            from .native_rvie import read_zip
            content, inner_name = read_zip(raw)
        documents, lines, metadata = parser(content, inner_name, issuer.ruc, period)
        metadata["content_sha256"] = hashlib.sha256(content).hexdigest()
        metadata["content_filename"] = inner_name
    except (BadZipFile, InvalidFileException, KeyError, ValueError, csv.Error, IndexError, ParseError,
            NotImplementedError, RuntimeError, EOFError, zlib.error):
        raise ValidationError("No se pudo leer el archivo. Verifique el formato y vuelva a exportarlo.") from None
    if not documents:
        raise ValidationError("El archivo no contiene comprobantes. No sustituirá la última carga.")
    dates = Counter(doc["issue_date"].isoformat() for doc in documents)
    metadata["issue_dates"] = {"first": min(dates), "last": max(dates), "counts": dict(sorted(dates.items()))}
    # Serialize two concurrent imports of the same issuer; a retry is harmless.
    with transaction.atomic():
        locked_issuer = Issuer.objects.select_for_update().get(pk=issuer.pk)
        if locked_issuer.ruc != issuer.ruc:
            raise ValidationError("El RUC cambió durante la carga. Seleccione de nuevo la empresa.")
        require_access(user, issuer, "accounting.import_reports")
        existing = ImportBatch.objects.filter(issuer=issuer, issuer_ruc=issuer.ruc, period=period, source=source,
            metadata__content_sha256=metadata["content_sha256"]).first() or ImportBatch.objects.filter(**filters).first()
        if existing:
            return existing, False
        batch = ImportBatch.objects.create(**filters, issuer_ruc=issuer.ruc, filename=filename,
            imported_by=user, document_count=len(documents), line_count=len(lines), metadata=metadata)
        Document.objects.bulk_create([Document(batch=batch, **doc) for doc in documents], batch_size=500)
        ids = {doc.key: doc.pk for doc in batch.documents.all()}
        Line.objects.bulk_create([Line(document_id=ids[line["key"]], **{k: v for k, v in line.items() if k != "key"})
                                  for line in lines], batch_size=500)
        AuditEvent.objects.create(actor=user, method="POST", route_name="accounting:import", path="/contabilidad/importar/",
            description="Importó evidencia para conciliación contable", changes={"batch": batch.pk, "source": source, "sha256": digest})
    return batch, True
