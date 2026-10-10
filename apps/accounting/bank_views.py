import csv
from decimal import Decimal
import hashlib
from io import BytesIO

from django import forms
from django.contrib import messages
from django.core.exceptions import ObjectDoesNotExist, ValidationError
from django.core.paginator import Paginator
from django.db.models import Exists, OuterRef, Sum
from django.http import FileResponse, Http404, HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.views.decorators.http import require_http_methods
from openpyxl import Workbook

from apps.payments.cash_forms import StyledForm
from .access import available_issuers
from .banking import (HEADERS, MAX_BYTES, accounts_for, annotate_line, confirm_match, create_account,
                      deposit_candidates, import_statement, payment_candidates, release_match)
from .models import BankAccount, BankEvent, BankMatch
from .views import protected


class AccountForm(StyledForm):
    issuer_id = forms.ModelChoiceField(label="Empresa titular", queryset=None)
    bank = forms.CharField(label="Banco", max_length=80, help_text="Use el mismo nombre del banco registrado en los depósitos de caja.")
    number = forms.CharField(label="Número completo de cuenta", max_length=60)
    label = forms.CharField(label="Nombre de la cuenta", max_length=120)

    def __init__(self, *args, user, **kwargs):
        super().__init__(*args, **kwargs)
        self.fields["issuer_id"].queryset = available_issuers(user)


class StatementForm(StyledForm):
    file = forms.FileField(label="Extracto normalizado CSV UTF-8 (máximo 2 MB)")

    def clean_file(self):
        upload = self.cleaned_data["file"]
        if upload.size > MAX_BYTES or not upload.name.lower().endswith(".csv"):
            raise forms.ValidationError("Seleccione un CSV de hasta 2 MB.")
        return upload


class MatchForm(StyledForm):
    kind = forms.ChoiceField(label="Origen en SICV", choices=[("PAYMENT", "Cobro de abonado"), ("DEPOSIT", "Depósito de caja al banco")])
    target_id = forms.IntegerField(label="ID del cobro o del movimiento de caja", min_value=1)
    note = forms.CharField(label="Comprobación realizada y motivo del vínculo", max_length=500, widget=forms.Textarea(attrs={"rows": 2}))


class ReleaseForm(StyledForm):
    match_id = forms.IntegerField(widget=forms.HiddenInput)
    reason = forms.CharField(label="Motivo para deshacer la conciliación", max_length=500)


class NoteForm(StyledForm):
    note = forms.CharField(label="Observación (comisión, transferencia, diferencia o pendiente de identificar)", max_length=500, widget=forms.Textarea(attrs={"rows": 2}))


def get_account(request, pk):
    return get_object_or_404(accounts_for(request.user), pk=pk)


@protected("accounting.view_bank")
@require_http_methods(["GET", "POST"])
def bank_workspace(request):
    form = AccountForm(request.POST or None, user=request.user)
    if request.method == "POST" and form.is_valid():
        try:
            values = dict(form.cleaned_data)
            values["issuer_id"] = values["issuer_id"].pk
            account = create_account(actor=request.user, **values)
        except ValidationError as exc:
            form.add_error(None, exc)
        else:
            messages.success(request, "Cuenta creada. Importe su extracto para empezar a conciliar.")
            return redirect("accounting:bank_account", pk=account.pk)
    return render(request, "accounting/bank/workspace.html", {"accounts": accounts_for(request.user), "form": form})


def filtered_lines(account, params):
    lines = account.lines.annotate(matched=Exists(BankMatch.objects.filter(line_id=OuterRef("pk"), active=True)))
    state = params.get("state", "pending")
    if state in ("pending", "matched"):
        lines = lines.filter(matched=state == "matched")
    query = params.get("q", "").strip()[:120]
    if query:
        from django.db.models import Q
        lines = lines.filter(Q(operation__icontains=query) | Q(external_id__icontains=query) | Q(description__icontains=query))
    return lines


@protected("accounting.view_bank")
@require_http_methods(["GET", "POST"])
def bank_account(request, pk):
    account = get_account(request, pk)
    form = StatementForm(request.POST or None, request.FILES or None)
    if request.method == "POST" and form.is_valid():
        upload = form.cleaned_data["file"]
        try:
            batch = import_statement(actor=request.user, account_id=pk, filename=upload.name, raw=upload.read(MAX_BYTES + 1))
        except ValidationError as exc:
            form.add_error(None, exc)
        else:
            messages.success(request, f"Extracto #{batch.pk}: {batch.row_count} filas, {batch.new_count} movimientos nuevos en esa importación. Las filas existentes no se duplican.")
            return redirect("accounting:bank_account", pk=pk)
    lines = filtered_lines(account, request.GET)
    pending = filtered_lines(account, {"state": "pending"})
    total = account.lines.aggregate(net=Sum("amount"))["net"] or Decimal("0.00")
    return render(request, "accounting/bank/account.html", {"account": account, "form": form,
        "page": Paginator(lines, 40).get_page(request.GET.get("page")), "net": total, "pending_count": pending.count(),
        "line_count": account.lines.count(), "state": request.GET.get("state", "pending"), "q": request.GET.get("q", ""),
        "statements": account.statements.select_related("imported_by")[:30], "headers": ";".join(HEADERS)})


@protected("accounting.view_bank")
@require_http_methods(["GET", "POST"])
def bank_line(request, account_pk, pk):
    account = get_account(request, account_pk)
    line = get_object_or_404(account.lines, pk=pk)
    form, release, note = MatchForm(), ReleaseForm(), NoteForm()
    if request.method == "POST":
        action = request.POST.get("action")
        chosen = {"match": MatchForm, "release": ReleaseForm, "note": NoteForm}.get(action)
        if chosen is None:
            raise Http404
        bound = chosen(request.POST)
        if bound.is_valid():
            try:
                if action == "match":
                    confirm_match(actor=request.user, line_id=pk, **bound.cleaned_data)
                elif action == "release":
                    get_object_or_404(line.matches, pk=bound.cleaned_data["match_id"])
                    release_match(actor=request.user, **bound.cleaned_data)
                else:
                    annotate_line(actor=request.user, line_id=pk, **bound.cleaned_data)
            except ValidationError as exc:
                bound.add_error(None, exc)
            except ObjectDoesNotExist:
                bound.add_error(None, "El registro indicado no está disponible. Revise el ID.")
            else:
                messages.success(request, "Revisión bancaria guardada.")
                return redirect("accounting:bank_line", account_pk=account_pk, pk=pk)
        if action == "match":
            form = bound
        elif action == "release":
            release = bound
        else:
            note = bound
    current = line.matches.filter(active=True).first()
    return render(request, "accounting/bank/line.html", {"account": account, "line": line, "current": current,
        "payments": payment_candidates(line) if line.amount > 0 and current is None else [],
        "deposits": deposit_candidates(line) if line.amount > 0 and current is None else [],
        "form": form, "release_form": release if release.is_bound else ReleaseForm(initial={"match_id": current.pk if current else None}),
        "note_form": note, "matches": line.matches.select_related("confirmed_by", "released_by"),
        "events": line.events.select_related("actor"), "statements": line.statements.all()})


@protected("accounting.view_bank")
@require_http_methods(["GET"])
def bank_template(request, pk):
    get_account(request, pk)
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = 'attachment; filename="plantilla-extracto-bancario.csv"'
    response.write("\ufeff")
    csv.writer(response, delimiter=";").writerow(HEADERS)
    return response


@protected("accounting.view_bank")
@require_http_methods(["GET"])
def bank_source(request, account_pk, pk):
    account = get_account(request, account_pk)
    batch = get_object_or_404(account.statements, pk=pk)
    raw = bytes(batch.original)
    if hashlib.sha256(raw).hexdigest() != batch.sha256:
        raise Http404("La evidencia no coincide con su huella guardada.")
    return FileResponse(BytesIO(raw), as_attachment=True, filename=f"extracto-{batch.pk}.csv", content_type="text/csv; charset=utf-8")


@protected("accounting.view_bank")
@require_http_methods(["GET"])
def bank_export(request, pk):
    account = get_account(request, pk)
    book = Workbook()
    sheet = book.active
    sheet.title = "Conciliación"
    def row(target, values):
        target.append(values)
        for cell in target[target.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
            elif isinstance(cell.value, Decimal):
                cell.number_format = '#,##0.00;[Red]-#,##0.00'
    row(sheet, ["Empresa", str(account.issuer), "RUC", account.issuer.ruc])
    row(sheet, ["Banco", account.bank, "Cuenta", account.number, "Moneda", account.currency])
    row(sheet, ["ID", "ID banco", "Fecha", "Operación", "Descripción", "Importe PEN", "Estado", "Vínculo", "Origen", "Confirmó", "Comentario"])
    for line in account.lines.prefetch_related("matches__confirmed_by").order_by("date", "pk"):
        active = next((m for m in line.matches.all() if m.active), None)
        row(sheet, [line.pk, line.external_id, str(line.date), line.operation, line.description, line.amount,
                   "Conciliado" if active else "Pendiente", active.pk if active else "",
                   f"{active.snapshot['target']['kind']} #{active.snapshot['target']['id']}" if active else "",
                   str(active.confirmed_by) if active else "", active.note if active else ""])
    history = book.create_sheet("Historial")
    row(history, ["Fecha", "Movimiento", "Acción", "Persona", "Nota"])
    for event in BankEvent.objects.filter(account=account).select_related("actor").order_by("pk"):
        row(history, [event.created_at.isoformat(), event.line_id, event.action, str(event.actor), event.note])
    sources = book.create_sheet("Fuentes")
    row(sources, ["ID", "Archivo", "SHA-256", "Filas", "Nuevos", "Importado"])
    for batch in account.statements.all():
        row(sources, [batch.pk, batch.filename, batch.sha256, batch.row_count, batch.new_count, batch.imported_at.isoformat()])
    for page in book:
        page.freeze_panes = "A4" if page is sheet else "A2"
        for col in page.columns:
            page.column_dimensions[col[0].column_letter].width = min(55, max(16, max(len(str(c.value or "")) for c in col) + 2))
    output = BytesIO()
    book.save(output)
    output.seek(0)
    return FileResponse(output, as_attachment=True, filename=f"conciliacion-bancaria-{account.pk}.xlsx")
