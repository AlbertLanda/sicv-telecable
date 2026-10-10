from decimal import Decimal
from io import BytesIO
from urllib.parse import urlencode
import uuid

from django import forms
from django.contrib import messages
from django.core.exceptions import ValidationError
from django.core.paginator import Paginator
from django.http import FileResponse, Http404
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_http_methods
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from apps.organization.models import Branch, Office
from apps.payments.cash_forms import StyledForm
from .models import RevenueRule
from .revenue import LIMA, RevenueReport, revenue_issuers, save_rule, scoped_payments
from .views import protected


class RevenueFilter(StyledForm):
    issuer = forms.ModelChoiceField(label="Empresa / RUC", queryset=None, required=False, empty_label="Todas mis empresas autorizadas")
    period = forms.ChoiceField(label="Periodo", choices=RevenueRule.Period.choices)
    anchor = forms.DateField(label="Fecha dentro del periodo", input_formats=["%Y-%m-%d"], widget=forms.DateInput(attrs={"type": "date"}))
    branch = forms.ModelChoiceField(label="Sede que recibió el cobro", queryset=None, required=False, empty_label="Todas las sedes")
    office = forms.ModelChoiceField(label="Oficina que recibió el cobro", queryset=None, required=False, empty_label="Todas las oficinas")

    def __init__(self, *args, user, **kwargs):
        super().__init__(*args, **kwargs)
        self.fields["issuer"].queryset = revenue_issuers(user)
        payments = scoped_payments(user)
        self.fields["branch"].queryset = Branch.objects.filter(pk__in=payments.values("branch_id")).order_by("name")
        self.fields["office"].queryset = Office.objects.filter(pk__in=payments.values("office_id")).select_related("branch").order_by("branch__name", "name")

    def clean_anchor(self):
        value = self.cleaned_data["anchor"]
        if not 2000 <= value.year <= 2100:
            raise forms.ValidationError("Seleccione un año entre 2000 y 2100.")
        return value

    def clean(self):
        data = super().clean()
        branch, office = data.get("branch"), data.get("office")
        if branch and office and office.branch_id != branch.pk:
            self.add_error("office", "La oficina no pertenece a la sede seleccionada.")
        return data


class RuleForm(StyledForm):
    period = forms.ChoiceField(choices=RevenueRule.Period.choices, widget=forms.HiddenInput)
    expected_pk = forms.IntegerField(min_value=0, widget=forms.HiddenInput)
    request_key = forms.UUIDField(initial=uuid.uuid4, widget=forms.HiddenInput)
    enabled = forms.BooleanField(label="Alerta activa", required=False, widget=forms.CheckboxInput(attrs={"class": "form-check-input"}))
    limit = forms.DecimalField(label="Umbral administrativo (S/)", max_digits=16, decimal_places=2, min_value=Decimal(".01"))
    warning_percent = forms.IntegerField(label="Avisar desde este porcentaje (%)", min_value=1, max_value=100)
    reason = forms.CharField(label="Motivo de la configuración o cambio", max_length=500, widget=forms.Textarea(attrs={"rows": 2}))


@protected("accounting.view_revenue")
@require_http_methods(["GET"])
def workspace(request):
    defaults = {"period": "MONTH", "anchor": timezone.localdate(timezone=LIMA).isoformat()}
    form = RevenueFilter(request.GET if request.GET else defaults, user=request.user)
    context = {"form": form, "has_issuers": revenue_issuers(request.user).exists()}
    if not form.is_valid():
        return render(request, "accounting/revenue/workspace.html", context, status=400)
    report = RevenueReport(user=request.user, **form.cleaned_data)
    if "export" in request.GET:
        if request.GET["export"] != "xlsx":
            raise Http404
        try:
            return export_report(report)
        except ValidationError as exc:
            form.add_error(None, exc)
            return render(request, "accounting/revenue/workspace.html", context, status=400)
    summary = report.summary()
    params = {"period": form.cleaned_data["period"], "anchor": form.cleaned_data["anchor"].isoformat()}
    for key in ("issuer", "branch", "office"):
        if form.cleaned_data[key]:
            params[key] = form.cleaned_data[key].pk
    peak = max((item["total"] for item in summary["daily"]), default=1)
    for item in summary["daily"]:
        item["width"] = int(item["total"] * 100 / peak) if peak else 0
    context.update({"report": report, **summary,
        "page": Paginator(report.detail(), 40).get_page(request.GET.get("page")),
        "query": urlencode(params), "export_url": "?" + urlencode({**params, "export": "xlsx"})})
    return render(request, "accounting/revenue/workspace.html", context)


@protected("accounting.view_revenue")
@require_http_methods(["GET", "POST"])
def rules(request, pk):
    issuer = get_object_or_404(revenue_issuers(request.user), pk=pk)
    period = request.POST.get("period") if request.method == "POST" else request.GET.get("period", "MONTH")
    if period not in RevenueRule.Period.values:
        raise Http404
    latest = RevenueRule.objects.filter(issuer=issuer, period=period).first()
    initial = {"period": period, "expected_pk": latest.pk if latest else 0,
        "enabled": latest.enabled if latest else True, "limit": latest.limit if latest else None,
        "warning_percent": latest.warning_percent if latest else 80}
    form = RuleForm(request.POST or None, initial=initial)
    if request.method == "POST":
        from .access import require_access
        require_access(request.user, issuer, "accounting.manage_revenue_rules")
        if form.is_valid():
            try:
                save_rule(actor=request.user, issuer_id=pk, **form.cleaned_data)
            except ValidationError as exc:
                form.add_error(None, exc)
            else:
                messages.success(request, "Nueva versión guardada. La alerta solo informa; no modifica cobros ni talonarios.")
                return redirect(reverse("accounting:revenue_rules", args=[pk]) + "?" + urlencode({"period": period}))
    return render(request, "accounting/revenue/rules.html", {"issuer": issuer, "form": form, "latest": latest,
        "period": period, "periods": RevenueRule.Period.choices, "period_label": RevenueRule.Period(period).label,
        "history": Paginator(RevenueRule.objects.filter(issuer=issuer, period=period).select_related("actor"), 20).get_page(request.GET.get("page"))},
        status=400 if request.method == "POST" else 200)


def export_report(report):
    # Una sola consulta materializa el detalle. Todos los totales del Excel
    # se calculan con esas mismas filas, aunque luego entre otro cobro.
    payments = list(report.detail()[:50001])
    if len(payments) > 50000:
        raise ValidationError("La exportación admite hasta 50000 cobros. Reduzca el periodo, empresa o sede.")
    issuers = {item.pk: item for item in report.issuers}
    book = Workbook()
    control = book.active
    control.title = "Control"
    def row(sheet, values):
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
            elif isinstance(cell.value, Decimal):
                cell.number_format = '#,##0.00;[Red]-#,##0.00'
    total = sum((p.amount for p in payments), Decimal("0.00"))
    for values in [
        ["Recaudación operativa", "PEN"], ["Desde", str(report.start), "Hasta", str(report.end)],
        ["Consulta", timezone.localtime(report.now, LIMA).isoformat()],
        ["Empresa", str(report.issuer) if report.issuer else "Todas las empresas autorizadas"],
        ["Sede", str(report.branch) if report.branch else "Todas", "Oficina", str(report.office) if report.office else "Todas"],
        ["Cobros", len(payments), "Total PEN", total],
        ["Criterio", "Cobros confirmados por fecha real de pago en Lima; sin fechas futuras. No es emisión fiscal ni cierre inmutable."],
        ["Atribución", "Caja conserva la empresa original. Sin caja se usa el talonario actual, cuya configuración puede cambiar. Empresas se identifican por su código; un RUC repetido no se fusiona."],
        ["Sin fecha de pago", report.missing_date, "Todas las fechas de las empresas seleccionadas"],
        ["Fecha futura en el periodo", report.future, "Todas sus sedes"],
        ["RUC de caja distinto del configurado", report.ruc_conflicts, "Excluidos del periodo; todas sus sedes"],
        ["Alcance", "No suma archivos SIRE/OSIPTEL, depósitos de caja, garantías, gastos ni ajustes comerciales de deuda."],
    ]:
        row(control, values)
    detail = book.create_sheet("Cobros")
    row(detail, ["ID pago", "Empresa", "RUC", "Fecha real de pago (Lima)", "Sede receptora", "Oficina", "Medio", "Importe PEN", "Referencia", "Abonado", "Usuario", "Comprobante actual", "Atribución", "Referencia conservada en caja"])
    company_totals, offices, days = {}, {}, {}
    for p in payments:
        issuer = issuers[p.revenue_issuer]
        local = timezone.localtime(p.paid_at, LIMA)
        receipt = p.receipt.full_number if hasattr(p, "receipt") else "Sin comprobante"
        row(detail, [p.pk, issuer.business_name, issuer.ruc, local.isoformat(), p.branch.name,
            p.office.name if p.office else "Sin oficina", p.get_method_display(), p.amount, p.reference, p.customer.code,
            p.received_by.get_username(), receipt, p.revenue_source, p.cash_reference or ""])
        for container, key in [(company_totals, (issuer.code, issuer.business_name, issuer.ruc)),
                (offices, (p.branch.code, p.branch.name, p.office.code if p.office else "", p.office.name if p.office else "Sin oficina", p.get_method_display())),
                (days, (str(local.date()),))]:
            current = container.setdefault(key, [0, Decimal("0.00")])
            current[0] += 1
            current[1] += p.amount
    for title, headers, container in [("Por empresa", ["Código empresa", "Empresa", "RUC"], company_totals),
            ("Por oficina y medio", ["Código sede", "Sede", "Código oficina", "Oficina", "Medio"], offices), ("Por día", ["Fecha"], days)]:
        sheet = book.create_sheet(title)
        row(sheet, headers + ["Cobros", "Total PEN"])
        for key, values in sorted(container.items()):
            row(sheet, [*key, *values])
    for sheet in book:
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="173B57")
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(55, max(16, max(len(str(c.value or "")) for c in column[:150]) + 2))
        if sheet is not control:
            sheet.auto_filter.ref = sheet.dimensions
    output = BytesIO()
    book.save(output)
    output.seek(0)
    return FileResponse(output, as_attachment=True, filename=f"recaudacion-{report.start}-{report.end}.xlsx")
