from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Review
from .native_rvie import NATIVE_HEADERS
from .reconciliation import classified_lines, osiptel_summary


def add_row(sheet, values):
    sheet.append(values)
    for cell, value in zip(sheet[sheet.max_row], values):
        if isinstance(value, str):
            # Spreadsheet text, including imported descriptions, cannot execute.
            cell.data_type = "s"


def export_workbook(*, issuer, period, legacy, rvie, rows, cutoff):
    book = Workbook()
    control = book.active
    control.title = "Control"
    add_row(control, ["Campo", "Valor"])
    for row in [("Empresa", issuer.business_name), ("RUC", issuer.ruc), ("Periodo", period.strftime("%Y-%m")),
                ("Alcance", "Comparación de archivos importados; no acredita aceptación SUNAT ni genera registros SIRE"),
                ("Importes", "Se conservan signo e importes originales. Sin recalcular IGV ni convertir monedas."),
                ("Notas de crédito", "No se deduce su ausencia del archivo. Los importes se suman con el signo recibido."),
                ("Vínculo ERP", "Consulta al momento de exportar, por empresa, tipo, serie y número del talonario actual. No acredita emisión fiscal."),
                ("Revisiones hasta ID", cutoff), ("Exportado (Lima)", timezone.localtime().strftime("%Y-%m-%d %H:%M:%S"))]:
        add_row(control, list(row))
    for batch in (legacy, rvie):
        if batch:
            for key, value in [("Versión", batch.pk), ("Fuente", batch.get_source_display()), ("Archivo", batch.filename),
                               ("RUC de origen", batch.issuer_ruc), ("SHA256", batch.sha256),
                               ("Importado por", batch.imported_by.username),
                               ("Importado (Lima)", timezone.localtime(batch.created_at).strftime("%Y-%m-%d %H:%M:%S"))]:
                add_row(control, [key, value])
            for warning in batch.metadata.get("warnings", []):
                add_row(control, ["Observación de archivo", warning])
            dates = batch.metadata.get("issue_dates", {})
            if dates:
                add_row(control, ["Fechas de emisión presentes", f"{dates['first']} a {dates['last']}"])
    sheet = book.create_sheet("Conciliación")
    add_row(sheet, ["Tipo", "Serie", "Número", "Resultado", "Diferencias", "Moneda SICV", "Total SICV", "Base SICV", "IGV SICV",
                    "Estado local SICV", "Código SUNAT en SICV", "Moneda SIRE", "Total SIRE", "Base SIRE", "IGV SIRE",
                    "Código SUNAT en plantilla", "Diferencia SICV-SIRE", "Alertas de origen",
                    "Est. Comp SIRE", "Receptor de origen", "Documento receptor", "Vínculo ERP",
                    "Código abonado", "Abonado ERP", "Documento abonado", "Sede abonado", "Sede servicio",
                    "Servicio / concepto", "Sede cobro", "Estado cobro", "Estado fiscal", "Alertas vínculo"])
    for row in rows:
        doc, local, external = row["document"], row["local"], row["rvie"]
        erp = row.get("erp", {})
        add_row(sheet, [doc.document_type, doc.series, doc.number, row["label"], "; ".join(row["differences"]),
                        local.currency if local else "", local.total if local else None, local.base if local else None,
                        local.tax if local else None, local.local_state if local else "", local.sunat_state if local else "",
                        external.currency if external else "", external.total if external else None, external.base if external else None,
                        external.tax if external else None, external.sunat_state if external else "", row["delta"], "; ".join(row["flags"]),
                        external.sire_state if external else "", doc.receiver_name, doc.receiver_document, erp.get("label", ""),
                        *[erp.get(field, "") for field in ("customer_code", "customer_name", "customer_document", "customer_branch",
                            "service_branches", "services", "collection_branch", "payment_state", "fiscal_state")],
                        "; ".join(erp.get("warnings", []))])
    native = book.create_sheet("Detalle SIRE")
    add_row(native, ["Fila de origen", *NATIVE_HEADERS])
    if rvie:
        for doc in rvie.documents.all():
            if doc.source_details.get("fields"):
                add_row(native, [doc.source_details["row"], *[doc.source_details["fields"].get(h, "") for h in NATIVE_HEADERS]])
    lines = classified_lines(legacy, cutoff=cutoff)
    detail = book.create_sheet("Detalle OSIPTEL")
    add_row(detail, ["Tipo", "Serie", "Número", "Fila origen", "Descripción", "Concepto original", "Tecnología original",
                     "Concepto revisado", "Tecnología revisada", "Cantidad", "Base original", "IGV original", "Total original",
                     "Moneda", "Revisión ID", "Motivo"])
    for row in lines:
        doc, line, review = row["document"], row["line"], row["review"]
        add_row(detail, [doc.document_type, doc.series, doc.number, line.source_row, line.description, line.concept, line.technology,
                         row["concept"], row["technology"], line.quantity, line.base, line.tax, line.total, doc.currency,
                         review.pk if review else None, review.note if review else ""])
    summary = book.create_sheet("Resumen OSIPTEL")
    add_row(summary, ["Moneda", "Concepto revisado", "Tecnología revisada", "Líneas", "Base original", "IGV original", "Total original"])
    for row in osiptel_summary(lines):
        add_row(summary, [row["currency"], row["concept"], row["technology"], row["count"], row["base"], row["tax"], row["total"]])
    reviews = book.create_sheet("Historial revisiones")
    add_row(reviews, ["ID", "Versión", "Comprobante", "Fila origen", "Usuario", "Fecha Lima", "Motivo", "Concepto", "Tecnología"])
    ids = [b.pk for b in (legacy, rvie) if b]
    for review in Review.objects.filter(document__batch_id__in=ids, pk__lte=cutoff).select_related("document", "line", "actor"):
        add_row(reviews, [review.pk, review.document.batch_id, review.document.label, review.line.source_row if review.line else None,
                          review.actor.username, timezone.localtime(review.created_at).strftime("%Y-%m-%d %H:%M:%S"),
                          review.note, review.concept, review.technology])
    for sheet in book:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="183B68")
        for index in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(index)].width = 24
    buffer = BytesIO()
    book.save(buffer)
    response = HttpResponse(buffer.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="conciliacion_{issuer.ruc}_{period:%Y%m}_L{legacy.pk if legacy else 0}_R{rvie.pk if rvie else 0}_V{cutoff}.xlsx"'
    return response
