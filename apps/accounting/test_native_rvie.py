"""Fictitious RVIE evidence only. No customer exports belong in this repository."""
import csv
from datetime import date
from decimal import Decimal
from io import BytesIO, StringIO
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED

from django.core.exceptions import PermissionDenied, ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from apps.accounts.models import User
from apps.payments.models import Issuer
from .importers import import_report
from .models import CompanyAccess, ImportBatch
from .native_rvie import NATIVE_HEADERS
from .reconciliation import compare
from .tests import RUC, PERIOD, FILENAME, legacy_bytes


def native_row(**changes):
    row = dict.fromkeys(NATIVE_HEADERS, "")
    row.update({"Ruc": RUC, "Razon Social": "Empresa de prueba", "Periodo": "202609",
        "CAR SUNAT": "CAR-FICTICIO", "Fecha de emisión": "05/09/2026", "Tipo CP/Doc.": "03",
        "Serie del CDP": "BQA1", "Nro CP o Doc. Nro Inicial (Rango)": "00001",
        "Tipo Doc Identidad": "1", "Nro Doc Identidad": "00000001",
        "Apellidos Nombres/ Razón Social": "Cliente ficticio, prueba",
        "BI Gravada": "100.00", "IGV / IPM": "18.00", "Total CP": "118.00", "Moneda": "PEN",
        "Est. Comp": "2"})
    row.update(changes)
    return [row[h] for h in NATIVE_HEADERS]


def native_csv(rows=None, headers=None):
    out = StringIO()
    writer = csv.writer(out)
    writer.writerow(NATIVE_HEADERS if headers is None else headers)
    writer.writerows([native_row()] if rows is None else rows)
    return out.getvalue().encode("utf-8-sig")


def zipped(raw, names=("RVIE_FICTICIO.csv",)):
    stream = BytesIO()
    with ZipFile(stream, "w", ZIP_DEFLATED) as archive:
        for name in names:
            archive.writestr(name, raw)
    return stream.getvalue()


class NativeRvieTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = User.objects.create_user(username="native_admin", role="ADMIN")
        cls.accountant = User.objects.create_user(username="native_accountant", role="ACCOUNTING")
        cls.issuer = Issuer.objects.create(code="RVIE-QA", ruc=RUC, business_name="Empresa ficticia RVIE")
        cls.access = CompanyAccess.objects.create(user=cls.accountant, issuer=cls.issuer, updated_by=cls.admin)

    def load(self, raw=None, name="RVIE_FICTICIO.zip"):
        return import_report(user=self.accountant, issuer=self.issuer, period=PERIOD, source="RVIE",
            upload=SimpleUploadedFile(name, zipped(native_csv()) if raw is None else raw))

    def test_native_zip_preserves_all_fields_and_separates_states(self):
        batch, created = self.load()
        self.assertTrue(created)
        doc = batch.documents.get()
        self.assertEqual(doc.key, ("03", "BQA1", "1"))
        self.assertEqual((doc.receiver_document, doc.receiver_name), ("00000001", "Cliente ficticio, prueba"))
        self.assertEqual((doc.base, doc.tax, doc.total), (Decimal("100"), Decimal("18"), Decimal("118")))
        self.assertEqual(doc.sunat_state, "")
        self.assertEqual(doc.sire_state, "2")
        self.assertEqual(doc.source_details["fields"], dict(zip(NATIVE_HEADERS, native_row())))
        self.assertEqual(batch.metadata["state_counts"], {"2": 1})
        self.assertEqual(batch.metadata["issue_dates"], {"first": "2026-09-05", "last": "2026-09-05", "counts": {"2026-09-05": 1}})

    def test_csv_and_repacked_zip_are_the_same_evidence(self):
        first, _ = self.load()
        again, created = self.load(native_csv(), "RVIE.csv")
        repacked, created_again = self.load(zipped(native_csv(), ("OTRO_NOMBRE.csv",)))
        self.assertEqual((first.pk, False, first.pk, False), (again.pk, created, repacked.pk, created_again))
        self.assertEqual(ImportBatch.objects.count(), 1)

    def test_sire_state_is_not_compared_as_sunat_acceptance_and_identity_is_checked(self):
        legacy, _ = import_report(user=self.accountant, issuer=self.issuer, period=PERIOD, source="OSIPTEL",
            upload=SimpleUploadedFile(FILENAME, legacy_bytes()))
        rvie, _ = self.load()
        self.assertEqual(compare(legacy, rvie)[0]["result"], "MATCH")
        self.assertEqual(rvie.documents.get().sunat_state, "")
        rvie.documents.update(receiver_document="00000002")
        self.assertIn("Documento receptor", compare(legacy, rvie)[0]["differences"])

    def test_raw_export_strings_cannot_execute_as_excel_formulas(self):
        batch, _ = self.load(zipped(native_csv([native_row(**{"Apellidos Nombres/ Razón Social": "=1+1"})])))
        self.client.force_login(self.accountant)
        response = self.client.get(reverse("accounting:workspace"),
            {"issuer": self.issuer.pk, "period": "2026-09", "rvie": batch.pk, "export": "xlsx"})
        book = load_workbook(BytesIO(response.content))
        self.assertEqual(book["Detalle SIRE"]["N2"].value, "=1+1")
        self.assertEqual(book["Detalle SIRE"]["N2"].data_type, "s")

    def test_all_or_nothing_rejects_schema_scope_ranges_duplicates_and_amounts(self):
        bad_rows = [native_row(**changes) for changes in (
            {"Ruc": "20999999992"}, {"Periodo": "202610"}, {"Nro Final (Rango)": "2"},
            {"Total CP": "NaN"}, {"IGV / IPM": "18.001"}, {"Moneda": "EUR"},
            {"Fecha de emisión": "31/09/2026"}, {"Nro CP o Doc. Nro Inicial (Rango)": "-1"})]
        for row in bad_rows + [native_row(), native_row()[:-1]]:
            with self.subTest(row=row):
                with self.assertRaises(ValidationError):
                    self.load(zipped(native_csv([native_row(), row])))
        for headers in (NATIVE_HEADERS[:-1], NATIVE_HEADERS[::-1], ["Ruc", "desconocida"]):
            with self.assertRaises(ValidationError):
                self.load(zipped(native_csv(headers=headers)))
        self.assertEqual(ImportBatch.objects.count(), 0)

    def test_zip_bounds_paths_bad_encoding_and_empty_files(self):
        for names in (("../escape.csv",), ("/absolute.csv",), ("folder/file.csv",), ("bad\\file.csv",),
                      ("one.csv", "two.csv"), ("invoice.xml",)):
            with self.subTest(names=names), self.assertRaises(ValidationError):
                self.load(zipped(native_csv(), names))
        with patch("apps.accounting.native_rvie.MAX_EXPANDED_BYTES", 10), self.assertRaises(ValidationError):
            self.load()
        with patch("apps.accounting.importers.MAX_ROWS", 1), self.assertRaises(ValidationError):
            self.load(zipped(native_csv([native_row(), native_row(**{"Nro CP o Doc. Nro Inicial (Rango)": "2"})])))
        for raw in (b"not a zip", zipped(b""), zipped(b"\xff\xfe"), zipped(native_csv([]))):
            with self.assertRaises(ValidationError):
                self.load(raw)
        self.assertEqual(ImportBatch.objects.count(), 0)

    def test_credit_note_other_components_and_prior_issue_date_are_preserved(self):
        fields = {"Tipo CP/Doc.": "07", "Fecha de emisión": "31/08/2026", "Fecha Vcto/Pago": "05/09/2026",
                  "BI Gravada": "-100", "IGV / IPM": "-18", "Total CP": "-118", "Mto Exonerado": "-10",
                  "Dscto BI": "2.00", "Dscto IGV / IPM": "0.36", "Tipo Cambio": "3.75000",
                  "Fecha Emisión Doc Modificado": "01/08/2026", "Tipo CP Modificado": "03",
                  "Serie CP Modificado": "BQA1", "Nro CP Modificado": "000009"}
        batch, _ = self.load(zipped(native_csv([native_row(**fields)])))
        doc = batch.documents.get()
        self.assertEqual(doc.total, Decimal("-118"))
        self.assertEqual(doc.issue_date, date(2026, 8, 31))
        for field, value in fields.items():
            self.assertEqual(doc.source_details["fields"][field], value)

    def test_native_upload_views_export_and_revocation(self):
        self.client.force_login(self.accountant)
        response = self.client.post(reverse("accounting:import"), {"issuer": self.issuer.pk, "period": "2026-09",
            "source": "RVIE", "confirm": "on", "file": SimpleUploadedFile("RVIE.zip", zipped(native_csv()))})
        self.assertEqual(response.status_code, 302)
        doc = ImportBatch.objects.get().documents.get()
        page = self.client.get(response.url)
        self.assertContains(page, "2026-09-05")
        self.assertContains(page, "Sin vínculo en el ERP")
        self.assertContains(page, "Est. Comp SIRE: 2")
        self.assertContains(self.client.get(reverse("accounting:document", args=[doc.pk])), "CAR-FICTICIO")
        export = self.client.get(response.url + "&export=xlsx")
        book = load_workbook(BytesIO(export.content))
        self.assertEqual(book["Detalle SIRE"]["M2"].value, "00000001")
        self.assertEqual(book["Detalle SIRE"].max_column, 41)
        self.access.enabled = False
        self.access.save()
        self.assertEqual(self.client.get(reverse("accounting:document", args=[doc.pk])).status_code, 404)
        with self.assertRaises(PermissionDenied):
            self.load()
