"""Synthetic fixtures only: scope, money, evidence and bounded imports."""
import csv
from datetime import date
from decimal import Decimal
from io import BytesIO, StringIO
from unittest.mock import patch
from zipfile import ZipFile

from django.core.exceptions import PermissionDenied, ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connections, close_old_connections
from django.test import Client, TestCase, TransactionTestCase, skipUnlessDBFeature
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from apps.accounts.models import User
from apps.audit.models import AuditEvent
from apps.payments.models import Issuer, Charge, Payment, Receipt
from .forms import ReviewForm
from .importers import LEGACY_HEADERS, RVIE_HEADERS, import_report, parse_osiptel, parse_rvie
from .models import CompanyAccess, Document, ImportBatch, Line, Review
from .reconciliation import classified_lines, compare, osiptel_summary, totals_by_currency

PERIOD = date(2026, 9, 1)
RUC = "20999999991"
FILENAME = f"OSIPTEL_{RUC}_20260901_20260930_DEMO.xlsx"


def legacy_row(**overrides):
    values = ["03", "Boleta", date(2026, 9, 5), "BQA1", 1, "00000001", "Cliente ficticio",
              0, 1, "", "Servicio ficticio", "Internet Fijo", "FTTH", "catálogo", 1, 100, 18, 118]
    indices = {"kind": 0, "date": 2, "series": 3, "number": 4, "name": 6, "local": 7, "sunat": 8,
               "description": 10, "concept": 11, "technology": 12, "quantity": 14, "base": 15, "tax": 16, "total": 17}
    for key, value in overrides.items():
        values[indices[key]] = value
    return values


def legacy_bytes(rows=None, *, notes=False, summary=None):
    book = Workbook()
    sheet = book.active
    sheet.title = "Detalle Ventas"
    sheet.append(LEGACY_HEADERS)
    for row in rows if rows is not None else [legacy_row()]:
        sheet.append(row)
    note_sheet = book.create_sheet("Notas Credito")
    note_sheet.append(["Tipo", "Serie", "Número"])
    if notes:
        note_sheet.append(["07", "BQA1", 9])
    if summary:
        sheet = book.create_sheet("Resumen OSIPTEL")
        for row in summary:
            sheet.append(row)
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def rvie_bytes(rows=None):
    stream = StringIO()
    writer = csv.writer(stream, delimiter=";")
    writer.writerow(RVIE_HEADERS)
    base = [RUC, "202609", "03", "BQA1", "00001", "05/09/2026", "PEN", "100.00", "18.00", "118.00", "1", "00000001", "Cliente ficticio"]
    writer.writerows(rows if rows is not None else [base])
    return stream.getvalue().encode("utf-8-sig")


class AccountingTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = User.objects.create_user(username="accounting_admin", role="ADMIN")
        cls.accountant = User.objects.create_user(username="accountant", role="ACCOUNTING")
        cls.other = User.objects.create_user(username="accountant_other", role="ACCOUNTING")
        cls.atc = User.objects.create_user(username="operator", role="ATC")
        cls.issuer = Issuer.objects.create(code="QA-CONT", business_name="Empresa ficticia", ruc=RUC)
        cls.second = Issuer.objects.create(code="QA-OTHER", business_name="Otra empresa ficticia", ruc="20999999992")
        cls.access = CompanyAccess.objects.create(user=cls.accountant, issuer=cls.issuer, updated_by=cls.admin)

    def load(self, raw=None, *, source="OSIPTEL", user=None, name=None):
        return import_report(user=user or self.accountant, issuer=self.issuer, period=PERIOD, source=source,
            upload=SimpleUploadedFile(name or (FILENAME if source == "OSIPTEL" else "rvie_demo.csv"),
                                      raw if raw is not None else legacy_bytes() if source == "OSIPTEL" else rvie_bytes()))[0]

    def panel(self, **params):
        return self.client.get(reverse("accounting:workspace"), {"issuer": self.issuer.pk, "period": "2026-09", **params})

    def test_authenticated_company_scope_and_landing(self):
        self.assertEqual(self.client.get(reverse("accounting:workspace")).status_code, 302)
        self.client.force_login(self.accountant)
        self.assertRedirects(self.client.get("/"), reverse("accounting:workspace"))
        self.assertContains(self.panel(), "Empresa ficticia")
        self.assertNotContains(self.panel(), "Otra empresa ficticia")
        self.assertEqual(self.panel(issuer=self.second.pk).status_code, 400)
        self.assertFalse(self.accountant.has_perm("payments.add_payment"))
        self.assertFalse(self.accountant.has_perm("payments.view_receipt"))
        self.assertFalse(self.accountant.has_perm("accounting.manage_access"))

    def test_atc_and_inactive_accountant_denied(self):
        self.client.force_login(self.atc)
        for name in ("workspace", "import", "access", "rvie_template", "demo"):
            self.assertEqual(self.client.get(reverse("accounting:" + name)).status_code, 403)
        self.accountant.is_active = False
        self.accountant.save(update_fields=["is_active"])
        with self.assertRaises(PermissionDenied):
            self.load()

    def test_grant_and_revoke_enforced_for_reads_downloads_and_writes(self):
        batch = self.load()
        doc = batch.documents.get()
        self.client.force_login(self.admin)
        url = reverse("accounting:access")
        self.assertEqual(self.client.post(url, {"user": self.accountant.pk, "issuer": self.issuer.pk}).status_code, 302)
        self.client.force_login(self.accountant)
        self.assertEqual(self.panel().status_code, 400)
        self.assertEqual(self.panel(export="xlsx").status_code, 400)
        self.assertEqual(self.client.get(reverse("accounting:document", args=[doc.pk])).status_code, 404)
        self.assertEqual(self.client.post(reverse("accounting:document", args=[doc.pk]), {"note": "Intento"}).status_code, 404)
        with self.assertRaises(PermissionDenied):
            self.load()
        self.assertTrue(AuditEvent.objects.filter(route_name="accounting:access").exists())

    def test_cannot_grant_own_access(self):
        self.client.force_login(self.accountant)
        self.assertEqual(self.client.post(reverse("accounting:access"), {"user": self.accountant.pk, "issuer": self.second.pk, "enabled": "on"}).status_code, 403)
        self.assertFalse(CompanyAccess.objects.filter(issuer=self.second, user=self.accountant).exists())

    def test_groups_items_by_type_series_and_number_without_deduplicating_lines(self):
        rows = [legacy_row(), legacy_row(), legacy_row(kind="14"), legacy_row(number=2, series="BQA2")]
        batch = self.load(legacy_bytes(rows))
        self.assertEqual((batch.document_count, batch.line_count), (3, 4))
        doc = batch.documents.get(document_type="03", series="BQA1")
        self.assertEqual(doc.total, Decimal("236.00"))
        self.assertEqual(doc.lines.count(), 2)
        self.assertIn("REPEATED_LINE", doc.flags)
        self.assertEqual(Charge.objects.count() + Payment.objects.count() + Receipt.objects.count(), 0)

    def test_import_is_idempotent_and_preserves_revisions(self):
        # Reuse the actual bytes: regenerating XLSX changes ZIP timestamps.
        raw = legacy_bytes()
        first = self.load(raw)
        again = self.load(raw)
        newer = self.load(legacy_bytes([legacy_row(total=119)]))
        self.assertEqual(first.pk, again.pk)
        self.assertNotEqual(first.pk, newer.pk)
        self.assertEqual(ImportBatch.objects.count(), 2)
        self.assertEqual(first.documents.get().total, Decimal("118"))

    def test_flags_tax_mapping_unknown_pending_blank_and_mixed_states(self):
        batch = self.load(legacy_bytes([legacy_row(base=0, tax=59, total=59, local=1, sunat=None, concept="Otros (Desconocido)"),
                                        legacy_row(base=0, tax=59, total=59, local=0, sunat=1)]))
        self.assertTrue({"TAX_MAPPING", "UNKNOWN_CLASS", "LOCAL_PENDING", "NO_SUNAT_STATE", "MIXED_METADATA"}.issubset(batch.documents.get().flags))

    def test_invalid_row_rolls_back_entire_import(self):
        for bad in (legacy_row(base="=1+1"), legacy_row(total="NaN"), legacy_row(date="31/09/2026"), legacy_row(number=-1),
                    legacy_row(base="100,00"), legacy_row(quantity="1.000001"), legacy_row(total="1.001"), legacy_row(date="2026-10-01")):
            with self.subTest(bad=bad):
                with self.assertRaises(ValidationError):
                    self.load(legacy_bytes([legacy_row(), bad]))
                self.assertEqual(ImportBatch.objects.count(), 0)

    def test_wrong_company_period_filename_extension_and_empty_rejected(self):
        for name in (FILENAME.replace(RUC, self.second.ruc), FILENAME.replace("20260901", "20260902"), "reporte.xlsx", "reporte.xls"):
            with self.assertRaises(ValidationError):
                self.load(name=name)
        with self.assertRaises(ValidationError):
            self.load(legacy_bytes([]))
        with self.assertRaises(ValidationError):
            self.load(b"")

    def test_credit_notes_not_silently_omitted(self):
        with self.assertRaisesMessage(ValidationError, "Notas Credito contiene datos"):
            self.load(legacy_bytes(notes=True))

    def test_bounded_file_rows_and_zip_xml(self):
        with patch("apps.accounting.importers.MAX_BYTES", 10):
            with self.assertRaises(ValidationError):
                self.load()
        with patch("apps.accounting.importers.MAX_ROWS", 1):
            with self.assertRaises(ValidationError):
                self.load(legacy_bytes([legacy_row(), legacy_row(number=2)]))
        data = BytesIO()
        with ZipFile(data, "w") as z:
            z.writestr("payload.xml", '<!DOCTYPE a [<!ENTITY x "test">]><a/>')
        with self.assertRaises(ValidationError):
            self.load(data.getvalue())
        with self.assertRaises(ValidationError):
            self.load(b"not a workbook")

    def test_manual_auxiliary_total_difference_is_visible(self):
        summary = [["Concepto", None, None, None, "Total"],
                   [None, None, None, None, None, None, None, 197, "INTERNET"],
                   ["TOTAL GENERAL", None, 100, 18, 118, None, None, 197, 166.95]]
        batch = self.load(legacy_bytes(summary=summary))
        self.assertEqual(batch.metadata["manual_totals"], [{"cell": "H3", "total": "197.00", "difference": "79.00"}])
        self.assertIn("79.00", batch.metadata["warnings"][0])

    def test_rvie_explicit_format_leading_zero_identity_and_missing_taxes(self):
        raw = rvie_bytes().replace(b"100.00;18.00", b";")
        batch = self.load(raw, source="RVIE")
        doc = batch.documents.get()
        self.assertEqual(doc.number, "1")
        self.assertIsNone(doc.base)
        self.assertIsNone(doc.tax)
        self.assertEqual(doc.total, Decimal("118"))

    def test_rvie_duplicate_ruc_period_format_currency_rejected(self):
        decoded = rvie_bytes().decode("utf-8-sig").splitlines()
        for bad in (("\n".join(decoded + [decoded[1]])).encode(), rvie_bytes().replace(RUC.encode(), self.second.ruc.encode()),
                    rvie_bytes().replace(b"202609", b"202608"), rvie_bytes().replace(b"PEN", b"EUR"), b"columna;desconocida\n1;2"):
            with self.assertRaises(ValidationError):
                self.load(bad, source="RVIE")
        self.assertEqual(ImportBatch.objects.count(), 0)

    def test_reconciliation_waits_for_missing_source(self):
        batch = self.load()
        rows = compare(batch, None)
        self.assertEqual(rows[0]["result"], "WAITING")
        self.assertTrue(rows[0]["attention"])

    def test_comparison_detects_missing_amount_currency_date_and_tax_differences(self):
        left = self.load(legacy_bytes([legacy_row(), legacy_row(number=2)]))
        raw = rvie_bytes().replace(b"118.00", b"117.00")
        right = self.load(raw, source="RVIE")
        rows = compare(left, right)
        self.assertEqual([row["result"] for row in rows], ["DIFFERENCE", "ONLY_LOCAL"])
        self.assertEqual(rows[0]["delta"], Decimal("1"))
        right.documents.update(number="3")
        self.assertEqual(compare(left, right)[2]["result"], "ONLY_RVIE")
        right.documents.update(number="1", total=118, currency="USD", issue_date=date(2026, 9, 6), tax=19)
        row = compare(left, right)[0]
        self.assertIsNone(row["delta"])
        self.assertIn("Moneda", row["differences"])
        self.assertIn("Fecha de emisión", row["differences"])
        self.assertIn("IGV", row["differences"])

    def test_matching_tax_errors_still_require_review_and_do_not_claim_acceptance(self):
        left = self.load(legacy_bytes([legacy_row(base=0, tax=118)]))
        right = self.load(rvie_bytes().replace(b"100.00;18.00", b"0.00;118.00"), source="RVIE")
        row = compare(left, right)[0]
        self.assertEqual(row["result"], "MATCH")
        self.assertTrue(row["attention"])
        self.assertIn("TAX_MAPPING", row["flag_codes"])

    def test_money_totals_do_not_mix_currencies_or_infer_missing_base(self):
        batch = self.load(source="RVIE")
        doc = batch.documents.get()
        clone = Document(batch=batch, document_type="14", series="SQA1", number="1", issue_date=PERIOD, total=10, currency="USD")
        totals = totals_by_currency([doc, clone])
        self.assertEqual([r["currency"] for r in totals], ["PEN", "USD"])
        self.assertTrue(totals[1]["incomplete"])

    def test_workspace_filters_versions_and_rejects_cross_company_batches(self):
        batch = self.load()
        self.client.force_login(self.accountant)
        response = self.panel()
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "BQA1-1")
        self.assertIn("no-store", response["Cache-Control"])
        self.assertEqual(self.panel(legacy="abc").status_code, 404)
        self.assertEqual(self.panel(rvie=batch.pk).status_code, 404)
        self.assertEqual(self.panel(period="2026-99").status_code, 400)
        self.assertContains(self.panel(q="no existe"), "No hay comprobantes")
        self.assertContains(self.panel(result="ATTENTION"), "BQA1-1")
        ImportBatch.objects.filter(pk=batch.pk).update(issuer=self.second)
        self.assertEqual(self.panel(legacy=batch.pk).status_code, 404)

    def test_review_requires_reason_and_line_in_same_document(self):
        batch = self.load(legacy_bytes([legacy_row(), legacy_row(number=2)]))
        doc, other = batch.documents.all()
        self.client.force_login(self.accountant)
        url = reverse("accounting:document", args=[doc.pk])
        self.assertEqual(self.client.post(url, {"note": "", "line": doc.lines.get().pk, "concept": "TV", "technology": "FTTH"}).status_code, 400)
        self.assertEqual(self.client.post(url, {"note": "Motivo", "line": other.lines.get().pk, "concept": "TV", "technology": "FTTH"}).status_code, 400)
        self.assertEqual(Review.objects.count(), 0)

    def test_append_only_review_changes_grouping_not_amounts_and_exports_can_pin_history(self):
        batch = self.load()
        doc, line = batch.documents.get(), Line.objects.get()
        self.client.force_login(self.accountant)
        url = reverse("accounting:document", args=[doc.pk])
        for concept in ("Instalación", "Reconexión"):
            self.assertEqual(self.client.post(url, {"note": "Validado contra concepto de servicio", "line": line.pk,
                                                   "concept": concept, "technology": "FTTH"}).status_code, 302)
        doc.refresh_from_db()
        line.refresh_from_db()
        self.assertEqual(line.concept, "Internet Fijo")
        self.assertEqual(doc.total, Decimal("118"))
        self.assertEqual(Review.objects.count(), 2)
        first = Review.objects.order_by("pk").first()
        self.assertEqual(classified_lines(batch, cutoff=first.pk)[0]["concept"], "Instalación")
        self.assertEqual(classified_lines(batch)[0]["concept"], "Reconexión")
        self.assertEqual(osiptel_summary(classified_lines(batch))[0]["total"], Decimal("118"))

    def test_export_is_pinned_complete_numeric_and_formula_safe(self):
        left = self.load(legacy_bytes([legacy_row(name="<script>alert(1)</script>", description="=HYPERLINK(\"https://example.invalid\")")]))
        right = self.load(source="RVIE")
        self.client.force_login(self.accountant)
        response = self.panel(export="xlsx", legacy=left.pk, rvie=right.pk, reviews=0, q="no existe")
        self.assertEqual(response.status_code, 200)
        book = load_workbook(BytesIO(response.content), data_only=False)
        self.assertEqual(book["Conciliación"].max_row, 2)
        self.assertEqual(book["Conciliación"]["G2"].value, 118)
        self.assertEqual(book["Detalle OSIPTEL"]["E2"].data_type, "s")
        self.assertIn("SHA256", [r[0].value for r in book["Control"]])
        self.assertNotContains(self.panel(), "<script>alert(1)</script>")
        self.assertContains(self.panel(), "&lt;script&gt;")
        self.assertEqual(self.panel(export="xlsx", reviews=9999).status_code, 404)

    def test_import_form_round_trip_and_duplicate_message(self):
        self.client.force_login(self.accountant)
        raw = legacy_bytes()
        for _ in range(2):
            response = self.client.post(reverse("accounting:import"), {"issuer": self.issuer.pk, "period": "2026-09",
                "source": "OSIPTEL", "confirm": "on", "file": SimpleUploadedFile(FILENAME, raw)})
            self.assertEqual(response.status_code, 302)
        self.assertEqual(ImportBatch.objects.count(), 1)
        self.assertEqual(self.client.get(reverse("accounting:demo")).status_code, 200)
        self.assertEqual(self.client.get(reverse("accounting:rvie_template")).status_code, 200)

    def test_post_requires_csrf_and_gets_do_not_write(self):
        client = Client(enforce_csrf_checks=True)
        client.force_login(self.accountant)
        self.assertEqual(client.post(reverse("accounting:import"), {}).status_code, 403)
        self.assertEqual(client.get(reverse("accounting:import")).status_code, 200)
        self.assertEqual(ImportBatch.objects.count(), 0)


class ConcurrentImportTests(TransactionTestCase):
    @skipUnlessDBFeature("has_select_for_update")
    def test_same_file_concurrent_imports_create_one_snapshot(self):
        from concurrent.futures import ThreadPoolExecutor
        admin = User.objects.create_user(username="import_admin", role="ADMIN")
        issuer = Issuer.objects.create(code="QA-CONCUR", business_name="Empresa ficticia concurrente", ruc=RUC)
        raw = legacy_bytes()
        def worker():
            close_old_connections()
            try:
                batch, created = import_report(user=User.objects.get(pk=admin.pk), issuer=Issuer.objects.get(pk=issuer.pk),
                    period=PERIOD, source="OSIPTEL", upload=SimpleUploadedFile(FILENAME, raw))
                return batch.pk
            finally:
                connections.close_all()
        with ThreadPoolExecutor(max_workers=2) as pool:
            ids = list(pool.map(lambda _: worker(), range(2)))
        self.assertEqual(ids[0], ids[1])
        self.assertEqual(ImportBatch.objects.count(), 1)
        self.assertEqual(Document.objects.count(), 1)
