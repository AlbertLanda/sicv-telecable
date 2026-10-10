from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from threading import Barrier
from unittest import skipUnless
import uuid

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import close_old_connections, connection, connections
from django.test import Client, TransactionTestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.customers.models import Customer
from apps.organization.models import Branch, Office
from apps.payments.models import CashEntry, Charge, Issuer, Payment, PaymentAllocation, ReceiptSequence
from apps.payments.tests.base import PaymentsTestCase
from apps.payments.tests.test_cash_sessions import CashFixture
from .models import CompanyAccess, RevenueRule
from .revenue import LIMA, RevenueReport, alert_state, period_dates, save_rule
from .revenue_views import export_report


class RevenueFixture(CashFixture):
    def setup_revenue(self):
        self.setup_cash()
        self.admin = get_user_model().objects.create_user(username="revenue-admin", role="ADMIN", branch=self.branch)
        CompanyAccess.objects.create(user=self.reviewer, issuer=self.issuer, updated_by=self.admin)
        self.foreign = Issuer.objects.create(code="REVENUE-OTHER", business_name="Empresa ajena", ruc="20000000002")
        self.other_sequence = ReceiptSequence.objects.create(code="REVENUE-OTHER", series="OTHER", label="Otra", issuer=self.foreign)

    def report(self, **kwargs):
        values = dict(user=self.reviewer, period="MONTH", anchor=timezone.localdate(timezone=LIMA))
        values.update(kwargs)
        return RevenueReport(**values)

    def rule(self, **kwargs):
        values = dict(actor=self.admin, issuer_id=self.issuer.pk, period="MONTH", limit="100.00", warning_percent=80,
                      enabled=True, reason="Control administrativo de prueba", expected_pk=0, request_key=uuid.uuid4())
        values.update(kwargs)
        return save_rule(**values)


class RevenueTests(RevenueFixture, PaymentsTestCase):
    def setUp(self):
        super().setUp()
        self.setup_revenue()

    def test_calendar_periods_and_leap_year(self):
        self.assertEqual(period_dates("DAY", date(2024, 2, 16)), (date(2024, 2, 16), date(2024, 2, 16)))
        self.assertEqual(period_dates("FORTNIGHT", date(2024, 2, 15)), (date(2024, 2, 1), date(2024, 2, 15)))
        self.assertEqual(period_dates("FORTNIGHT", date(2024, 2, 16)), (date(2024, 2, 16), date(2024, 2, 29)))
        self.assertEqual(period_dates("MONTH", date(2025, 2, 16)), (date(2025, 2, 1), date(2025, 2, 28)))
        self.assertEqual(period_dates("YEAR", date(2026, 10, 10)), (date(2026, 1, 1), date(2026, 12, 31)))
        for period, anchor in [("BAD", date(2026, 1, 1)), ("DAY", date(1999, 1, 1))]:
            with self.assertRaises(ValidationError):
                period_dates(period, anchor)

    def test_counts_cash_and_noncash_once_without_allocations_duplication(self):
        payment = self.payment()
        for value in (30, 50):
            charge = Charge.objects.create(customer=self.customer, amount=value, concept="OTHER", due_date=timezone.localdate())
            PaymentAllocation.objects.create(payment=payment, charge=charge, amount=value)
        self.payment(amount=20, method="YAPE", reference="000042")
        summary = self.report().summary()
        self.assertEqual(summary["totals"], {"total": Decimal("100"), "count": 2, "fallback": 2})
        self.assertEqual(sum(r["total"] for r in summary["breakdown"]), 100)
        self.assertEqual(sum(r["total"] for r in summary["daily"]), 100)

    def test_pending_voided_missing_date_future_and_foreign_not_in_total(self):
        self.payment(amount=11)
        self.payment(amount=22, settled=False)
        voided = self.payment(amount=33)
        voided.void(user=self.admin, reason="Prueba")
        missing = self.payment(amount=44)
        Payment.objects.filter(pk=missing.pk).update(paid_at=None)
        now = timezone.now()
        self.payment(amount=55, paid_at=now + timedelta(minutes=1))
        self.payment(amount=66, series=self.other_sequence)
        report = self.report(now=now)
        self.assertEqual(report.summary()["totals"]["total"], 11)
        self.assertEqual((report.missing_date, report.future), (1, 1))

    def test_real_payment_date_lima_midnight_and_month_boundaries(self):
        # 05:00 UTC = medianoche en Lima. La fecha de emisión no decide.
        self.payment(amount=1, paid_at=datetime(2026, 3, 1, 4, 59, tzinfo=timezone.get_fixed_timezone(0)))
        payment = self.payment(amount=2, paid_at=datetime(2026, 3, 1, 5, 0, tzinfo=timezone.get_fixed_timezone(0)))
        payment.receipt.issued_at = datetime(2026, 2, 28, 10, tzinfo=LIMA)
        payment.receipt.save(update_fields=["issued_at"])
        self.payment(amount=4, paid_at=datetime(2026, 4, 1, 5, 0, tzinfo=timezone.get_fixed_timezone(0)))
        result = self.report(anchor=date(2026, 3, 8), now=datetime(2026, 5, 1, tzinfo=LIMA)).summary()
        self.assertEqual(result["totals"]["total"], 2)
        self.assertEqual(result["daily"][0]["day"], date(2026, 3, 1))

    def test_cash_issuer_preserved_when_series_changes_and_no_double_count_deposit(self):
        session = self.open()
        payment = self.payment()
        self.movement(session, kind="DEPOSIT", amount=50, bank="Banco prueba", account="00012345", reference="OP-TEST")
        self.movement(session, kind="INCOME", amount=10)
        self.movement(session, kind="GUARANTEE", amount=20)
        self.sequence.issuer = self.foreign
        self.sequence.save(update_fields=["issuer"])
        summary = self.report().summary()
        self.assertEqual(summary["totals"], {"total": Decimal("80"), "count": 1, "fallback": 0})
        self.assertEqual(self.report().detail().get().pk, payment.pk)
        self.assertEqual(self.report(user=self.admin, issuer=self.foreign).summary()["totals"]["total"], 0)

    def test_cash_without_issuer_does_not_fall_back_to_new_series_issuer(self):
        self.sequence.issuer = None
        self.sequence.save(update_fields=["issuer"])
        self.open()
        self.payment()
        self.sequence.issuer = self.issuer
        self.sequence.save(update_fields=["issuer"])
        self.assertEqual(self.report().summary()["totals"]["count"], 0)

    def test_changed_ruc_requires_review_instead_of_mixing_identity(self):
        self.open()
        self.payment()
        self.issuer.ruc = "20000000003"
        self.issuer.save(update_fields=["ruc"])
        report = self.report()
        self.assertEqual(report.ruc_conflicts, 1)
        self.assertEqual(report.summary()["totals"]["total"], 0)

    def test_legacy_attribution_is_explicit_and_tracks_current_series(self):
        self.payment()
        self.assertEqual(self.report().summary()["totals"]["fallback"], 1)
        self.sequence.issuer = self.foreign
        self.sequence.save(update_fields=["issuer"])
        self.assertEqual(self.report().summary()["totals"]["count"], 0)
        self.assertEqual(self.report(user=self.admin, issuer=self.foreign).summary()["totals"]["total"], 80)

    def test_branch_filter_does_not_reduce_company_alert_total(self):
        other = Branch.objects.create(code="RV-B", name="Sede segunda")
        office = Office.objects.create(branch=other, code="RV-O", name="Oficina segunda")
        self.payment(amount=70)
        self.payment(amount=30, branch=other, office=office)
        self.rule()
        report = self.report(branch=self.branch, office=self.office).summary()
        self.assertEqual(report["totals"]["total"], 70)
        self.assertEqual(report["companies"][0]["company_total"], 100)
        self.assertEqual(report["companies"][0]["alert"]["level"], "reached")
        with self.assertRaises(ValidationError):
            self.report(branch=self.branch, office=office)

    def test_null_office_and_inactive_branch_remain_in_history(self):
        self.payment(office=None)
        self.branch.is_active = False
        self.branch.save(update_fields=["is_active"])
        result = self.report().summary()
        self.assertEqual(result["totals"]["total"], 80)
        self.assertIsNone(result["breakdown"][0]["office__name"])

    def test_companies_without_payments_and_disabled_rule_are_visible(self):
        self.rule(enabled=False)
        result = self.report().summary()
        self.assertEqual(len(result["companies"]), 1)
        self.assertEqual(result["companies"][0]["alert"]["level"], "none")
        self.assertEqual(result["totals"]["total"], 0)

    def test_alert_threshold_boundaries_and_changed_ruc(self):
        rule = self.rule()
        for amount, level in [(0, "normal"), (Decimal("79.99"), "normal"), (80, "warning"), (Decimal("99.99"), "warning"), (100, "reached"), (120, "reached")]:
            self.assertEqual(alert_state(rule, Decimal(amount), self.issuer.ruc)["level"], level)
        self.assertEqual(alert_state(rule, Decimal(80), self.foreign.ruc)["level"], "review")

    def test_rule_revision_replay_stale_form_and_no_payment_side_effect(self):
        self.payment()
        key = uuid.uuid4()
        first = self.rule(request_key=key)
        self.assertEqual(self.rule(request_key=key).pk, first.pk)
        with self.assertRaises(ValidationError):
            self.rule(request_key=key, limit=200)
        with self.assertRaises(ValidationError):
            self.rule(limit=200)
        second = self.rule(expected_pk=first.pk, enabled=False)
        self.assertEqual(second.revision, 2)
        first.refresh_from_db()
        self.assertTrue(first.enabled)
        self.assertEqual(self.report().summary()["companies"][0]["alert"]["level"], "none")
        self.assertEqual(self.report().summary()["totals"]["total"], 80)
        self.assertEqual(Payment.objects.count(), 1)

    def test_rule_validation_permissions_and_explicit_grant(self):
        for values in [{"limit": "NaN"}, {"limit": "Infinity"}, {"limit": "1.001"}, {"limit": 0}, {"limit": -1},
                {"limit": "100000000000000"}, {"warning_percent": 0}, {"warning_percent": 101}, {"warning_percent": "1.5"},
                {"period": "BAD"}, {"reason": ""}, {"expected_pk": -1}, {"request_key": "bad"}, {"enabled": "yes"}]:
            with self.subTest(values=values), self.assertRaises(ValidationError):
                self.rule(**values)
        with self.assertRaises(PermissionDenied):
            self.rule(actor=self.reviewer)
        self.reviewer.user_permissions.add(Permission.objects.get(content_type__app_label="accounting", codename="manage_revenue_rules"))
        reviewer = type(self.reviewer).objects.get(pk=self.reviewer.pk)
        self.rule(actor=reviewer)
        with self.assertRaises(PermissionDenied):
            self.rule(actor=reviewer, issuer_id=self.foreign.pk)

    def test_company_access_revocation_blocks_data_exports_and_rule_urls(self):
        self.payment()
        self.payment(series=self.other_sequence)
        self.client.force_login(self.reviewer)
        url = reverse("accounting:revenue")
        params = {"period": "MONTH", "anchor": timezone.localdate().isoformat()}
        response = self.client.get(url, params)
        self.assertContains(response, self.issuer.business_name)
        self.assertNotContains(response, self.foreign.business_name)
        self.assertEqual(self.client.get(url, {**params, "issuer": self.foreign.pk}).status_code, 400)
        self.assertEqual(self.client.get(reverse("accounting:revenue_rules", args=[self.foreign.pk])).status_code, 404)
        CompanyAccess.objects.filter(user=self.reviewer).update(enabled=False)
        self.assertEqual(self.client.get(url, {**params, "issuer": self.issuer.pk, "export": "xlsx"}).status_code, 400)
        self.assertEqual(self.client.get(reverse("accounting:revenue_rules", args=[self.issuer.pk])).status_code, 404)
        self.assertContains(self.client.get(url), "No tiene empresas activas")

    def test_atc_inactive_and_ungranted_users_cannot_see_dashboard(self):
        for user in (self.actor, get_user_model().objects.create_user(username="rv-tech", role="TECHNICIAN")):
            self.client.force_login(user)
            self.assertEqual(self.client.get(reverse("accounting:revenue")).status_code, 403)
            with self.assertRaises(PermissionDenied):
                self.report(user=user)
        self.reviewer.is_active = False
        self.reviewer.save()
        with self.assertRaises(PermissionDenied):
            self.report()

    def test_filters_reject_invalid_period_dates_and_mismatched_office(self):
        other = Branch.objects.create(code="RV-C", name="Otra sede")
        self.payment(branch=other)
        self.client.force_login(self.reviewer)
        base = {"period": "MONTH", "anchor": timezone.localdate().isoformat()}
        for params in [{"period": "BAD"}, {"anchor": "2026-02-30"}, {"anchor": "1999-01-01"},
                       {"issuer": "x"}, {"branch": other.pk, "office": self.office.pk}]:
            with self.subTest(params=params):
                self.assertEqual(self.client.get(reverse("accounting:revenue"), {**base, **params}).status_code, 400)

    def test_rule_web_history_csrf_no_self_granted_configuration(self):
        url = reverse("accounting:revenue_rules", args=[self.issuer.pk])
        self.client.force_login(self.admin)
        data = dict(period="MONTH", expected_pk=0, request_key=uuid.uuid4(), enabled="on", limit="100", warning_percent=80, reason="Prueba QA")
        self.assertContains(self.client.get(url), "Sin alerta configurada")
        self.assertEqual(self.client.post(url, data).status_code, 302)
        self.assertContains(self.client.get(url), "Prueba QA")
        self.assertEqual(self.client.post(url, {**data, "request_key": uuid.uuid4()}).status_code, 400)
        self.client.force_login(self.reviewer)
        self.assertContains(self.client.get(url), "Administración configura")
        self.assertEqual(self.client.post(url, data).status_code, 403)
        client = Client(enforce_csrf_checks=True)
        client.force_login(self.admin)
        self.assertEqual(client.post(url, data).status_code, 403)

    def test_pagination_total_export_and_literal_spreadsheet_text(self):
        self.customer.code = "=HYPERLINK(1)"
        self.customer.save(update_fields=["code"])
        for _ in range(41):
            self.payment(amount=1)
        self.client.force_login(self.reviewer)
        params = {"period": "MONTH", "anchor": timezone.localdate().isoformat(), "page": 2}
        response = self.client.get(reverse("accounting:revenue"), params)
        self.assertEqual(response.context["totals"]["total"], 41)
        self.assertEqual(len(response.context["page"]), 1)
        self.assertIn("no-store", response.headers["Cache-Control"])
        response = self.client.get(reverse("accounting:revenue"), {**params, "export": "xlsx"})
        book = load_workbook(BytesIO(b"".join(response.streaming_content)))
        self.assertEqual(book["Control"]["D6"].value, 41)
        self.assertEqual(book["Cobros"].max_row, 42)
        self.assertEqual(book["Cobros"]["J2"].data_type, "s")
        self.assertEqual(book["Cobros"]["J2"].value, "=HYPERLINK(1)")
        self.assertEqual(book["Por empresa"]["E2"].value, 41)
        self.assertEqual(book["Por oficina y medio"]["G2"].value, 41)

    def test_duplicate_ruc_blocks_alert_without_exposing_other_company(self):
        self.payment()
        self.rule()
        self.foreign.ruc = self.issuer.ruc
        self.foreign.save(update_fields=["ruc"])
        self.payment(amount=500, series=self.other_sequence)
        result = self.report().summary()
        self.assertEqual(result["totals"]["total"], 80)
        self.assertEqual(len(result["companies"]), 1)
        self.assertEqual(result["companies"][0]["alert"]["level"], "review")
        with self.assertRaises(ValidationError):
            self.rule()
        self.issuer.ruc = ""
        self.issuer.save(update_fields=["ruc"])
        with self.assertRaises(ValidationError):
            self.rule()

    def test_same_named_offices_are_not_merged(self):
        other = Office.objects.create(branch=self.branch, code="RV-SAME", name=self.office.name)
        self.payment(amount=10)
        self.payment(amount=20, office=other)
        rows = self.report().summary()["breakdown"]
        self.assertEqual(len(rows), 2)
        self.assertEqual({row["office_id"] for row in rows}, {self.office.pk, other.pk})

    def test_summary_query_count_does_not_grow_per_payment(self):
        self.payment()
        report = self.report()
        with CaptureQueriesContext(connection) as first:
            report.summary()
            list(report.detail())
        for _ in range(8):
            self.payment()
        report = self.report()
        with CaptureQueriesContext(connection) as second:
            report.summary()
            list(report.detail())
        self.assertEqual(len(first), len(second))


@skipUnless(connection.vendor == "postgresql", "Bloqueos verificados en CI PostgreSQL")
class RevenueRuleConcurrencyTests(RevenueFixture, TransactionTestCase):
    def setUp(self):
        self.branch = Branch.objects.create(code="RV-CON", name="Sede de pruebas")
        self.customer = Customer.objects.create(code="RV-CON", branch=self.branch, first_name="Prueba")
        self.setup_revenue()

    def race(self, *actions):
        barrier = Barrier(len(actions))
        def run(action):
            close_old_connections()
            try:
                with connection.cursor() as cursor:
                    cursor.execute("SET lock_timeout = '5s'")
                barrier.wait(timeout=5)
                try:
                    return action().pk
                except ValidationError:
                    return "rejected"
            finally:
                connections.close_all()
        with ThreadPoolExecutor(max_workers=len(actions)) as pool:
            futures = [pool.submit(run, action) for action in actions]
            return [future.result(timeout=15) for future in futures]

    def test_same_request_is_idempotent(self):
        key = uuid.uuid4()
        results = self.race(lambda: self.rule(request_key=key), lambda: self.rule(request_key=key))
        self.assertEqual(results[0], results[1])
        self.assertEqual(RevenueRule.objects.count(), 1)

    def test_two_editors_cannot_overwrite_an_unseen_version(self):
        results = self.race(self.rule, lambda: self.rule(limit=200))
        self.assertEqual(results.count("rejected"), 1)
        self.assertEqual(RevenueRule.objects.count(), 1)
