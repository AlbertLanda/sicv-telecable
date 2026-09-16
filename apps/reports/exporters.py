"""Las cuatro salidas del reporte: PDF, Excel, Word y HTML.

Todas leen las mismas filas y la misma lista de columnas que arma
`materials.build_report()`. Aquí solo se decide cómo se dibuja cada una, nunca
qué entra: si un formato pudiera recortar por su cuenta, el Excel y el PDF del
mismo periodo acabarían sumando distinto y nadie sabría cuál de los dos vale.

La cabecera es la misma en los tres formatos de papel -hora de impresión a la
izquierda, título al centro, sede a la derecha- porque es la que identifica la
hoja cuando alguien la encuentra suelta encima de una mesa.
"""

from io import BytesIO

from django.utils import timezone
from django.utils.formats import number_format

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


# El azul corporativo del SICV, el mismo del sistema visual de pantalla. Se
# repite aquí en vez de importarse del CSS porque son dos medios distintos:
# uno se sirve al navegador y el otro se incrusta en el papel.
TC_BLUE = "0F4C7F"

# Fecha y hora al minuto. Lo comparten la emision, la atencion y la hora de
# impresion de la cabecera, porque son la misma clase de dato y leerlos con
# precisiones distintas en la misma hoja solo hace dudar de si significan
# cosas distintas.
#
# La hora no sobra en un reporte de materiales: dos ordenes del mismo dia en
# el mismo domicilio solo se distinguen por ella, y es lo que permite
# reconstruir en que visita salio cada material. Los segundos si sobran -nadie
# cuadra un almacen al segundo- y solo alargan la celda.
MOMENT_FORMAT = "%d/%m/%Y %H:%M"


def printed_at():
    """La hora de impresión que sale en la esquina de la hoja."""
    return timezone.localtime().strftime(MOMENT_FORMAT)


def branch_label(branch):
    return branch.name if branch is not None else "Todas las sedes"


def cell_text(row, column):
    """El valor de una celda, ya convertido a texto imprimible.

    Las fechas se formatean aquí y no en cada salida para que el PDF, el Word
    y el HTML no puedan discrepar en cómo escriben el mismo día. La cantidad
    arrastra su unidad: «12.00» sin saber si son metros o unidades no sirve
    para validar un retiro.
    """
    value = row.get(column["key"])

    if value in (None, ""):
        return ""

    if column["key"] in ("issued_on", "attended_on"):
        return timezone.localtime(value).strftime(MOMENT_FORMAT)

    if column["key"] == "quantity":
        # Con el formato del idioma activo -«45,50» en español-, que es como
        # la pantalla ya escribe cualquier cifra del sistema. El reporte es un
        # mismo documento con cuatro entregas: si el papel dijera «45.50» y la
        # pantalla «45,50», quien compara la impresión con la consulta que
        # acaba de hacer tendría que decidir cuál de las dos cifras vale.
        cantidad = number_format(value, decimal_pos=2)

        return f"{cantidad} {row.get('unit', '')}".strip()

    return str(value)


def filename(report, extension):
    """Nombre del archivo que se descarga.

    Lleva el periodo dentro: quien guarda tres reportes en la misma carpeta
    los distingue sin abrirlos, que es justo lo que no pasaba cuando todos se
    llamaban «reporte».
    """
    desde = report["date_from"].strftime("%Y%m%d")
    hasta = report["date_to"].strftime("%Y%m%d")

    return f"materiales_{desde}_{hasta}.{extension}"


# ---------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------

def render_pdf(report, buffer):
    """El reporte en A4 apaisado.

    Apaisado porque son trece columnas: en vertical, «Dirección» y «Abonado»
    se parten en cuatro líneas cada una y la hoja deja de poder leerse de un
    vistazo, que es para lo único que se imprime.
    """
    page = landscape(A4)
    margin = 10 * mm

    documento = SimpleDocTemplate(
        buffer,
        pagesize=page,
        leftMargin=margin,
        rightMargin=margin,
        topMargin=margin,
        bottomMargin=margin,
        title=report["title"],
    )

    ancho = page[0] - margin * 2

    encabezado = ParagraphStyle(
        "encabezado",
        fontName="Helvetica-Bold",
        fontSize=10,
        textColor=colors.black,
    )
    normal = ParagraphStyle(
        "normal",
        fontName="Helvetica",
        fontSize=7,
        leading=8.5,
    )
    cabecera_celda = ParagraphStyle(
        "cabecera_celda",
        fontName="Helvetica-Bold",
        fontSize=7,
        leading=8.5,
        textColor=colors.white,
    )

    # La franja de identificación: hora, título y sede en una sola fila de
    # tres celdas, que es como se reparte el ancho sin calcular posiciones.
    franja = Table(
        [[
            Paragraph(printed_at(), ParagraphStyle("hora", fontSize=7.5)),
            Paragraph(
                f"<para align='center'>{report['title']}</para>", encabezado
            ),
            Paragraph(
                f"<para align='right'>{branch_label(report['branch'])}</para>",
                encabezado,
            ),
        ]],
        colWidths=[ancho * 0.25, ancho * 0.5, ancho * 0.25],
    )
    franja.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))

    columnas = report["columns"]
    peso_total = sum(columna["width"] for columna in columnas)
    anchos = [ancho * columna["width"] / peso_total for columna in columnas]

    datos = [[
        Paragraph(columna["label"], cabecera_celda) for columna in columnas
    ]]

    for fila in report["rows"]:
        datos.append([
            Paragraph(cell_text(fila, columna), normal) for columna in columnas
        ])

    tabla = Table(datos, colWidths=anchos, repeatRows=1)
    estilo = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{TC_BLUE}")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D9E4EF")),
        ("LINEBELOW", (0, 0), (-1, 0), 0.8, colors.HexColor(f"#{TC_BLUE}")),
    ]

    # Filas alternas. En trece columnas el ojo pierde la línea a mitad de la
    # hoja, y la banda es lo que la devuelve sin añadir tinta al papel.
    for indice in range(1, len(datos)):
        if indice % 2 == 0:
            estilo.append(
                ("BACKGROUND", (0, indice), (-1, indice), colors.HexColor("#F5F9FD"))
            )

    tabla.setStyle(TableStyle(estilo))

    total = Paragraph(
        f"<para align='right'><b>Total: {report['total']}</b></para>",
        ParagraphStyle("total", fontName="Helvetica-Bold", fontSize=8),
    )

    documento.build([franja, Spacer(1, 8), tabla, Spacer(1, 8), total])

    return filename(report, "pdf")


# ---------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------

def _tercios(total):
    """Reparte las columnas en izquierda, centro y derecha para la franja.

    Se calcula del número de columnas en vez de escribirse a mano -«A:D»,
    «E:I», «J:M»- para que añadir o quitar una columna no deje el título
    descentrado sin que nadie lo note.
    """
    lado = max(1, total // 3)

    if total - lado * 2 < 1:
        # Con tan pocas columnas no hay tres bloques que repartir: todo a una.
        return (1, total), None, None

    return (1, lado), (lado + 1, total - lado), (total - lado + 1, total)


def _franja_excel(hoja, bloque, texto, fuente, alineacion):
    """Escribe un bloque de la franja, combinado sobre sus columnas."""
    if bloque is None:
        return

    desde, hasta = bloque

    if hasta > desde:
        hoja.merge_cells(
            start_row=1, start_column=desde, end_row=1, end_column=hasta
        )

    celda = hoja.cell(row=1, column=desde, value=texto)
    celda.font = fuente
    celda.alignment = Alignment(horizontal=alineacion, vertical="center")


def render_excel(report, buffer):
    """El reporte como `.xlsx` real, no como tabla HTML renombrada.

    Se entrega con las columnas dimensionadas, la cabecera congelada y el
    autofiltro puesto. Es el formato que logística usa para cuadrar, y cuadrar
    significa ordenar por material y filtrar por acción: entregarlo sin eso
    obliga a rehacer a mano lo mismo en cada descarga.
    """
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Materiales"

    columnas = report["columns"]
    ultima = get_column_letter(len(columnas))

    azul = PatternFill("solid", fgColor=TC_BLUE)
    blanco = Font(bold=True, color="FFFFFF", size=9)
    borde = Border(
        left=Side(style="thin", color="D9E4EF"),
        right=Side(style="thin", color="D9E4EF"),
        top=Side(style="thin", color="D9E4EF"),
        bottom=Side(style="thin", color="D9E4EF"),
    )

    # La misma franja de una línea que la hoja en pantalla: hora a la
    # izquierda, título al centro y sede a la derecha.
    #
    # Se reparte en tres bloques combinados en vez de tres filas apiladas para
    # que sea la misma cabecera y no una versión propia del Excel. Tres filas
    # dejaban además la tabla empezando en la quinta, y quien abre la hoja para
    # cuadrar quiere ver datos sin desplazarse.
    #
    # El periodo desaparece de la franja pero no del archivo: viaja en el
    # nombre (`materiales_20260915_20260915.xlsx`), que es lo que queda a la
    # vista cuando alguien lo archiva o lo reenvía.
    izquierda, centro, derecha = _tercios(len(columnas))

    _franja_excel(
        hoja, izquierda, printed_at(),
        Font(size=9, color="5A6B80"), "left",
    )
    _franja_excel(
        hoja, centro, report["title"],
        Font(bold=True, size=14, color="17355F"), "center",
    )
    _franja_excel(
        hoja, derecha, branch_label(report["branch"]),
        Font(bold=True, size=11, color="17355F"), "right",
    )

    hoja.row_dimensions[1].height = 22

    fila_cabecera = 3

    for indice, columna in enumerate(columnas, start=1):
        celda = hoja.cell(row=fila_cabecera, column=indice, value=columna["label"])
        celda.fill = azul
        celda.font = blanco
        celda.border = borde
        celda.alignment = Alignment(horizontal="center", vertical="center")
        hoja.column_dimensions[get_column_letter(indice)].width = columna["width"] + 4

    for desplazamiento, fila in enumerate(report["rows"], start=1):
        for indice, columna in enumerate(columnas, start=1):
            valor = fila.get(columna["key"])

            # La cantidad viaja como número, no como texto: es la única
            # columna que alguien va a sumar en la propia hoja, y un texto
            # «12.00 Metro» no se suma.
            #
            # La unidad no se pierde por eso: se pinta con el formato de la
            # celda, así que la hoja sigue leyéndose «12.00 Metro» y la suma
            # sigue funcionando. Poner la unidad en una columna aparte habría
            # dado un Excel con catorce columnas frente a las trece del PDF,
            # y el reporte dejaría de ser el mismo documento.
            if columna["key"] == "quantity":
                valor = float(valor) if valor is not None else None

            else:
                valor = cell_text(fila, columna)

            celda = hoja.cell(
                row=fila_cabecera + desplazamiento, column=indice, value=valor
            )

            if columna["key"] == "quantity" and fila.get("unit"):
                celda.number_format = f'0.00" {fila["unit"]}"'

            celda.font = Font(size=9)
            celda.border = borde
            celda.alignment = Alignment(
                horizontal="right" if columna.get("numeric") else "left",
                vertical="top",
                wrap_text=columna["key"] in ("address", "customer_name"),
            )

    # El total, abajo a la derecha, sobre el mismo bloque de columnas que
    # ocupa la sede en la franja de arriba. La hoja queda así leyéndose por su
    # borde derecho -sede arriba, total abajo-, igual que en la pantalla, el
    # PDF y el Word. En la columna A quedaba pegado al primer dato de la
    # tabla, donde se confunde con una fila más.
    fila_total = fila_cabecera + len(report["rows"]) + 1
    desde_total, hasta_total = derecha or (1, len(columnas))

    if hasta_total > desde_total:
        hoja.merge_cells(
            start_row=fila_total,
            start_column=desde_total,
            end_row=fila_total,
            end_column=hasta_total,
        )

    celda_total = hoja.cell(
        row=fila_total, column=desde_total, value=f"Total: {report['total']}"
    )
    celda_total.font = Font(bold=True, size=10, color="17355F")
    celda_total.alignment = Alignment(horizontal="right", vertical="center")

    # Una raya encima lo separa de la última fila de datos: va pegado a ella,
    # como el pie de la tabla en pantalla, y sin la raya se leería como una
    # fila más de material.
    for columna in range(desde_total, hasta_total + 1):
        hoja.cell(row=fila_total, column=columna).border = Border(
            top=Side(style="thin", color="8AA4C0")
        )

    hoja.freeze_panes = hoja.cell(row=fila_cabecera + 1, column=1)
    hoja.auto_filter.ref = (
        f"A{fila_cabecera}:{ultima}{fila_cabecera + len(report['rows'])}"
    )

    libro.save(buffer)

    return filename(report, "xlsx")


# ---------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------

def _sombrear(celda, color):
    """Pinta el fondo de una celda de Word.

    `python-docx` no expone el relleno de una celda, así que se añade el
    elemento `w:shd` a mano. Es la única forma de que la cabecera del Word se
    lea igual que la de los otros tres formatos.
    """
    relleno = OxmlElement("w:shd")
    relleno.set(qn("w:val"), "clear")
    relleno.set(qn("w:color"), "auto")
    relleno.set(qn("w:fill"), color)

    celda._tc.get_or_add_tcPr().append(relleno)


def render_word(report, buffer):
    """El reporte como `.docx` real, apaisado como el PDF.

    Word es el formato de quien tiene que anotar sobre el reporte antes de
    firmarlo -logística marca a mano lo que no cuadra-, así que la tabla se
    entrega editable y no como una imagen del PDF.
    """
    documento = Document()

    seccion = documento.sections[0]

    # A4 apaisado, el mismo papel del PDF. Word abre en Carta por defecto, y
    # dos formatos del mismo reporte que salen en hojas distintas no se pueden
    # archivar juntos ni fotocopiar con el mismo ajuste.
    #
    # El tamaño se fija a mano además de la orientación: cambiar `orientation`
    # no toca el ancho ni el alto, así que la tabla saldría cortada por la
    # derecha sobre una hoja que sigue siendo vertical.
    seccion.orientation = WD_ORIENT.LANDSCAPE
    seccion.page_width = Pt(841.89)
    seccion.page_height = Pt(595.28)
    seccion.left_margin = seccion.right_margin = Pt(28)
    seccion.top_margin = seccion.bottom_margin = Pt(28)

    # La franja de una línea, como la pantalla y el Excel: hora a la
    # izquierda, título al centro y sede a la derecha.
    #
    # Se consigue con dos tabulaciones y no con una tabla invisible de tres
    # celdas: la tabla la heredaría el cursor de quien edita el documento -y
    # Word es justo el formato de quien anota encima- y acabaría escribiendo
    # dentro de la cabecera sin querer.
    ancho = seccion.page_width - seccion.left_margin - seccion.right_margin

    franja = documento.add_paragraph()
    tabulaciones = franja.paragraph_format.tab_stops
    tabulaciones.add_tab_stop(int(ancho / 2), WD_TAB_ALIGNMENT.CENTER)
    tabulaciones.add_tab_stop(ancho, WD_TAB_ALIGNMENT.RIGHT)

    hora = franja.add_run(printed_at())
    hora.font.size = Pt(8.5)
    hora.font.color.rgb = RGBColor(0x6F, 0x81, 0x97)

    marca = franja.add_run(f"	{report['title']}")
    marca.bold = True
    marca.font.size = Pt(14)
    marca.font.color.rgb = RGBColor(0x17, 0x35, 0x5F)

    sede = franja.add_run(f"	{branch_label(report['branch'])}")
    sede.bold = True
    sede.font.size = Pt(11)
    sede.font.color.rgb = RGBColor(0x17, 0x35, 0x5F)

    columnas = report["columns"]
    tabla = documento.add_table(rows=1, cols=len(columnas))
    tabla.style = "Table Grid"

    for celda, columna in zip(tabla.rows[0].cells, columnas):
        # Cabecera azul con letra blanca, la misma de la pantalla, el PDF y el
        # Excel. Era el unico de los cuatro formatos donde la cabecera se
        # confundia con la primera fila de datos.
        _sombrear(celda, TC_BLUE)

        etiqueta = celda.paragraphs[0].add_run(columna["label"])
        etiqueta.bold = True
        etiqueta.font.size = Pt(7)
        etiqueta.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for indice, fila in enumerate(report["rows"], start=1):
        celdas = tabla.add_row().cells

        for celda, columna in zip(celdas, columnas):
            # Filas alternas: en trece columnas el ojo pierde la linea a mitad
            # de la hoja, y la banda es lo que la devuelve.
            if indice % 2 == 0:
                _sombrear(celda, "F5F9FD")

            escrito = celda.paragraphs[0].add_run(cell_text(fila, columna))
            escrito.font.size = Pt(6.5)

    total = documento.add_paragraph()
    total.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    suma = total.add_run(f"Total: {report['total']}")
    suma.bold = True
    suma.font.size = Pt(10)
    suma.font.color.rgb = RGBColor(0x17, 0x35, 0x5F)

    documento.save(buffer)

    return filename(report, "docx")


# Qué función atiende cada formato. El HTML no está aquí: no produce un
# archivo que descargar sino una pantalla, y la resuelve la vista con su
# plantilla.
RENDERERS = {
    "PDF": render_pdf,
    "EXCEL": render_excel,
    "WORD": render_word,
}

CONTENT_TYPES = {
    "PDF": "application/pdf",
    "EXCEL": (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    ),
    "WORD": (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ),
}


def render(report, export_format):
    """Dibuja el reporte en el formato pedido y devuelve (buffer, nombre)."""
    buffer = BytesIO()
    nombre = RENDERERS[export_format](report, buffer)
    buffer.seek(0)

    return buffer, nombre
