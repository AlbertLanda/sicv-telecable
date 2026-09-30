"""El cierre de caja: el consolidado de emisión que usa Contabilidad.

Todo el escenario es sintético: empresas, talonarios, abonados y montos
inventados para la prueba. Ningún dato sale del padrón real.
"""

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from unittest import mock

from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from openpyxl import load_workbook

from apps.accounts.models import User
from apps.customers.models import Customer
from apps.organization.context_processors import (
    ACTIVE_BRANCH_SESSION_KEY,
    ACTIVE_OFFICE_SESSION_KEY,
)
from apps.organization.models import Branch, Office
from apps.payments.models import (
    Issuer,
    OfficeSequence,
    Payment,
    Receipt,
    ReceiptSequence,
)
from apps.reports import cash_closing
from apps.reports.cash_closing import build_cash_closing, series_label, series_name

from .pdf_text import pdf_text


FACTURA = ReceiptSequence.SunatCode.FACTURA
BOLETA = ReceiptSequence.SunatCode.BOLETA
RECIBO = ReceiptSequence.SunatCode.RECIBO_SERVICIO


def moment(day, hour=10, minute=0):
    """Una hora de Lima, que es la que usa el recorte por fecha."""
    return timezone.make_aware(datetime(2026, 9, day, hour, minute))


class CashClosingTestCase(TestCase):
    def setUp(self):
        self.branch = Branch.objects.create(code="SEDA", name="Sede Alfa")
        self.other_branch = Branch.objects.create(code="SEDB", name="Sede Beta")

        self.office = Office.objects.create(
            branch=self.branch, code="SEDA-CAJA1", name="Caja Uno"
        )
        self.second_office = Office.objects.create(
            branch=self.branch, code="SEDA-CAJA2", name="Caja Dos"
        )
        self.other_office = Office.objects.create(
            branch=self.other_branch, code="SEDB-CAJA1", name="Caja Beta"
        )

        self.issuer = Issuer.objects.create(
            code="TSTA", business_name="EMPRESA ALFA S.A.C.", ruc="20000000011"
        )
        self.other_issuer = Issuer.objects.create(
            code="TSTB", business_name="EMPRESA BETA E.I.R.L.", ruc="20000000022"
        )

        self.factura = self.make_sequence(
            "TST-F901", "F901", "F:F901 - ALFA", FACTURA, self.issuer
        )
        self.boleta = self.make_sequence(
            "TST-B902", "B902", "B:B902 - BETA", BOLETA, self.other_issuer
        )
        # La misma serie impresa en dos talonarios de marca distinta, como
        # los tres «S010» del padrón real.
        self.recibo_gamma = self.make_sequence(
            "TST-S903-G", "S903", "S:S903 - GAMMA", RECIBO, self.issuer
        )
        self.recibo_delta = self.make_sequence(
            "TST-S903-D", "S903", "S:S903 - DELTA", RECIBO, self.other_issuer
        )

        self.customer = Customer.objects.create(
            code="CLI-T001",
            branch=self.branch,
            document_type=Customer.DocumentType.DNI,
            document_number="10000001",
            first_name="Abonado",
            paternal_surname="Prueba",
            maternal_surname="Uno",
        )

        self.cashier = User.objects.create_user(
            username="cajero_prueba",
            password="test1234",
            first_name="Cajero",
            last_name="Prueba",
            role=User.Role.ATC,
            branch=self.branch,
        )
        self.second_cashier = User.objects.create_user(
            username="cajero_dos",
            password="test1234",
            role=User.Role.ATC,
            branch=self.branch,
        )
        self.other_cashier = User.objects.create_user(
            username="cajero_beta",
            password="test1234",
            role=User.Role.ATC,
            branch=self.other_branch,
        )
        self.accountant = User.objects.create_user(
            username="contable_prueba",
            password="test1234",
            role=User.Role.ACCOUNTING,
            branch=self.branch,
        )

        self.numbers = {}

    def make_sequence(self, code, series, label, sunat_code, issuer):
        return ReceiptSequence.objects.create(
            code=code,
            series=series,
            label=label,
            sunat_code=sunat_code,
            issuer=issuer,
        )

    def emit(
        self,
        amount,
        sequence,
        *,
        issued_at=None,
        status=Payment.Status.REGISTERED,
        office=None,
        branch=None,
        user=None,
        method=Payment.Method.CASH,
        reference="",
    ):
        """Un cobro con su comprobante, directo a la base."""
        issued_at = issued_at or moment(10)
        number = self.numbers.get(sequence.pk, 0) + 1
        self.numbers[sequence.pk] = number

        payment = Payment.objects.create(
            customer=self.customer,
            amount=Decimal(amount),
            method=method,
            reference=reference,
            branch=branch or self.branch,
            office=office or self.office,
            received_by=user or self.cashier,
            received_at=issued_at,
            paid_at=issued_at if status == Payment.Status.REGISTERED else None,
            status=status,
        )

        return Receipt.objects.create(
            payment=payment,
            sequence=sequence,
            series=sequence.series,
            number=number,
            issued_at=issued_at,
        )

    def build(self, **kwargs):
        options = {
            "branch": self.branch,
            "date_from": date(2026, 9, 1),
            "date_to": date(2026, 9, 30),
        }
        options.update(kwargs)

        return build_cash_closing(**options)

    def amounts(self, report):
        """Texto -> monto de la hoja, para comprobar filas por su etiqueta."""
        return {line["text"]: line["amount"] for line in report["lines"]}


class SeriesLabelTests(CashClosingTestCase):
    def test_brand_from_the_label_and_number_without_the_letter(self):
        self.assertEqual(series_label(self.factura), "Serie ALFA - 901")
        self.assertEqual(series_name(self.factura), "ALFA - 901")

    def test_a_series_that_is_not_letter_and_number_is_kept_whole(self):
        cobrador = self.make_sequence(
            "TST-ZZ1", "ZZ1", "B: ZZ1 - COBRADOR PRUEBA", BOLETA, self.issuer
        )
        vale = self.make_sequence(
            "TST-VC", "V.COND", "V.COND - ALFA", RECIBO, self.issuer
        )

        self.assertEqual(series_label(cobrador), "Serie COBRADOR PRUEBA - ZZ1")
        self.assertEqual(series_label(vale), "Serie ALFA - V.COND")

    def test_without_brand_in_the_label_uses_the_issuer(self):
        suelto = self.make_sequence(
            "TST-R904", "R904", "R904", RECIBO, self.issuer
        )

        # Sin la forma societaria: la hoja nombra la marca.
        self.assertEqual(series_label(suelto), "Serie EMPRESA ALFA - 904")


class ConsolidatedTotalsTests(CashClosingTestCase):
    def test_sales_are_grouped_by_document_type_and_by_talonario(self):
        self.emit("100.00", self.factura)
        self.emit("50.50", self.factura)
        self.emit("20.00", self.boleta)
        self.emit("30.00", self.recibo_gamma)
        self.emit("45.00", self.recibo_delta)

        report = self.build()
        rows = self.amounts(report)

        self.assertEqual(report["sales_total"], Decimal("245.50"))
        self.assertEqual(rows["Total de Ventas:"], Decimal("245.50"))
        self.assertEqual(rows[". Facturas:"], Decimal("150.50"))
        self.assertEqual(rows[".. Serie ALFA - 901 :"], Decimal("150.50"))
        self.assertEqual(rows[". Boletas:"], Decimal("20.00"))
        self.assertEqual(rows[".. Serie BETA - 902 :"], Decimal("20.00"))
        self.assertEqual(rows[". Recibos de servicios públicos:"], Decimal("75.00"))
        self.assertEqual(report["operations"], 5)

    def test_same_printed_series_of_two_brands_are_two_lines(self):
        self.emit("30.00", self.recibo_gamma)
        self.emit("45.00", self.recibo_delta)

        rows = self.amounts(self.build())

        self.assertEqual(rows[".. Serie GAMMA - 903 :"], Decimal("30.00"))
        self.assertEqual(rows[".. Serie DELTA - 903 :"], Decimal("45.00"))

    def test_the_sheet_keeps_the_sicav_rows_in_order(self):
        self.emit("100.00", self.factura)

        texts = [line["text"] for line in self.build()["lines"]]

        self.assertEqual(
            texts,
            [
                "SALDO ANTERIOR",
                "Total:",
                "INGRESOS",
                "Total de Ventas:",
                ". Facturas:",
                ".. Serie ALFA - 901 :",
                ". Boletas:",
                ". Recibos de servicios públicos:",
                "Garantias:",
                "Otros:",
                "... TOTAL:",
                "EGRESOS",
                "(-)Garantias:",
                "(-)Deposito MN:",
                "(-)Gastos RECIBO:",
                "... TOTAL:",
                "SALDO EN CAJA:",
                "",
                "COMPOSICION",
                "Billetes de 200",
                "Billetes de 100",
                "Billetes de 50",
                "Billetes de 20",
                "Billetes de 10",
                "Monedas",
                "Cheques",
                "Vales Personal",
                "Otros",
                "... TOTAL",
            ],
        )

    def test_the_composition_is_zero_while_sicv_has_no_cash_count(self):
        self.emit("80.00", self.boleta)

        report = self.build()

        self.assertTrue(all(item["total"] == 0 for item in report["composition"]))
        self.assertEqual(report["composition_total"], 0)

    def test_rows_sicv_does_not_register_yet_are_zero(self):
        self.emit("80.00", self.boleta)

        report = self.build()

        self.assertEqual(report["opening_balance"], Decimal("0"))
        self.assertEqual(report["guarantees"], Decimal("0"))
        self.assertEqual(report["other_income"], Decimal("0"))
        self.assertEqual(report["expense_total"], Decimal("0"))
        self.assertTrue(all(e["total"] == 0 for e in report["expenses"]))
        self.assertEqual(report["income_total"], Decimal("80.00"))
        self.assertEqual(report["cash_balance"], Decimal("80.00"))

    def test_an_empty_period_still_prints_the_three_groups_in_zero(self):
        report = self.build()
        rows = self.amounts(report)

        self.assertEqual(rows[". Facturas:"], 0)
        self.assertEqual(rows[". Boletas:"], 0)
        self.assertEqual(rows[". Recibos de servicios públicos:"], 0)
        self.assertEqual(report["cash_balance"], 0)
        self.assertEqual(report["rows"], [])


class ListedSequencesTests(CashClosingTestCase):
    """Cada talonario de la sede figura en la hoja aunque no haya emitido."""

    def offer(self, sequence, office):
        OfficeSequence.objects.create(office=office, sequence=sequence)

    def test_a_talonario_of_the_branch_without_sales_is_listed_in_zero(self):
        self.offer(self.recibo_gamma, self.office)

        rows = self.amounts(self.build())

        self.assertEqual(rows[".. Serie GAMMA - 903 :"], 0)

    def test_talonarios_of_another_branch_are_not_listed(self):
        self.offer(self.recibo_gamma, self.other_office)

        rows = self.amounts(self.build())

        self.assertNotIn(".. Serie GAMMA - 903 :", rows)

    def test_retired_talonarios_and_closed_offices_are_not_listed(self):
        self.recibo_gamma.is_active = False
        self.recibo_gamma.save()
        self.offer(self.recibo_gamma, self.office)

        self.second_office.is_active = False
        self.second_office.save()
        self.offer(self.recibo_delta, self.second_office)

        rows = self.amounts(self.build())

        self.assertNotIn(".. Serie GAMMA - 903 :", rows)
        self.assertNotIn(".. Serie DELTA - 903 :", rows)

    def test_without_consolidating_only_the_office_talonarios_are_listed(self):
        self.offer(self.recibo_gamma, self.office)
        self.offer(self.recibo_delta, self.second_office)

        rows = self.amounts(self.build(office=self.office))

        self.assertIn(".. Serie GAMMA - 903 :", rows)
        self.assertNotIn(".. Serie DELTA - 903 :", rows)

    def test_the_issuer_filter_also_trims_the_listed_talonarios(self):
        self.offer(self.recibo_gamma, self.office)
        self.offer(self.recibo_delta, self.office)

        rows = self.amounts(self.build(issuer=self.issuer))

        self.assertIn(".. Serie GAMMA - 903 :", rows)
        self.assertNotIn(".. Serie DELTA - 903 :", rows)

    def test_a_collector_block_is_listed_only_when_it_sells(self):
        """El block de un cobrador se ofrece en todas las oficinas y la hoja
        de SICAV no lo trae en cero: figura en cuanto emite algo."""
        cobrador = self.make_sequence(
            "TST-ZZ1", "ZZ1", "B: ZZ1 - COBRADOR PRUEBA", BOLETA, self.issuer
        )
        cobrador.autonumber = False
        cobrador.save()
        self.offer(cobrador, self.office)

        self.assertNotIn(
            ".. Serie COBRADOR PRUEBA - ZZ1 :", self.amounts(self.build())
        )

        self.emit("15.00", cobrador)

        rows = self.amounts(self.build())

        self.assertEqual(rows[".. Serie COBRADOR PRUEBA - ZZ1 :"], Decimal("15.00"))
        self.assertEqual(rows[". Boletas:"], Decimal("15.00"))


class SicavSeriesTests(CashClosingTestCase):
    """Las series de la hoja son las de SICAV de cada sede, en su orden."""

    def with_series(self, rows):
        return mock.patch.dict(cash_closing.SICAV_SERIES, {"SEDA": rows})

    def series_lines(self, report):
        return [
            line["text"] for line in report["lines"] if line["style"] == "series"
        ]

    def test_jauja_lists_exactly_the_series_of_its_sicav_sheet(self):
        jauja = Branch.objects.get(code="JAUJA")

        report = build_cash_closing(
            branch=jauja, date_from=date(2026, 9, 1), date_to=date(2026, 9, 29)
        )

        self.assertEqual(
            self.series_lines(report),
            [
                ".. Serie CABLE LOS ANDES - 001 :",
                ".. Serie INVERSIONES - 002 :",
                ".. Serie INVERSIONES - 003 :",
                ".. Serie CABLE LOS ANDES - 001 :",
                ".. Serie INVERSIONES - 002 :",
                ".. Serie CABLE LOS ANDES - 003 :",
                ".. Serie INVERSIONES - 003 :",
                ".. Serie INVERSIONES - 004 :",
                ".. Serie CABLE LOS ANDES - 004 :",
                ".. Serie RED OPTICA - 011 :",
                ".. Serie SPEEDY - 011 :",
                ".. Serie RED OPTICA - 012 :",
                ".. Serie SPEEDY - 012 :",
                ".. Serie VELOCIDAD - 012 :",
            ],
        )

    def test_the_series_keep_the_sicav_order_and_show_in_zero(self):
        rows = [
            (RECIBO, "GAMMA", "903", "TSTA"),
            (FACTURA, "ALFA", "901", "TSTA"),
            (RECIBO, "DELTA", "903", "TSTB"),
        ]

        with self.with_series(rows):
            report = self.build()

        self.assertEqual(
            self.series_lines(report),
            [
                ".. Serie ALFA - 901 :",
                ".. Serie GAMMA - 903 :",
                ".. Serie DELTA - 903 :",
            ],
        )
        self.assertTrue(
            all(line["amount"] == 0 for line in report["lines"] if line["style"] == "series")
        )

    def test_a_sale_adds_up_in_its_sicav_row(self):
        with self.with_series([(FACTURA, "ALFA", "901", "TSTA")]):
            self.emit("100.00", self.factura)
            rows = self.amounts(self.build())

        self.assertEqual(rows[".. Serie ALFA - 901 :"], Decimal("100.00"))

    def test_a_talonario_outside_the_list_that_sells_goes_last_in_its_group(self):
        with self.with_series([
            (RECIBO, "GAMMA", "903", "TSTA"),
            (RECIBO, "OMEGA", "999", "TSTB"),
        ]):
            self.emit("45.00", self.recibo_delta)
            report = self.build()

        self.assertEqual(
            self.series_lines(report),
            [
                ".. Serie GAMMA - 903 :",
                ".. Serie OMEGA - 999 :",
                ".. Serie DELTA - 903 :",
            ],
        )
        self.assertEqual(report["sales_total"], Decimal("45.00"))

    def test_the_listed_series_ignore_the_padron(self):
        """Con lista de SICAV, un talonario del padrón que no vende no se
        agrega en cero."""
        OfficeSequence.objects.create(office=self.office, sequence=self.boleta)

        with self.with_series([(FACTURA, "ALFA", "901", "TSTA")]):
            report = self.build()

        self.assertEqual(self.series_lines(report), [".. Serie ALFA - 901 :"])

    def test_the_issuer_filter_trims_the_sicav_rows(self):
        rows = [
            (FACTURA, "ALFA", "901", "TSTA"),
            (BOLETA, "BETA", "902", "TSTB"),
        ]

        with self.with_series(rows):
            report = self.build(issuer=self.other_issuer)

        self.assertEqual(self.series_lines(report), [".. Serie BETA - 902 :"])

    def test_the_series_filter_leaves_only_its_row(self):
        rows = [
            (RECIBO, "GAMMA", "903", "TSTA"),
            (RECIBO, "DELTA", "903", "TSTB"),
        ]

        with self.with_series(rows):
            report = self.build(sequence=self.recibo_delta)

        self.assertEqual(self.series_lines(report), [".. Serie DELTA - 903 :"])


class PendingAndVoidedTests(CashClosingTestCase):
    def test_pending_and_voided_do_not_inflate_the_sales(self):
        self.emit("100.00", self.factura)
        self.emit("40.00", self.factura, status=Payment.Status.PENDING)
        self.emit("25.00", self.boleta, status=Payment.Status.VOIDED)

        report = self.build()

        self.assertEqual(report["sales_total"], Decimal("100.00"))
        self.assertEqual(report["cash_balance"], Decimal("100.00"))
        self.assertEqual(report["operations"], 1)
        self.assertEqual(
            report["pending"], {"operations": 1, "total": Decimal("40.00")}
        )
        self.assertEqual(
            report["voided"], {"operations": 1, "total": Decimal("25.00")}
        )

    def test_the_detail_lists_every_receipt_with_its_status(self):
        self.emit("100.00", self.factura)
        self.emit("40.00", self.factura, status=Payment.Status.PENDING)
        self.emit("25.00", self.boleta, status=Payment.Status.VOIDED)

        statuses = [row["status_code"] for row in self.build()["rows"]]

        self.assertEqual(
            sorted(statuses),
            sorted([
                Payment.Status.REGISTERED,
                Payment.Status.PENDING,
                Payment.Status.VOIDED,
            ]),
        )


class PeriodAndScopeTests(CashClosingTestCase):
    def test_the_period_is_cut_by_issue_date_both_ends_included(self):
        self.emit("1.00", self.boleta, issued_at=moment(9, 23, 59))
        self.emit("10.00", self.boleta, issued_at=moment(10, 0, 5))
        self.emit("20.00", self.boleta, issued_at=moment(12, 23, 55))
        self.emit("2.00", self.boleta, issued_at=moment(13, 0, 1))

        report = self.build(date_from=date(2026, 9, 10), date_to=date(2026, 9, 12))

        self.assertEqual(report["sales_total"], Decimal("30.00"))

    def test_the_issue_date_rules_over_the_payment_date(self):
        receipt = self.emit("15.00", self.boleta, issued_at=moment(10))
        Payment.objects.filter(pk=receipt.payment_id).update(paid_at=moment(20))

        report = self.build(date_from=date(2026, 9, 10), date_to=date(2026, 9, 10))

        self.assertEqual(report["sales_total"], Decimal("15.00"))

    def test_another_branch_never_enters(self):
        self.emit("100.00", self.factura)
        self.emit(
            "999.00",
            self.factura,
            branch=self.other_branch,
            office=self.other_office,
            user=self.other_cashier,
        )

        report = self.build()

        self.assertEqual(report["sales_total"], Decimal("100.00"))
        self.assertEqual(len(report["rows"]), 1)

    def test_without_consolidating_only_the_office_counts(self):
        self.emit("100.00", self.factura, office=self.office)
        self.emit("70.00", self.factura, office=self.second_office)

        report = self.build(office=self.second_office)

        self.assertEqual(report["sales_total"], Decimal("70.00"))
        self.assertEqual(report["place_label"], str(self.second_office))

    def test_filter_by_issuer(self):
        self.emit("100.00", self.factura)
        self.emit("20.00", self.boleta)
        self.emit("45.00", self.recibo_delta)

        report = self.build(issuer=self.other_issuer)

        self.assertEqual(report["sales_total"], Decimal("65.00"))

    def test_filter_by_series(self):
        self.emit("30.00", self.recibo_gamma)
        self.emit("45.00", self.recibo_delta)

        report = self.build(sequence=self.recibo_gamma)

        self.assertEqual(report["sales_total"], Decimal("30.00"))

    def test_filter_by_user_who_registered(self):
        self.emit("100.00", self.factura, user=self.cashier)
        self.emit("60.00", self.factura, user=self.second_cashier)

        report = self.build(user=self.second_cashier)

        self.assertEqual(report["sales_total"], Decimal("60.00"))


class CashClosingPermissionTests(CashClosingTestCase):
    def test_accounting_role_brings_the_permission(self):
        self.assertTrue(self.accountant.has_perm("payments.view_cash_closing"))

    def test_atc_and_admin_do_not(self):
        admin = User.objects.create_user(
            username="admin_prueba", role=User.Role.ADMIN, branch=self.branch
        )

        self.assertFalse(self.cashier.has_perm("payments.view_cash_closing"))
        self.assertFalse(admin.has_perm("payments.view_cash_closing"))

    def test_accounting_still_does_not_see_payments_or_register_them(self):
        """El rol solo trae el cierre. Las pruebas de cobranza cuentan con
        que Contabilidad no herede los permisos de pagos."""
        for permission in (
            "payments.view_payment",
            "payments.add_payment",
            "payments.view_receipt",
            "payments.void_payment",
        ):
            with self.subTest(permission=permission):
                self.assertFalse(self.accountant.has_perm(permission))


class CashClosingWebTestCase(CashClosingTestCase):
    def login(self, user, office=None):
        self.client.force_login(user)

        session = self.client.session
        session[ACTIVE_BRANCH_SESSION_KEY] = self.branch.pk

        if office is not None:
            session[ACTIVE_OFFICE_SESSION_KEY] = office.pk

        session.save()

    def page(self, **extra):
        data = {
            "date_from": "2026-09-01",
            "date_to": "2026-09-30",
            "report_type": "EMISSION",
            "all_offices": "on",
        }
        data.update(extra)

        return self.client.get(reverse("reports:cash_closing"), data)

    def body(self, response):
        return b"".join(response.streaming_content)


class CashClosingAccessTests(CashClosingWebTestCase):
    def test_anonymous_is_sent_to_login(self):
        response = self.client.get(reverse("reports:cash_closing"))

        self.assertEqual(response.status_code, 302)
        self.assertIn("/accounts/login/", response["Location"])

    def test_atc_is_rejected_from_the_screen_and_the_exports(self):
        self.login(self.cashier)

        self.assertEqual(self.page().status_code, 403)
        self.assertEqual(self.page(export_format="EXCEL").status_code, 403)

    def test_accounting_sees_the_filters_of_sicav(self):
        self.login(self.accountant)

        response = self.client.get(reverse("reports:cash_closing"))

        self.assertEqual(response.status_code, 200)

        for label in (
            "Cierre de caja",
            "Empresa",
            "Desde",
            "Hasta",
            "Consolidado de emisión",
            "Usuario",
            "Serie",
            "Formato",
            "Consolidado oficinas",
            "Exportar",
        ):
            with self.subTest(label=label):
                self.assertContains(response, label)

    def test_the_screen_opens_with_the_month_so_far(self):
        """Sin filtros en la dirección el formulario abre con el mes en curso."""
        self.login(self.accountant)

        response = self.client.get(reverse("reports:cash_closing"))
        form = response.context["form"]
        today = timezone.localdate()

        self.assertEqual(form.cleaned_data["date_from"], today.replace(day=1))
        self.assertEqual(form.cleaned_data["date_to"], today)
        self.assertTrue(form.cleaned_data["all_offices"])

    def test_the_user_list_only_offers_who_registered_in_the_branch(self):
        self.emit("10.00", self.boleta, user=self.cashier)
        self.emit(
            "10.00",
            self.boleta,
            user=self.other_cashier,
            branch=self.other_branch,
            office=self.other_office,
        )
        self.login(self.accountant)

        response = self.client.get(reverse("reports:cash_closing"))
        users = list(response.context["form"].fields["user"].queryset)

        self.assertEqual(users, [self.cashier])

    def test_the_menu_offers_the_report_to_accounting(self):
        self.login(self.accountant)

        response = self.client.get(reverse("reports:cash_closing"))

        self.assertContains(response, reverse("reports:cash_closing"))
        self.assertContains(response, "Cierre de caja</a>")


class CashClosingScreenTests(CashClosingWebTestCase):
    def setUp(self):
        super().setUp()
        self.emit("100.00", self.factura)
        self.emit("20.00", self.boleta, office=self.second_office)
        self.emit("40.00", self.boleta, status=Payment.Status.PENDING)
        self.emit(
            "999.00",
            self.factura,
            branch=self.other_branch,
            office=self.other_office,
            user=self.other_cashier,
        )
        self.login(self.accountant, office=self.office)

    def consolidated(self, response):
        """Etiqueta -> monto de la pestaña Consolidado del Excel."""
        sheet = load_workbook(BytesIO(self.body(response)))["Consolidado"]

        return {
            sheet.cell(row=row, column=1).value: sheet.cell(row=row, column=2).value
            for row in range(1, sheet.max_row + 1)
        }

    def test_the_screen_shows_only_the_filters(self):
        """Como en SICAV: la pantalla pide los datos y el reporte sale en
        archivo. Ni el consolidado ni el detalle se pintan aquí."""
        response = self.page()

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'id="reportForm"')
        self.assertNotContains(response, "Consolidado de emisión</h2>")
        self.assertNotContains(response, "Saldo en caja")
        self.assertNotContains(response, "S/ 120,00")
        self.assertNotContains(response, "CLI-T001")

    def test_the_form_looks_like_the_materials_report(self):
        """Mismo trazo que materiales: formulario en filas etiqueta / campo,
        «Formato» y un solo botón «Exportar» al pie."""
        response = self.page()

        self.assertContains(response, 'class="tc-form tc-report-form"')
        self.assertContains(response, 'class="tc-card-footer"')
        self.assertContains(response, "Exportar")
        self.assertContains(response, 'method="get"')

    def test_the_format_offers_excel_and_pdf_and_opens_in_excel(self):
        response = self.page()
        form = response.context["form"]

        self.assertEqual(
            [value for value, _label in form.fields["export_format"].choices],
            ["EXCEL", "PDF"],
        )
        self.assertEqual(form["export_format"].value(), "EXCEL")

    def test_pdf_opens_in_a_new_tab_and_excel_stays(self):
        """La misma regla que materiales: la lista viene del servidor."""
        response = self.page()

        self.assertEqual(response.context["viewed_formats"], ("HTML", "PDF"))
        self.assertContains(response, '"PDF"')

    def test_unchecking_all_offices_exports_only_the_active_office(self):
        response = self.page(all_offices="", export_format="EXCEL")

        values = self.consolidated(response)

        self.assertEqual(values["Total de Ventas:"], 100.0)

    def test_an_inverted_period_returns_to_the_filters_with_its_error(self):
        response = self.page(
            date_from="2026-09-30", date_to="2026-09-01", export_format="EXCEL"
        )

        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "reports/cash_closing.html")
        self.assertContains(
            response, "La fecha inicial no puede ser posterior a la final."
        )

    def test_excel_keeps_the_sicav_sheet_and_adds_the_detail(self):
        response = self.page(export_format="EXCEL")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn(
            "cierre_caja_seda_20260901_20260930.xlsx",
            response["Content-Disposition"],
        )

        book = load_workbook(BytesIO(self.body(response)))
        self.assertEqual(book.sheetnames, ["Consolidado", "Detalle"])

        sheet = book["Consolidado"]
        self.assertEqual(sheet["A2"].value, "CONTROL ADMINISTRATIVO EMISIÓN")
        self.assertEqual(sheet["A3"].value, "Desde 01/09/2026 Hasta 30/09/2026")
        self.assertEqual(sheet["B1"].value, "Sede Alfa")

        values = {
            sheet.cell(row=row, column=1).value: sheet.cell(row=row, column=2).value
            for row in range(1, sheet.max_row + 1)
        }
        self.assertEqual(values["Total de Ventas:"], 120.0)
        self.assertEqual(values["SALDO EN CAJA:"], 120.0)
        self.assertEqual(values["COMPOSICION"], None)
        self.assertEqual(values["Billetes de 200"], 0.0)
        self.assertEqual(values["... TOTAL"], 0.0)
        self.assertRegex(
            sheet["A1"].value,
            r"^contable_prueba \d{2}/\d{2}/\d{4} \d{2}:\d{2}:\d{2}$",
        )
        # Nada que la hoja de SICAV no tenga.
        self.assertNotIn("INFORMATIVO (no suma en caja)", values)
        self.assertNotIn(999.0, values.values())

    def test_the_excel_detail_adds_up_to_the_consolidated(self):
        """La pestaña Detalle suma lo mismo que el Total de Ventas."""
        book = load_workbook(BytesIO(self.body(self.page(export_format="EXCEL"))))
        detail = book["Detalle"]
        sheet = book["Consolidado"]

        headers = [cell.value for cell in detail[1]]
        amount_col = headers.index("Monto") + 1
        status_col = headers.index("Estado") + 1
        rows = [
            row for row in range(2, detail.max_row + 1)
            if detail.cell(row=row, column=status_col).value
        ]

        cancelled = sum(
            detail.cell(row=row, column=amount_col).value
            for row in rows
            if detail.cell(row=row, column=status_col).value == "Cancelado"
        )
        sales = next(
            sheet.cell(row=row, column=2).value
            for row in range(1, sheet.max_row + 1)
            if sheet.cell(row=row, column=1).value == "Total de Ventas:"
        )

        self.assertEqual(len(rows), 3)
        self.assertEqual(cancelled, sales)

    def test_pdf_is_a_real_pdf_with_the_consolidated(self):
        response = self.page(export_format="PDF")

        data = self.body(response)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(data.startswith(b"%PDF"))

        text = pdf_text(data)
        self.assertIn("CONTROL ADMINISTRATIVO EMISIÓN", text)
        self.assertIn("SALDO EN CAJA:", text)
        self.assertIn("120.00", text)
        self.assertIn("COMPOSICION", text)

    def test_an_unknown_format_returns_to_the_form_with_its_error(self):
        response = self.page(export_format="WORD")

        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "reports/cash_closing.html")
        self.assertTrue(response.context["form"].errors["export_format"])
