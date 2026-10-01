"""«Ingresos por usuario»: el reporte de ingresos de SICAV, una fila por
concepto cobrado y el total al pie.

Todo el escenario es sintético, como el del cierre de caja: empresas,
talonarios, abonados, direcciones y montos inventados para la prueba.
"""

from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from apps.customers.models import Customer, CustomerAddress
from apps.payments.models import Charge, Payment, PaymentAllocation
from apps.reports.cash_closing import build_cash_closing
from apps.reports.user_income import (
    build_user_income,
    customer_name,
    document_label,
)

from . import test_cash_closing as base
from .pdf_text import pdf_text


XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class UserIncomeTestCase(base.CashClosingWebTestCase):
    def setUp(self):
        super().setUp()
        CustomerAddress.objects.create(
            customer=self.customer,
            address="Calle Prueba 123",
            district="Distrito Prueba",
            is_primary=True,
        )

    def charge(self, description, amount, *, period_end=None):
        return Charge.objects.create(
            customer=self.customer,
            concept=Charge.Concept.OTHER,
            description=description,
            amount=Decimal(amount),
            due_date=date(2026, 9, 30),
            period_end=period_end,
        )

    def collect(self, sequence, *charges, extra="0.00", **kwargs):
        """Un cobro que cubre esos cargos enteros, más `extra` a favor."""
        total = sum((charge.amount for charge in charges), Decimal(extra))
        receipt = self.emit(str(total), sequence, **kwargs)

        for charge in charges:
            PaymentAllocation.objects.create(
                payment=receipt.payment, charge=charge, amount=charge.amount
            )

        return receipt

    def income(self, **kwargs):
        options = {
            "branch": self.branch,
            "date_from": date(2026, 9, 1),
            "date_to": date(2026, 9, 30),
        }
        options.update(kwargs)

        return build_user_income(**options)


class UserIncomeRowsTests(UserIncomeTestCase):
    def test_one_row_per_charge_the_receipt_covered(self):
        plan = self.charge("PLAN PRUEBA 100MG", "70.00", period_end=date(2026, 9, 30))
        anexo = self.charge("ANEXO", "5.00")
        self.collect(self.factura, plan, anexo)

        rows = self.income()["rows"]

        self.assertEqual(
            [(row["detail"], row["amount"], row["paid_until"]) for row in rows],
            [
                ("PLAN PRUEBA 100MG", Decimal("70.00"), date(2026, 9, 30)),
                ("ANEXO", Decimal("5.00"), None),
            ],
        )

        first = rows[0]
        self.assertEqual(first["customer_code"], "CLI-T001")
        self.assertEqual(first["customer_name"], "Prueba Uno, Abonado")
        self.assertEqual(first["address"], "Calle Prueba 123")
        self.assertEqual(first["issued_on"], date(2026, 9, 10))
        self.assertEqual(first["paid_on"], date(2026, 9, 10))
        self.assertEqual(first["document"], "F:F901 0000001")
        self.assertEqual(first["user"], "cajero_prueba")

    def test_what_was_not_applied_goes_as_credit(self):
        """El total tiene que ser el dinero que entró: lo que no cubrió ningún
        cargo va en su propia fila."""
        self.collect(self.factura, self.charge("ANEXO", "5.00"), extra="15.00")
        self.emit("30.00", self.boleta)

        rows = self.income()["rows"]

        self.assertEqual(
            [(row["detail"], row["amount"]) for row in rows],
            [
                ("Saldo a favor", Decimal("30.00")),
                ("ANEXO", Decimal("5.00")),
                ("Saldo a favor", Decimal("15.00")),
            ],
        )

    def test_only_cancelled_payments_count(self):
        self.collect(self.factura, self.charge("ANEXO", "5.00"))
        self.emit("40.00", self.factura, status=Payment.Status.PENDING)
        self.emit("50.00", self.factura, status=Payment.Status.VOIDED)

        report = self.income()

        self.assertEqual(len(report["rows"]), 1)
        self.assertEqual(report["total"], Decimal("5.00"))

    def test_the_total_is_the_sales_total_of_the_consolidated(self):
        """Mismas reglas que el consolidado: con los mismos filtros, los dos
        reportes dicen la misma cifra."""
        self.collect(self.factura, self.charge("PLAN", "70.00"), extra="10.00")
        self.collect(
            self.boleta, self.charge("ANEXO", "5.00"), office=self.second_office
        )
        self.emit("20.00", self.recibo_gamma)
        self.emit("40.00", self.boleta, status=Payment.Status.PENDING)

        for filters in ({}, {"issuer": self.issuer}, {"office": self.office}):
            with self.subTest(filters=sorted(filters)):
                consolidated = build_cash_closing(
                    branch=self.branch,
                    date_from=date(2026, 9, 1),
                    date_to=date(2026, 9, 30),
                    **filters,
                )

                self.assertEqual(
                    self.income(**filters)["total"], consolidated["sales_total"]
                )

    def test_without_user_everyone_and_with_user_only_theirs(self):
        self.collect(self.factura, self.charge("ANEXO", "5.00"))
        self.collect(
            self.factura, self.charge("PLAN", "70.00"), user=self.second_cashier
        )

        everyone = self.income()
        theirs = self.income(user=self.second_cashier)

        self.assertEqual(everyone["users_label"], "Usuarios")
        self.assertEqual(
            {row["user"] for row in everyone["rows"]},
            {"cajero_prueba", "cajero_dos"},
        )
        self.assertEqual(theirs["users_label"], "cajero_dos")
        self.assertEqual([row["detail"] for row in theirs["rows"]], ["PLAN"])

    def test_rows_follow_the_document_and_its_number(self):
        """Como la hoja de SICAV: por talonario y número, no por fecha."""
        self.collect(self.factura, self.charge("UNO", "1.00"), issued_at=base.moment(5))
        self.collect(self.factura, self.charge("DOS", "2.00"), issued_at=base.moment(3))
        self.collect(self.boleta, self.charge("TRES", "3.00"), issued_at=base.moment(20))

        rows = self.income()["rows"]

        self.assertEqual(
            [row["document"] for row in rows],
            ["B:B902 0000001", "F:F901 0000001", "F:F901 0000002"],
        )


class UserIncomeLabelTests(UserIncomeTestCase):
    def test_a_company_goes_by_its_business_name(self):
        empresa = Customer(
            person_type=Customer.PersonType.LEGAL,
            business_name="EMPRESA CLIENTE S.A.C.",
            document_number="20000000099",
        )

        self.assertEqual(customer_name(empresa), "EMPRESA CLIENTE S.A.C.")

    def test_a_series_without_letter_in_its_label_goes_alone(self):
        vale = self.make_sequence(
            "TST-VC", "V.COND", "V.COND - ALFA", base.RECIBO, self.issuer
        )

        self.assertEqual(document_label(self.emit("5.00", vale)), "V.COND 0000001")


class UserIncomeExportTests(UserIncomeTestCase):
    def setUp(self):
        super().setUp()
        self.login(self.accountant)
        self.collect(
            self.factura,
            self.charge("PLAN PRUEBA", "70.00", period_end=date(2026, 9, 30)),
            self.charge("ANEXO", "5.00"),
        )
        self.collect(
            self.factura, self.charge("RECONEXION", "10.00"), user=self.second_cashier
        )

    def export(self, **extra):
        data = {"report_type": "INCOME_BY_USER", "export_format": "EXCEL"}
        data.update(extra)

        return self.page(**data)

    def sheet(self, response):
        return load_workbook(BytesIO(self.body(response)))["Ingresos"]

    def test_excel_has_the_sicav_columns_and_the_total_at_the_bottom(self):
        response = self.export()

        self.assertEqual(response["Content-Type"], XLSX)
        self.assertIn(
            "ingresos_usuario_seda_20260901_20260930.xlsx",
            response["Content-Disposition"],
        )

        sheet = self.sheet(response)

        self.assertRegex(sheet["A1"].value, r"^\d{2}/\d{2}/\d{4} \d{2}:\d{2}:\d{2}$")
        self.assertEqual(
            sheet["B1"].value,
            "Reporte de ingresos: 01/09/2026 - 30/09/2026 Usuarios",
        )
        self.assertEqual(sheet["J1"].value, "Sede Alfa")
        self.assertEqual(
            [cell.value for cell in sheet[3]],
            [
                "Código", "Fecha", "Abonado", "Dirección", "Fecha pago",
                "Detalle", "Pagó hasta", "Monto", "Documento", "Usuario",
            ],
        )
        self.assertEqual(
            [sheet.cell(row=row, column=6).value for row in (4, 5, 6)],
            ["PLAN PRUEBA", "ANEXO", "RECONEXION"],
        )
        self.assertEqual(sheet["G4"].value.date(), date(2026, 9, 30))
        self.assertIsNone(sheet["G5"].value)
        self.assertEqual(sheet["H4"].value, 70.0)
        self.assertEqual(sheet["I4"].value, "F:F901 0000001")
        self.assertEqual(sheet["J6"].value, "cajero_dos")

        # El total al pie, bajo «Monto», y nada después.
        self.assertEqual(sheet["G7"].value, "TOTAL")
        self.assertEqual(sheet["H7"].value, 85.0)
        self.assertEqual(sheet.max_row, 7)

    def test_with_a_user_only_theirs_named_in_the_title_and_the_file(self):
        response = self.export(user=self.second_cashier.pk)

        self.assertIn(
            "ingresos_usuario_seda_20260901_20260930_cajero_dos.xlsx",
            response["Content-Disposition"],
        )

        sheet = self.sheet(response)

        self.assertEqual(
            sheet["B1"].value,
            "Reporte de ingresos: 01/09/2026 - 30/09/2026 cajero_dos",
        )
        self.assertEqual(sheet["F4"].value, "RECONEXION")
        self.assertEqual(sheet["G5"].value, "TOTAL")
        self.assertEqual(sheet["H5"].value, 10.0)

    def test_excel_forces_formula_like_text_to_text(self):
        dangerous = '=HYPERLINK("https://example.invalid","x")'
        CustomerAddress.objects.filter(customer=self.customer).update(address=dangerous)

        cell = self.sheet(self.export())["D4"]

        self.assertEqual(cell.value, "'" + dangerous)
        self.assertEqual(cell.data_type, "s")

    def test_pdf_has_the_columns_the_rows_and_the_total(self):
        response = self.export(export_format="PDF")
        data = self.body(response)

        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertIn(
            "ingresos_usuario_seda_20260901_20260930.pdf",
            response["Content-Disposition"],
        )
        self.assertTrue(data.startswith(b"%PDF"))

        text = pdf_text(data)

        for expected in (
            "Reporte de ingresos: 01/09/2026 - 30/09/2026 Usuarios",
            "Sede Alfa",
            "Código",
            "Pagó hasta",
            "Documento",
            "Prueba Uno, Abonado",
            "PLAN PRUEBA",
            "F:F901 0000001",
            "cajero_dos",
            "TOTAL",
            "85.00",
        ):
            with self.subTest(expected=expected):
                self.assertIn(expected, text)
