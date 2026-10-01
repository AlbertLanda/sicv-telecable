"""Las salidas de «Ingresos por usuario»: Excel y PDF.

Leen las filas que arma `user_income.build_user_income()` y solo deciden cómo
se dibujan. El total también llega armado: si cada salida sumara por su
cuenta, el Excel y el PDF del mismo periodo podrían decir cifras distintas.

Las dos tienen la forma de la hoja de SICAV: la hora en la esquina, el título
con el periodo al centro, la sede a la derecha, las columnas en su orden y el
total al pie, bajo «Monto».
"""

from io import BytesIO
from xml.sax.saxutils import escape

from django.utils import timezone
from django.utils.text import slugify

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .cash_closing_exporters import (
    AMOUNT_FORMAT,
    CONTENT_TYPES,
    DATE_FORMAT,
    STAMP_FORMAT,
    amount_text,
    safe_excel_text,
)


TOTAL_LABEL = "TOTAL"

# Fecha de verdad y no texto en el Excel, para que Contabilidad pueda ordenar
# y filtrar por día; se ve igual que en la hoja de SICAV.
EXCEL_DATE_FORMAT = "DD/MM/YYYY"


def title_text(report):
    """«Reporte de ingresos: 01/09/2026 - 30/09/2026 Usuarios»."""
    desde = report["date_from"].strftime(DATE_FORMAT)
    hasta = report["date_to"].strftime(DATE_FORMAT)

    return f"Reporte de ingresos: {desde} - {hasta} {report['users_label']}"


def stamp():
    """La esquina de la hoja: solo la hora en que se sacó, como en SICAV."""
    return timezone.localtime().strftime(STAMP_FORMAT)


def cell_text(row, column):
    """Una celda escrita como texto, para el PDF."""
    value = row.get(column["key"])

    if value in (None, ""):
        return ""

    if column.get("date"):
        return value.strftime(DATE_FORMAT)

    if column.get("numeric"):
        return amount_text(value)

    return str(value)


def _amount_index(columns):
    """La posición de «Monto»: el total va debajo y su etiqueta a la izquierda."""
    return next(
        index for index, column in enumerate(columns) if column.get("numeric")
    )


def filename(report, extension):
    """Lleva la sede, el periodo y, si lo hay, el usuario: sacar el reporte
    de cada cajero no debe dar archivos con el mismo nombre."""
    desde = report["date_from"].strftime("%Y%m%d")
    hasta = report["date_to"].strftime("%Y%m%d")
    sede = report["branch"].code.lower()
    name = f"ingresos_usuario_{sede}_{desde}_{hasta}"

    if report["user"] is not None:
        name = f"{name}_{slugify(report['user'].username)}"

    return f"{name}.{extension}"


# ---------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------

def render_excel(report, buffer):
    book = Workbook()
    sheet = book.active
    sheet.title = "Ingresos"

    columns = report["columns"]
    last = len(columns)
    bold = Font(bold=True, size=10)
    normal = Font(size=10)
    right = Alignment(horizontal="right")
    line = Side(style="thin", color="000000")

    corner = sheet.cell(row=1, column=1, value=stamp())
    corner.font = Font(bold=True, size=8)

    sheet.merge_cells(start_row=1, start_column=2, end_row=1, end_column=last - 1)
    title = sheet.cell(row=1, column=2, value=safe_excel_text(title_text(report)))
    title.font = Font(bold=True, size=12)
    title.alignment = Alignment(horizontal="center")

    place = sheet.cell(row=1, column=last, value=safe_excel_text(report["place_label"]))
    place.font = Font(bold=True, size=11)
    place.alignment = right

    header_row = 3

    for index, column in enumerate(columns, start=1):
        cell = sheet.cell(row=header_row, column=index, value=column["label"])
        cell.font = bold
        cell.border = Border(bottom=line)
        cell.alignment = Alignment(
            horizontal="right" if column.get("numeric") else "left",
            vertical="bottom",
            wrap_text=True,
        )
        sheet.column_dimensions[get_column_letter(index)].width = column["width"]

    row = header_row

    for detail in report["rows"]:
        row += 1

        for index, column in enumerate(columns, start=1):
            value = detail.get(column["key"])

            if column.get("numeric"):
                # Número y no texto: Contabilidad suma y compara estas celdas.
                cell = sheet.cell(row=row, column=index, value=float(value))
                cell.number_format = AMOUNT_FORMAT
                cell.alignment = right
            elif column.get("date"):
                cell = sheet.cell(row=row, column=index, value=value)
                cell.number_format = EXCEL_DATE_FORMAT
            else:
                cell = sheet.cell(row=row, column=index, value=safe_excel_text(value))

            cell.font = normal

    amount_column = _amount_index(columns) + 1
    total_row = row + 1

    label = sheet.cell(row=total_row, column=amount_column - 1, value=TOTAL_LABEL)
    label.font = bold
    label.alignment = right

    total = sheet.cell(row=total_row, column=amount_column, value=float(report["total"]))
    total.font = bold
    total.number_format = AMOUNT_FORMAT
    total.alignment = right
    total.border = Border(top=line)

    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)

    if report["rows"]:
        sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(last)}{row}"

    book.save(buffer)

    return filename(report, "xlsx")


# ---------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------

def render_pdf(report, buffer):
    """La hoja en A4 apaisado: diez columnas no caben de pie. La cabecera de
    la tabla se repite en cada página y el total cierra la última."""
    pagesize = landscape(A4)
    margin = 10 * mm
    document = SimpleDocTemplate(
        buffer,
        pagesize=pagesize,
        leftMargin=margin,
        rightMargin=margin,
        topMargin=margin,
        bottomMargin=margin,
        title=title_text(report),
    )
    width = pagesize[0] - margin * 2

    small = ParagraphStyle("small", fontName="Helvetica-Bold", fontSize=7)
    title = ParagraphStyle(
        "title", fontName="Helvetica-Bold", fontSize=11, alignment=1, leading=14
    )
    place = ParagraphStyle(
        "place", fontName="Helvetica-Bold", fontSize=10, alignment=2
    )
    header = ParagraphStyle(
        "header", fontName="Helvetica-Bold", fontSize=8, leading=9.5
    )
    header_right = ParagraphStyle("header_right", parent=header, alignment=2)
    body = ParagraphStyle("body", fontName="Helvetica", fontSize=7.5, leading=9)
    body_right = ParagraphStyle("body_right", parent=body, alignment=2)
    total_style = ParagraphStyle(
        "total", parent=body_right, fontName="Helvetica-Bold"
    )

    strip = Table(
        [[
            Paragraph(escape(stamp()), small),
            Paragraph(escape(title_text(report)), title),
            Paragraph(escape(report["place_label"]), place),
        ]],
        colWidths=[width * 0.18, width * 0.64, width * 0.18],
    )
    strip.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
    ]))

    columns = report["columns"]
    units = sum(column["width"] for column in columns)
    col_widths = [width * column["width"] / units for column in columns]
    amount_index = _amount_index(columns)

    data = [[
        Paragraph(
            escape(column["label"]),
            header_right if column.get("numeric") else header,
        )
        for column in columns
    ]]

    for row in report["rows"]:
        data.append([
            Paragraph(
                escape(cell_text(row, column)),
                body_right if column.get("numeric") else body,
            )
            for column in columns
        ])

    total_row = [""] * len(columns)
    total_row[amount_index - 1] = Paragraph(TOTAL_LABEL, total_style)
    total_row[amount_index] = Paragraph(amount_text(report["total"]), total_style)
    data.append(total_row)

    last = len(data) - 1
    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("VALIGN", (0, 0), (-1, 0), "BOTTOM"),
        ("LEFTPADDING", (0, 0), (-1, -1), 2),
        ("RIGHTPADDING", (0, 0), (-1, -1), 2),
        ("TOPPADDING", (0, 0), (-1, -1), 1.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5),
        ("LINEBELOW", (0, 0), (-1, 0), 0.8, colors.black),
        ("LINEABOVE", (amount_index, last), (amount_index, last), 0.8, colors.black),
        # El total no empieza página solo: baja con la última fila.
        ("NOSPLIT", (0, last - 1), (-1, last)),
    ]))

    document.build([strip, Spacer(1, 8), table])

    return filename(report, "pdf")


RENDERERS = {
    "EXCEL": render_excel,
    "PDF": render_pdf,
}


def render(report, export_format):
    """Dibuja el reporte en el formato pedido y devuelve (buffer, nombre)."""
    buffer = BytesIO()
    name = RENDERERS[export_format](report, buffer)
    buffer.seek(0)

    return buffer, name
