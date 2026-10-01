"""Las salidas del cierre de caja: Excel y PDF.

Leen las filas que arma `cash_closing.build_cash_closing()` y solo deciden cómo
se dibujan. Ninguna suma por su cuenta: si el Excel pudiera recalcular, el
Excel y el PDF del mismo periodo acabarían diciendo cifras distintas y
Contabilidad no sabría cuál vale.

Las dos salidas son la hoja de SICAV tal cual -cabecera, filas y textos-,
porque Contabilidad la compara renglón por renglón con la suya: nada que la
hoja de SICAV no tenga. El Excel agrega el detalle comprobante por comprobante
en una segunda pestaña, para cuadrar sin tocar la primera.
"""

from io import BytesIO
from xml.sax.saxutils import escape

from django.utils import timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .exporters import MOMENT_FORMAT, TC_BLUE


DATE_FORMAT = "%d/%m/%Y"

# La esquina de la hoja de SICAV lleva la hora con segundos.
STAMP_FORMAT = "%d/%m/%Y %H:%M:%S"

# Montos con dos decimales y separador de miles, que es como se leen en la
# hoja de SICAV y como Contabilidad los copia a sus propios cuadros.
AMOUNT_FORMAT = "#,##0.00"

# Excel interpreta como fórmulas las cadenas que empiezan por ciertos
# caracteres. Los reportes incluyen datos de usuarios/clientes y referencias
# externas, así que toda celda textual se fuerza a texto antes de escribirla.
_EXCEL_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r", "\n")


def safe_excel_text(value):
    """Devuelve texto que Excel no pueda interpretar como fórmula."""
    text = "" if value is None else str(value)

    if text.startswith(_EXCEL_FORMULA_PREFIXES):
        return "'" + text

    return text


def stamp(report):
    """Quién imprimió y cuándo, como la esquina de la hoja de SICAV."""
    moment = timezone.localtime().strftime(STAMP_FORMAT)

    return f"{report['printed_by']} {moment}".strip()


def amount_text(amount):
    """Un monto del PDF escrito como en la hoja de SICAV: «300,939.80»."""
    return f"{amount:,.2f}"


def period_text(report):
    desde = report["date_from"].strftime(DATE_FORMAT)
    hasta = report["date_to"].strftime(DATE_FORMAT)

    return f"Desde {desde} Hasta {hasta}"


def detail_text(row, column):
    """Una celda del detalle, ya escrita como texto.

    Las fechas se escriben aquí y no en cada salida para que la pantalla y el
    Excel no discrepen en cómo ponen el mismo día.
    """
    value = row.get(column["key"])

    if value in (None, ""):
        return ""

    if column["key"] in ("issued_at", "paid_at"):
        return timezone.localtime(value).strftime(MOMENT_FORMAT)

    return str(value)


def filename(report, extension):
    """Lleva la sede y el periodo, para distinguir archivos sin abrirlos."""
    desde = report["date_from"].strftime("%Y%m%d")
    hasta = report["date_to"].strftime("%Y%m%d")
    sede = report["branch"].code.lower()

    return f"cierre_caja_{sede}_{desde}_{hasta}.{extension}"


# ---------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------

def _consolidated_sheet(sheet, report):
    sheet.title = "Consolidado"
    sheet.column_dimensions["A"].width = 48
    sheet.column_dimensions["B"].width = 16

    bold = Font(bold=True, size=10)
    normal = Font(size=10)

    stamp_cell = sheet.cell(row=1, column=1, value=safe_excel_text(stamp(report)))
    stamp_cell.font = Font(bold=True, size=8)

    place = sheet.cell(row=1, column=2, value=safe_excel_text(report["place_label"]))
    place.font = Font(bold=True, size=11)
    place.alignment = Alignment(horizontal="right")

    heading = [
        (report["title"].upper(), Font(bold=True, size=12)),
        (period_text(report), bold),
    ]

    row = 2

    for text, font in heading:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        cell = sheet.cell(row=row, column=1, value=safe_excel_text(text))
        cell.font = font
        cell.alignment = Alignment(horizontal="center")
        row += 1

    for line in report["lines"]:
        text, amount, style = line["text"], line["amount"], line["style"]

        # La fila en blanco que separa el saldo de la composición.
        if style == "blank":
            row += 1
            continue

        label = sheet.cell(row=row, column=1, value=safe_excel_text(text))
        label.font = bold if style in ("section", "total", "balance") else normal

        if amount is not None:
            # Número y no texto: Contabilidad suma y compara estas celdas.
            value = sheet.cell(row=row, column=2, value=float(amount))
            value.number_format = AMOUNT_FORMAT
            value.alignment = Alignment(horizontal="right")
            value.font = bold if style in ("total", "balance") else normal

            if style == "balance":
                side = Side(style="thin", color="000000")
                value.border = Border(left=side, right=side, top=side, bottom=side)

        row += 1


def _detail_sheet(sheet, report):
    sheet.title = "Detalle"

    columns = report["columns"]
    blue = PatternFill("solid", fgColor=TC_BLUE)
    white = Font(bold=True, color="FFFFFF", size=9)
    border = Border(
        left=Side(style="thin", color="D9E4EF"),
        right=Side(style="thin", color="D9E4EF"),
        top=Side(style="thin", color="D9E4EF"),
        bottom=Side(style="thin", color="D9E4EF"),
    )

    for index, column in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index, value=column["label"])
        cell.fill = blue
        cell.font = white
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.column_dimensions[get_column_letter(index)].width = column["width"]

    for offset, detail in enumerate(report["rows"], start=2):
        for index, column in enumerate(columns, start=1):
            if column["key"] == "amount":
                value = float(detail["amount"])
            else:
                value = safe_excel_text(detail_text(detail, column))

            cell = sheet.cell(row=offset, column=index, value=value)
            cell.font = Font(size=9)
            cell.border = border

            if column["key"] == "amount":
                cell.number_format = AMOUNT_FORMAT
                cell.alignment = Alignment(horizontal="right")

    last_row = 1 + len(report["rows"])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{last_row}"

    # Los totales por estado, bajo la columna del monto. El cancelado es el
    # mismo número que el Total de Ventas de la primera pestaña; pendientes y
    # anulados no suman allí.
    amount_column = next(
        index
        for index, column in enumerate(columns, start=1)
        if column["key"] == "amount"
    )
    totals = [
        ("Total cancelado", report["sales_total"]),
        ("Total pendiente", report["pending"]["total"]),
        ("Total anulado", report["voided"]["total"]),
    ]

    for offset, (text, amount) in enumerate(totals, start=last_row + 2):
        label = sheet.cell(row=offset, column=amount_column - 1, value=text)
        label.font = Font(bold=True, size=9)
        label.alignment = Alignment(horizontal="right")

        value = sheet.cell(row=offset, column=amount_column, value=float(amount))
        value.font = Font(bold=True, size=9)
        value.number_format = AMOUNT_FORMAT
        value.alignment = Alignment(horizontal="right")


def render_excel(report, buffer):
    book = Workbook()

    _consolidated_sheet(book.active, report)
    _detail_sheet(book.create_sheet(), report)

    book.save(buffer)

    return filename(report, "xlsx")


# ---------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------

def render_pdf(report, buffer):
    """La hoja consolidada en A4 vertical, lista para firmar y archivar.

    Lleva solo el consolidado: el detalle es para cuadrar en el Excel, no
    para imprimirse.
    """
    margin = 15 * mm
    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=margin,
        rightMargin=margin,
        topMargin=margin,
        bottomMargin=margin,
        title=report["title"],
    )
    width = A4[0] - margin * 2

    small = ParagraphStyle("small", fontName="Helvetica-Bold", fontSize=7.5)
    place = ParagraphStyle(
        "place", fontName="Helvetica-Bold", fontSize=10, alignment=2
    )
    centered = ParagraphStyle(
        "centered", fontName="Helvetica-Bold", fontSize=10, alignment=1
    )
    title = ParagraphStyle(
        "title", fontName="Helvetica-Bold", fontSize=13, alignment=1, leading=16
    )

    strip = Table(
        [[
            Paragraph(escape(stamp(report)), small),
            Paragraph(escape(report["place_label"]), place),
        ]],
        colWidths=[width * 0.5, width * 0.5],
    )
    strip.setStyle(TableStyle([
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
    ]))

    story = [
        strip,
        Spacer(1, 4),
        Paragraph(report["title"].upper(), title),
        Paragraph(period_text(report), centered),
    ]

    story.append(Spacer(1, 8))

    rows = []
    styles = [
        ("LEFTPADDING", (0, 0), (-1, -1), 2),
        ("RIGHTPADDING", (0, 0), (-1, -1), 2),
        ("TOPPADDING", (0, 0), (-1, -1), 1.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ALIGN", (1, 0), (1, -1), "RIGHT"),
    ]

    for index, line in enumerate(report["lines"]):
        text, amount, style = line["text"], line["amount"], line["style"]
        rows.append([text, "" if amount is None else amount_text(amount)])

        if style in ("section", "total", "balance"):
            styles.append(("FONTNAME", (0, index), (-1, index), "Helvetica-Bold"))

        if style == "balance":
            styles.append(("BOX", (1, index), (1, index), 0.8, colors.black))

    table = Table(rows, colWidths=[width * 0.72, width * 0.28])
    table.setStyle(TableStyle(styles))
    story.append(table)

    document.build(story)

    return filename(report, "pdf")


RENDERERS = {
    "EXCEL": render_excel,
    "PDF": render_pdf,
}

CONTENT_TYPES = {
    "PDF": "application/pdf",
    "EXCEL": (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    ),
}


def render(report, export_format):
    """Dibuja el cierre en el formato pedido y devuelve (buffer, nombre)."""
    buffer = BytesIO()
    name = RENDERERS[export_format](report, buffer)
    buffer.seek(0)

    return buffer, name
